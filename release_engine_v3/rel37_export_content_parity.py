"""REL37 actual export content-parity acceptance.

Download/parse success is not content parity. This checker compares the
validated CanonicalDocument to preview, DOCX, and PDF independently.
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from xml.etree import ElementTree as ET

from release_engine_v3.rel37_canonical_document import CanonicalDocument
from release_engine_v3.rel37_professional_projection import projection_tables
from release_engine_v3.rel37_sector_context import (
    UI_SECTOR_PAIRS,
    cover_sector_from_hashed_narrative,
    environment_narrative_paragraphs,
    sector_reference_aliases,
)

_ENV_HEAD_RE = re.compile(
    r'(البيئة التنظيمية والتهديدات|البيئة والمحركات|'
    r'Regulatory Environment|Business Environment|'
    r'Environment and Drivers|Threat Landscape)',
    re.I,
)
_ENV_NEXT_RE = re.compile(
    r'(تحليل الفجوات|Gap Analysis|خارطة الطريق|Roadmap|'
    r'الركائز الاستراتيجية|Strategic Pillars)',
    re.I,
)
_APPENDIX_RE = re.compile(
    r'(Appendix|Annex|الملحق|الملاحق|ملحق)',
    re.I,
)
_TOC_LINE_RE = re.compile(r'^\d{1,2}\s+\S')
_TOC_NUM_RE = re.compile(r'^\d{1,2}(\s+)?\S')
_CHROME_RE = re.compile(
    r'(CONFIDENTIAL|Prepared by Mizan|^\s*Page\s+\d+$|'
    r'^\d{1,2}\s+[A-Za-z]+\s+\d{4}$)',
    re.I,
)
_COVER_LABELS = ('القطاع', 'Sector')
_GENERIC_ENV_AR = 'تعمل الجهة في بيئة تنظيمية'
_REL37_NARRATIVE_DOMAINS = frozenset({'data', 'ai', 'dt'})

_AR_RE = re.compile(r'[\u0600-\u06FF]')
_W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


@dataclass
class ParityResult:
    persist_success: str = 'not_run'
    download_parse_success: str = 'not_run'
    content_parity: str = 'not_run'
    language_correctness: str = 'not_run'
    rendering_readability: str = 'not_run'
    security_acceptance: str = 'not_run'
    blockers: List[str] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)
    file_sha256: str = ''
    model_hash: str = ''
    source_hash: str = ''
    details: Dict[str, Any] = field(default_factory=dict)

    @property
    def passed(self) -> bool:
        dims = (
            self.persist_success,
            self.download_parse_success,
            self.content_parity,
            self.language_correctness,
            self.rendering_readability,
        )
        return all(item == 'passed' for item in dims) and not self.blockers


def arabic_char_count(text: str, *, org_name: str = '') -> int:
    blob = str(text or '').replace(org_name or '', '')
    return len(_AR_RE.findall(blob))


def inventory_docx_bytes(raw: bytes) -> Dict[str, Any]:
    paragraphs: List[str] = []
    tables: List[List[List[str]]] = []
    headers: List[str] = []
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        for name in zf.namelist():
            if not name.startswith('word/') or not name.endswith('.xml'):
                continue
            root = ET.fromstring(zf.read(name))
            is_hf = 'header' in name or 'footer' in name
            for p in root.iter(f'{_W_NS}p'):
                if list(p.iter(f'{_W_NS}tbl')):
                    continue
                text = ''.join(
                    (t.text or '') + (t.tail or '')
                    for t in p.iter(f'{_W_NS}t')
                ).strip()
                if not text:
                    continue
                if is_hf:
                    headers.append(text)
                else:
                    paragraphs.append(text)
            for tbl in root.iter(f'{_W_NS}tbl'):
                rows: List[List[str]] = []
                for tr in tbl.findall(f'{_W_NS}tr'):
                    cells = []
                    for tc in tr.findall(f'{_W_NS}tc'):
                        cells.append(''.join(
                            (t.text or '') + (t.tail or '')
                            for t in tc.iter(f'{_W_NS}t')
                        ).strip())
                    if any(cells):
                        rows.append(cells)
                if rows:
                    tables.append(rows)
    return {
        'paragraphs': paragraphs,
        'tables': tables,
        'headers_footers': headers,
        'text': '\n'.join(paragraphs + [
            ' | '.join(cell for cell in row)
            for table in tables for row in table
        ] + headers),
    }


def _normalize_extracted_pdf_text(text: str) -> str:
    """Keep extractor output, but recover logical Arabic from presentation forms.

    This does not invent missing text. It only NFKC-normalizes glyphs that
    were actually extracted, and drops NUL placeholders left by missing
    Latin glyphs inside Arabic runs.
    """
    import unicodedata
    cleaned = unicodedata.normalize('NFKC', text or '').replace('\x00', '')
    if cleaned.startswith('\ufeff'):
        cleaned = cleaned[1:]
    return cleaned.replace('\ufeff', '')


def _decode_actual_text_hex(hx: str) -> str:
    raw_hx = str(hx or '')
    if raw_hx.upper().startswith('FEFF'):
        raw_hx = raw_hx[4:]
    try:
        return bytes.fromhex(raw_hx).decode('utf-16-be')
    except Exception:  # noqa: BLE001
        return ''


def _actual_text_spans_from_stream(stream: bytes) -> List[str]:
    parts: List[str] = []
    if not stream or b'ActualText' not in stream:
        return parts
    for match in re.finditer(br'/ActualText\s*<([0-9A-Fa-f]+)>', stream):
        text = _decode_actual_text_hex(match.group(1).decode('ascii'))
        if text:
            parts.append(text)
    return parts


def _page_content_xrefs(page) -> List[int]:
    xrefs: List[int] = []
    try:
        contents = page.get_contents() or []
        if isinstance(contents, int):
            contents = [contents]
        xrefs.extend(int(item) for item in contents if item)
    except Exception:  # noqa: BLE001
        pass
    try:
        xobjects = page.get_xobjects() or []
        for item in xobjects:
            if isinstance(item, int):
                xrefs.append(item)
                continue
            if isinstance(item, (tuple, list)) and item:
                xref = item[0]
                if isinstance(xref, int):
                    xrefs.append(xref)
    except Exception:  # noqa: BLE001
        pass
    return list(dict.fromkeys(xrefs))


def _unescape_pdf_literal(data: bytes) -> str:
    out = bytearray()
    idx = 0
    while idx < len(data):
        if data[idx] != 0x5C or idx + 1 >= len(data):
            out.append(data[idx])
            idx += 1
            continue
        nxt = data[idx + 1]
        if nxt == 0x6E:
            out.append(0x0A)
        elif nxt == 0x72:
            out.append(0x0D)
        elif nxt == 0x74:
            out.append(0x09)
        elif 0x30 <= nxt <= 0x37:
            octal = bytearray()
            look = idx + 1
            while look < len(data) and len(octal) < 3 and 0x30 <= data[look] <= 0x37:
                octal.append(data[look])
                look += 1
            out.append(int(octal.decode('ascii'), 8) & 0xFF)
            idx = look
            continue
        else:
            out.append(nxt)
        idx += 2
    try:
        return out.decode('latin-1')
    except Exception:  # noqa: BLE001
        return out.decode('utf-8', 'replace')


def _pdf_content_events(stream: bytes) -> List[Dict[str, Any]]:
    """Marked-content order: painted Tj strings and ActualText spans.

    Text render mode 3 is a non-displayed logical layer. That provenance
    is carried on the event. Matching ActualText is not enough to call
    a painted line an overlay.
    """
    events: List[Dict[str, Any]] = []
    idx = 0
    data = stream or b''
    render_mode = 0
    stack: List[int] = [0]
    current_actual = ''
    span_has_tr3 = False

    def _ws_before(pos: int) -> bool:
        return pos == 0 or data[pos - 1] in b' \t\r\n'

    while idx < len(data):
        if data.startswith(b'/ActualText', idx):
            match = re.match(br'/ActualText\s*<([0-9A-Fa-f]+)>', data[idx:])
            if match:
                current_actual = _decode_actual_text_hex(
                    match.group(1).decode('ascii'))
                events.append({
                    'kind': 'actual',
                    'text': current_actual,
                })
                idx += match.end()
                continue
        if data.startswith(b'EMC', idx) and _ws_before(idx):
            if span_has_tr3 and current_actual:
                events.append({
                    'kind': 'hidden_logical',
                    'text': current_actual,
                    'reason': 'non_displayed_render_mode_3',
                })
            current_actual = ''
            span_has_tr3 = False
            idx += 3
            continue
        tr = re.match(br'(\d+)\s+Tr\b', data[idx:idx + 8])
        if tr and _ws_before(idx):
            render_mode = int(tr.group(1))
            if render_mode == 3:
                span_has_tr3 = True
            idx += tr.end()
            continue
        if data[idx:idx + 1] == b'q' and _ws_before(idx) and (
                idx + 1 >= len(data) or data[idx + 1] in b' \t\r\n'):
            stack.append(render_mode)
            idx += 1
            continue
        if data[idx:idx + 1] == b'Q' and _ws_before(idx) and (
                idx + 1 >= len(data) or data[idx + 1] in b' \t\r\n'):
            render_mode = stack.pop() if stack else 0
            if not stack:
                stack = [0]
            idx += 1
            continue
        if data[idx] == 0x28:
            cursor = idx + 1
            buf = bytearray()
            while cursor < len(data):
                if data[cursor] == 0x5C and cursor + 1 < len(data):
                    buf.append(0x5C)
                    buf.append(data[cursor + 1])
                    cursor += 2
                    continue
                if data[cursor] == 0x29:
                    break
                buf.append(data[cursor])
                cursor += 1
            rest = data[cursor + 1:cursor + 16].lstrip()
            if rest.startswith(b'Tj') or rest.startswith(b"'") or rest.startswith(b'"'):
                events.append({
                    'kind': 'visible',
                    'text': _unescape_pdf_literal(bytes(buf)),
                    'non_displayed': render_mode == 3,
                    'render_mode': render_mode,
                })
            idx = cursor + 1
            continue
        if data[idx] == 0x3C:
            match = re.match(br'<([0-9A-Fa-f]+)>', data[idx:])
            if match:
                rest = data[idx + match.end():idx + match.end() + 16].lstrip()
                hex_body = match.group(1).decode('ascii')
                if (
                        rest.startswith(b'Tj')
                        and hex_body.upper().startswith('FEFF')
                ):
                    events.append({
                        'kind': 'visible',
                        'text': _decode_actual_text_hex(hex_body),
                        'logical_paint': True,
                        'non_displayed': render_mode == 3,
                        'render_mode': render_mode,
                    })
                idx += match.end()
                continue
        idx += 1
    return events


def _hidden_logicals_from_events(events: Sequence[Dict[str, Any]]) -> List[str]:
    """Logical strings proven non-displayed by render mode 3 in-file."""
    hidden: List[str] = []
    for event in events:
        text = str(event.get('text') or '').strip()
        if not text:
            continue
        if event.get('kind') == 'hidden_logical':
            hidden.append(text)
            continue
        if event.get('kind') == 'visible' and event.get('non_displayed'):
            if any('\u0600' <= ch <= '\u06FF' for ch in text) or any(
                    ch.isalpha() for ch in text):
                hidden.append(text)
    return hidden


def _page_line_blocks(page, *, ignore_actualtext: bool) -> List[Dict[str, Any]]:
    import pymupdf
    flags = int(getattr(pymupdf, 'TEXTFLAGS_DICT', 0) or 0)
    ignore_at = int(getattr(pymupdf, 'TEXT_IGNORE_ACTUALTEXT', 0) or 0)
    if ignore_actualtext and ignore_at:
        flags |= ignore_at
    payload = page.get_text('dict', flags=flags) or {}
    blocks: List[Dict[str, Any]] = []
    for block in payload.get('blocks') or []:
        for line in block.get('lines') or []:
            text = ''.join(
                str(span.get('text') or '') for span in line.get('spans') or [])
            bbox = line.get('bbox') or (0, 0, 0, 0)
            if not str(text).strip():
                continue
            blocks.append({
                'text': text,
                'y0': float(bbox[1]),
                'y1': float(bbox[3]),
                'x0': float(bbox[0]),
            })
    return blocks


def _align_painted_and_actual(
        painted: Sequence[Dict[str, Any]],
        logical: Sequence[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    aligned: List[Dict[str, Any]] = []
    used = set()
    for item in painted:
        match_idx = None
        best = 1e9
        for idx, other in enumerate(logical):
            if idx in used:
                continue
            dist = abs(float(other['y0']) - float(item['y0'])) + (
                abs(float(other['x0']) - float(item['x0'])) * 0.25)
            if dist < best:
                best = dist
                match_idx = idx
        actual = ''
        if match_idx is not None and best <= 10:
            actual = str(logical[match_idx].get('text') or '')
            used.add(match_idx)
        aligned.append({
            'visible': item.get('text') or '',
            'actual': actual,
            'y0': item.get('y0'),
            'y1': item.get('y1'),
        })
    return aligned


def extract_pdf_pages(raw: bytes) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Page-bound painted text and ActualText. Never a global bag.

    Painted text ignores ActualText. Logical ActualText is aligned to the
    same line boxes. Cover is page 0.
    """
    pages: List[Dict[str, Any]] = []
    meta: Dict[str, Any] = {
        'pages': 0,
        'reliable': False,
        'associated': False,
        'extractor': '',
        'actual_text_spans': 0,
    }
    if not raw.startswith(b'%PDF'):
        return pages, meta
    try:
        import pymupdf
        doc = pymupdf.open(stream=raw, filetype='pdf')
        meta['pages'] = len(doc)
        span_count = 0
        ignore_at = int(getattr(pymupdf, 'TEXT_IGNORE_ACTUALTEXT', 0) or 0)
        text_flags = int(getattr(pymupdf, 'TEXTFLAGS_TEXT', 0) or 0) | ignore_at
        for index, page in enumerate(doc):
            if ignore_at:
                visible = page.get_text('text', flags=text_flags) or ''
            else:
                visible = page.get_text() or ''
            painted = _page_line_blocks(page, ignore_actualtext=True)
            logical = _page_line_blocks(page, ignore_actualtext=False)
            actual_parts: List[str] = []
            events: List[Dict[str, str]] = []
            for xref in _page_content_xrefs(page):
                try:
                    stream = doc.xref_stream(xref)
                except Exception:  # noqa: BLE001
                    continue
                events.extend(_pdf_content_events(stream))
                spans = _actual_text_spans_from_stream(stream)
                span_count += len(spans)
                actual_parts.extend(spans)
            pages.append({
                'index': index,
                'visible': _normalize_extracted_pdf_text(visible),
                'actual': _normalize_extracted_pdf_text('\n'.join(actual_parts)),
                'events': events,
                'blocks': _align_painted_and_actual(painted, logical),
            })
        meta['actual_text_spans'] = span_count
        meta['extractor'] = 'pymupdf+page-actualtext+line-boxes'
        meta['associated'] = bool(pages)
        meta['reliable'] = bool(pages) and any(
            (item['visible'] or item['actual']).strip() for item in pages)
        return pages, meta
    except Exception as exc:  # noqa: BLE001
        meta['extractor_error'] = type(exc).__name__
        return pages, meta


def _extract_pdf_actual_text(raw: bytes) -> Tuple[str, int]:
    """Read producer-emitted PDF ActualText spans, page-associated first.

    Global xref fallback is retained only for inventory/readability, not
    for narrative or cover comparison.
    """
    pages, meta = extract_pdf_pages(raw)
    associated = '\n'.join(
        item['actual'] for item in pages if item.get('actual'))
    if associated.strip():
        return associated, int(meta.get('actual_text_spans') or 0)
    parts: List[str] = []
    try:
        import pymupdf
        doc = pymupdf.open(stream=raw, filetype='pdf')
        for xref in range(1, doc.xref_length()):
            try:
                stream = doc.xref_stream(xref)
            except Exception:  # noqa: BLE001
                continue
            parts.extend(_actual_text_spans_from_stream(stream))
    except Exception:  # noqa: BLE001
        return '', 0
    return '\n'.join(parts), len(parts)


def extract_pdf_text(raw: bytes) -> Tuple[str, Dict[str, Any]]:
    meta: Dict[str, Any] = {'pages': 0, 'extractor': '', 'reliable': False}
    if not raw.startswith(b'%PDF'):
        return '', meta
    actual, actual_count = _extract_pdf_actual_text(raw)
    meta['actual_text_spans'] = actual_count
    try:
        import pymupdf
        doc = pymupdf.open(stream=raw, filetype='pdf')
        pages = [page.get_text() or '' for page in doc]
        meta['pages'] = len(doc)
        meta['extractor'] = 'pymupdf+actualtext' if actual else 'pymupdf'
        text = _normalize_extracted_pdf_text('\n'.join(pages + [actual]))
        meta['reliable'] = bool(text.strip()) and meta['pages'] > 0
        return text, meta
    except Exception as exc:  # noqa: BLE001
        meta['extractor_error'] = type(exc).__name__
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        pages = [(page.extract_text() or '') for page in reader.pages]
        meta['pages'] = len(reader.pages)
        meta['extractor'] = 'pypdf2+actualtext' if actual else 'pypdf2'
        text = _normalize_extracted_pdf_text('\n'.join(pages + [actual]))
        meta['reliable'] = bool(text.strip()) and meta['pages'] > 0
        return text, meta
    except Exception as exc:  # noqa: BLE001
        meta['extractor_error'] = type(exc).__name__
        if actual:
            meta['extractor'] = 'actualtext'
            text = _normalize_extracted_pdf_text(actual)
            meta['reliable'] = bool(text.strip())
            return text, meta
        return '', meta


def _norm(value: Any) -> str:
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def _layout_norm(value: Any) -> str:
    """Whitespace, presentation forms, and mixed-script word boundaries.

    Does not drop Latin tokens or rewrite their identity.
    """
    import unicodedata
    text = unicodedata.normalize('NFKC', str(value or '')).replace('\x00', '')
    text = re.sub(r'([A-Za-z0-9])(?=[\u0600-\u06FF])', r'\1 ', text)
    text = re.sub(r'([\u0600-\u06FF])(?=[A-Za-z0-9])', r'\1 ', text)
    return re.sub(r'\s+', ' ', text).strip()


_AR_RUN_CHAR_RE = re.compile(
    r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF،؛؟ـ]'
)
_MIXED_RUN_RE = re.compile(
    r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF،؛؟ـ\s]+'
    r'|[^\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF،؛؟ـ\s]+'
)


def _mixed_run_kind(char: str) -> str:
    if char.isspace() or char in '.,;:|/()-[]{}':
        return 'sp'
    if _AR_RUN_CHAR_RE.match(char):
        return 'ar'
    return 'lat'


def _mixed_script_runs(text: str) -> List[str]:
    """Same-script runs. Internal spaces stay with the current script.

    Treating whitespace as Arabic split ``REL33 P1 Data Management Org``
    into per-word runs and reversed them. Latin phrases stay one run.
    """
    runs: List[str] = []
    buf: List[str] = []
    current = None
    pending_sp = ''
    for char in str(text or ''):
        kind = _mixed_run_kind(char)
        if kind == 'sp':
            pending_sp += char
            continue
        if current is None:
            buf.append(pending_sp + char)
            pending_sp = ''
            current = kind
            continue
        if kind == current:
            buf.append(pending_sp + char)
            pending_sp = ''
            continue
        runs.append(''.join(buf))
        buf = [pending_sp + char]
        pending_sp = ''
        current = kind
    if buf:
        runs.append(''.join(buf) + pending_sp)
    elif pending_sp:
        runs.append(pending_sp)
    return [run for run in runs if run]


def _undo_visual_rtl_line(line: str) -> str:
    """Recover logical run order from a visual-LTR extraction of RTL text.

    Arabic-only lines stay unchanged. Latin-only lines stay unchanged.
    Mixed lines reverse script runs, not words inside a Latin phrase.
    """
    text = str(line or '')
    if not any('\u0600' <= ch <= '\u06FF' for ch in text):
        return text
    if not any(ch.isascii() and (ch.isalpha() or ch.isdigit()) for ch in text):
        return text
    runs = _mixed_script_runs(text)
    if len(runs) <= 1:
        return text
    return ''.join(reversed(runs))


def _is_mixed_arabic_latin(text: str) -> bool:
    raw = str(text or '')
    has_ar = any('\u0600' <= ch <= '\u06FF' for ch in raw)
    has_lat = any(ch.isascii() and ch.isalpha() for ch in raw)
    return has_ar and has_lat


def _undo_visible_lines(text: str) -> str:
    return '\n'.join(
        _undo_visual_rtl_line(line) for line in str(text or '').splitlines())


_DECIMAL_GUARD = '\u241E'
_DECIMAL_NUMBER_RE = re.compile(
    r'(?<![\w])(\d+)[.\u066B](\d+)(?![\w])'
)
_COMPLETE_DECIMAL_RE = re.compile(r'^\d+[.\u066B]\d+%?$')


def _protect_numeric_decimals(text: str) -> str:
    """Keep complete numeric identity, including decimal components."""
    return _DECIMAL_NUMBER_RE.sub(
        lambda match: f'{match.group(1)}{_DECIMAL_GUARD}{match.group(2)}',
        str(text or ''),
    )


def _restore_numeric_decimals(text: str) -> str:
    return str(text or '').replace(_DECIMAL_GUARD, '.')


def _is_complete_numeric_value(word: str) -> bool:
    raw = str(word or '').strip()
    if raw.endswith('%') and raw[:-1]:
        raw = raw[:-1]
    return bool(_COMPLETE_DECIMAL_RE.match(raw))


def _content_words(text: str) -> List[str]:
    """Layout-normalized words, punctuation split away. Digits stay.

    A decimal value such as ``10.33`` stays one token. Splitting it into
    ``10`` and ``33`` would let leftover identity digits authorize a
    changed quantitative value.
    """
    cleaned = _protect_numeric_decimals(_layout_norm(text))
    cleaned = re.sub(r'[^\w\u0600-\u06FF%\u241E]+', ' ', cleaned)
    words: List[str] = []
    for word in cleaned.split():
        word = _restore_numeric_decimals(word)
        if re.sub(r'[\u060C\u061B\u061F\u0640]', '', word):
            words.append(word)
    return words


def _undo_extracted_arabic_visual(text: str) -> str:
    """Logical words from LTR-extracted Arabic presentation order.

    Each Arabic word is un-reversed, then Arabic words are restored
    right-to-left. Latin tokens keep their extracted identity.
    """
    arabic: List[str] = []
    latin: List[str] = []
    for word in _layout_norm(text).split():
        if any('\u0600' <= ch <= '\u06FF' for ch in word):
            arabic.append(word[::-1])
        else:
            latin.append(word)
    if not arabic:
        return _layout_norm(text)
    return _layout_norm(' '.join(list(reversed(arabic)) + latin))


def _same_complete_words(painted: str, actual: str) -> bool:
    """True when layout-normalized words match identity, multiplicity, and order.

    A complete bag of words is not a complete ordered statement.
    """
    want = _content_words(painted)
    have = _content_words(actual)
    return bool(want) and bool(have) and want == have


def _complete_layout_equivalent(painted: str, actual: str) -> bool:
    """True only when painted content is the complete associated text.

    Whitespace, line boundaries, presentation forms, and a proven
    mixed-script visual-to-logical conversion may differ. Sorted word
    lists, set inclusion, opening-word prefixes, overlap percentages,
    digit deletion, and acronym membership without the surrounding
    sentence are not enough.
    """
    vis = _layout_norm(painted)
    act = _layout_norm(actual)
    if vis and act and vis == act:
        return True
    if not (vis and act):
        return False
    candidates = [_undo_extracted_arabic_visual(painted)]
    if not _arabic_already_logical(painted, actual):
        candidates.extend((
            _layout_norm(_undo_visible_lines(painted)),
            _undo_extracted_arabic_visual(_undo_visible_lines(painted)),
        ))
    act_words = _content_words(act)
    if not act_words:
        return False
    if (
            not _arabic_already_logical(painted, actual)
            and _mixed_script_layout_equivalent(painted, actual)
    ):
        return True
    for candidate in candidates:
        if not candidate:
            continue
        if candidate == act:
            return True
        if _content_words(candidate) == act_words:
            return True
    return False


def _token_only_latin_overlay(text: str) -> bool:
    """Latin identity tokens only; not an ordinary painted sentence."""
    vis = _layout_norm(text)
    if not vis or any('\u0600' <= ch <= '\u06FF' for ch in vis):
        return False
    line_lat = _latin_relation(vis)
    if not line_lat:
        return False
    return all(
        (word in line_lat) or not any(ch.isalpha() for ch in word)
        for word in vis.split()
    )


def _arabic_word_overlap(text: str, actual: str) -> float:
    words = [
        word for word in _content_words(text)
        if any('\u0600' <= ch <= '\u06FF' for ch in word)
    ]
    have = {
        word for word in _content_words(actual)
        if any('\u0600' <= ch <= '\u06FF' for ch in word)
    }
    if not words:
        return 0.0
    return sum(1 for word in words if word in have) / len(words)


_LEADING_IDENTITY_RE = re.compile(
    r'^(?:[A-Z]{2,}[A-Za-z0-9_-]*\s*)+'
)


def _split_visual_reverse_prefix(text: str, actual: str) -> List[str]:
    """Separate a character-reversed leftover from later logical paint.

    Official Arabic cells mash the extracted visual run onto the
    invisible logical draw. The leftover is a presentation of ActualText,
    not a new sentence. A logical sentence missing words is not split
    away from itself because it is not a contiguous ActualText span.
    """
    raw = str(text or '')
    act = _layout_norm(actual)
    if len(raw) < 24 or not act:
        return [raw] if raw else []
    for index in range(8, len(raw) - 11):
        prefix, rest = raw[:index], raw[index:]
        rest_norm = _layout_norm(rest)
        rest_core = _LEADING_IDENTITY_RE.sub('', rest_norm).strip()
        if len(rest_core) < 16 or rest_core not in act:
            continue
        prefix_norm = _layout_norm(prefix)
        prefix_rev = _layout_norm(prefix[::-1])
        prefix_undo = _undo_extracted_arabic_visual(prefix)
        visual = (
                len(prefix_rev) >= 12
                and (
                    prefix_rev in act
                    or rest_core.startswith(prefix_rev[:16])
                    or prefix_rev.startswith(rest_core[:16])
                )
        ) or (
                prefix_undo
                and prefix_undo != prefix_norm
                and (prefix_undo in act or rest_core.startswith(prefix_undo[:16]))
        )
        if visual and prefix_norm not in rest_core:
            return [part for part in (prefix, rest) if _layout_norm(part)]
    return [raw]


def _overlay_line_fragments(text: str, actual: str) -> List[str]:
    """Recover associated paint fragments from one extracted line."""
    raw = str(text or '')
    if not raw:
        return []
    parts: List[str] = []
    for sentence in re.split(r'(?<=[.۔])(?=\S)', raw):
        if not _layout_norm(sentence):
            continue
        parts.extend(_split_visual_reverse_prefix(sentence, actual))
    return parts or [raw]


def _non_identity_words(text: str) -> List[str]:
    return [
        word for word in _content_words(text)
        if not _is_identity_latin_token(word)
    ]


def _mixed_script_layout_equivalent(painted: str, actual: str) -> bool:
    """Arabic visual undo plus independent Latin identity/order.

    Helvetica identity tokens are a separate draw. They do not license
    ignoring the surrounding ordered Arabic statement or values.
    """
    undone = _undo_extracted_arabic_visual(painted)
    if _non_identity_words(undone) != _non_identity_words(actual):
        return False
    if not _non_identity_words(actual):
        return False
    return _latin_relation(painted) == _latin_relation(actual)


def _arabic_already_logical(painted: str, actual: str) -> bool:
    """True when Arabic words are already logical, so run-undo must not rewrite Latin order."""
    want = [
        word for word in _content_words(actual)
        if any('\u0600' <= ch <= '\u06FF' for ch in word)
    ]
    have = [
        word for word in _content_words(painted)
        if any('\u0600' <= ch <= '\u06FF' for ch in word)
    ]
    complete, _extras = _walk_complete_statement(
        want, have, set(want), actual, allow_out_of_order_spans=False)
    return complete


def _is_identity_latin_token(word: str) -> bool:
    tokens = _semantic_latin_tokens(word)
    return tokens == [word]


def _compact_letters(text: str) -> str:
    return re.sub(r'[^\w\u0600-\u06FF%]+', '', _layout_norm(text))


def _bare_word(word: str) -> str:
    """Word identity without layout punctuation. Digits, decimals, and % stay."""
    text = _protect_numeric_decimals(str(word or ''))
    text = re.sub(r'[^\w\u0600-\u06FF%\u241E]+', '', text)
    text = _restore_numeric_decimals(text)
    return re.sub(r'[\u060C\u061B\u061F\u0640]', '', text)


def _identity_digit_remnants(actual: str) -> set:
    """Leftover digits of omitted Helvetica identity tokens (REL33→33, P1→1).

    These are presentation chips of mixed letter-digit identifiers.
    They are not permission to drop a quantitative value, and a decimal
    component is not a separate harmless chip.
    """
    remnants = set()
    for token in _semantic_latin_tokens(actual):
        if not any(ch.isalpha() for ch in token):
            continue
        digits = ''.join(ch for ch in token if ch.isdigit())
        if digits:
            remnants.add(digits)
    return remnants


def _is_quantitative_number_word(word: str) -> bool:
    """True for a quantitative integer or decimal, not an identifier token.

    ``10`` and ``10.33`` are values. ``REL33`` is not. A standalone
    integer is real content until overlay provenance proves otherwise.
    """
    if _is_complete_numeric_value(word):
        return True
    compact = _bare_word(word)
    if not compact or not compact.isdigit():
        return False
    return not _is_identity_latin_token(word)


def _is_identity_digit_remnant(word: str, actual: str) -> bool:
    """Standalone leftover identifier digits, never a complete numeric value.

    Text equality with digits inside REL33/P1 is not overlay provenance
    and is not permission to drop the chip. Callers must associate the
    chip with its source identifier token in the same extracted run.
    """
    if _is_complete_numeric_value(word):
        return False
    compact = _bare_word(word)
    if not compact or not compact.isdigit():
        return False
    return compact in _identity_digit_remnants(actual)


def _identifier_owns_digit_chip(token: str, chip: str) -> bool:
    """True when chip digits are this identifier token's leftover draw.

    ``P1`` owns ``1``. ``REL33`` owns ``33``. A preceding Arabic word
    or another quantitative integer does not. This is run-adjacency to
    the source identifier, not file font/draw/span proof.
    """
    if not token or not chip:
        return False
    if not _is_identity_latin_token(token):
        return False
    digits = ''.join(ch for ch in _bare_word(token) if ch.isdigit())
    return bool(digits) and digits == _bare_word(chip)


def _is_helvetica_overlay_word(word: str) -> bool:
    """True for separately drawn Latin overlay chips of identity tokens.

    REL33/P1 stay identity tokens. Title-case org words that already
    appear in ActualText may be leftover overlay chips. A quantitative
    number, ``not``, or ``must`` is not discardable merely because it
    is Latin or shares digits with an identifier.
    """
    compact = _bare_word(word)
    if not compact:
        return False
    if _is_complete_numeric_value(word) or compact.isdigit():
        return False
    lowered = compact.lower()
    if lowered in {'not', 'must', 'no', 'never'}:
        return False
    if _is_identity_latin_token(compact):
        return True
    return any(ch.isascii() and ch.isalpha() for ch in compact)


def _compact_without_identity(text: str, actual: str = '') -> str:
    """Ordered letters/digits after dropping Helvetica overlay and punctuation.

    Official leftover runs omit separately drawn Helvetica tokens and may
    keep only their leftover digits. Skipping those tokens is not
    permission to drop values or reorder the statement.
    """
    remnants = _identity_digit_remnants(actual)
    kept: List[str] = []
    for word in _content_words(text):
        bare = _bare_word(word)
        if not bare or _is_helvetica_overlay_word(word):
            continue
        if bare in remnants:
            prev = kept[-1] if kept else ''
            if _identifier_owns_digit_chip(prev, bare):
                continue
            kept.append(bare)
            continue
        kept.append(bare)
    return ''.join(kept)


def _arabic_words(text: str | Sequence[str]) -> List[str]:
    words = text if isinstance(text, (list, tuple)) else _content_words(text)
    return [
        word for word in words
        if any('\u0600' <= ch <= '\u06FF' for ch in word)
    ]


def _arabic_logical_subsequence(painted: str, actual: str) -> bool:
    """True when painted Arabic words already occur in ActualText order."""
    have = [_bare_word(word) for word in _arabic_words(painted)]
    want = [_bare_word(word) for word in _arabic_words(actual)]
    if not have or not want:
        return False
    idx = 0
    for word in have:
        while idx < len(want) and want[idx] != word:
            idx += 1
        if idx >= len(want):
            return False
        idx += 1
    return True


def _contiguous_span_index(run: Sequence[str], want: Sequence[str]) -> int:
    """Start index in want when run is a contiguous 3+ word span."""
    needles = [_bare_word(word) for word in run]
    hay = [_bare_word(word) for word in want]
    size = len(needles)
    if size < 3 or any(not item for item in needles):
        return -1
    for index in range(len(hay) - size + 1):
        if hay[index:index + size] == needles:
            return index
    return -1


def _is_short_reverse_debris(word: str, known: Sequence[str] | set) -> bool:
    """Two-letter reverse chip of a known word, not a free 2-letter match.

    Official leftover extraction can leave the last two letters of a
    reversed word. A two-letter particle that only occurs inside a longer
    reverse is not debris.
    """
    compact = _bare_word(word)
    if len(compact) != 2:
        return False
    for item in known:
        rev = _bare_word(item)[::-1]
        if len(rev) >= 3 and (rev.startswith(compact) or rev.endswith(compact)):
            return True
    return False


def _is_reverse_fragment(word: str, known: Sequence[str] | set) -> bool:
    compact = _bare_word(word)
    if not compact:
        return False
    known_bare = {_bare_word(item) for item in known if _bare_word(item)}
    if compact[::-1] in known_bare:
        return True
    for item in known_bare:
        rev = item[::-1]
        if len(compact) >= 3 and len(rev) >= 3 and (
                compact in rev or rev in compact):
            return True
    return False


def _arabic_leftover_rotation_of_actual(text: str, actual: str) -> bool:
    """True when a leftover run is the visual reverse of an ActualText prefix.

    Official mashed leftovers omit separately drawn Helvetica identity and
    keep those tokens' leftover digits inside the reversed Arabic
    presentation. After undo, the first ActualText Arabic word sits at
    the end and the rest is the following ActualText Arabic prefix.
    Leftover identity chips are Latin and are not part of this Arabic
    rotation. One-letter reverse debris of a known word is not a new
    statement.

    This associates leftover chips with the leftover run they were
    extracted from. It is not previous-word permission to drop a
    quantitative integer, and it is not font/draw/span proof.
    """
    undone = _undo_extracted_arabic_visual(text)
    known = _content_words(actual)
    words: List[str] = []
    for word in _arabic_words(undone):
        bare = _bare_word(word)
        if not bare:
            continue
        if len(bare) == 1 and any(
                len(_bare_word(item)) >= 3
                and _bare_word(item)[::-1].endswith(bare)
                for item in known):
            continue
        if _is_short_reverse_debris(word, known):
            continue
        words.append(bare)
    act_ar = [_bare_word(word) for word in _arabic_words(actual)]
    if len(words) < 3 or not act_ar:
        return False
    if words[-1] == act_ar[0] and words[:-1] == act_ar[1:len(words)]:
        return True
    if words == act_ar[:len(words)] or words[::-1] == act_ar[:len(words)]:
        return True
    return False


def _is_visual_leftover_run(text: str, actual: str) -> bool:
    """True when a run is a visual conversion of ActualText, not the logical run."""
    act = _layout_norm(actual)
    raw = _layout_norm(text)
    if not raw or not act or raw == act or raw in act:
        return False
    if _arabic_logical_subsequence(raw, act):
        return False
    raw_rev = _layout_norm(str(text or '')[::-1])
    act_c = _compact_without_identity(act, act)
    raw_rev_c = _compact_without_identity(raw_rev, act)
    if raw_rev == act or (len(raw) >= 12 and raw_rev in act):
        return True
    if len(raw_rev_c) >= 12 and raw_rev_c in act_c:
        return True
    candidates = (
        _undo_extracted_arabic_visual(text),
    )
    for candidate in candidates:
        if not candidate or candidate == raw:
            continue
        if candidate == act:
            return True
        if len(candidate) >= 12 and candidate in act:
            return True
        cand_c = _compact_without_identity(candidate, act)
        if len(cand_c) >= 12 and cand_c in act_c:
            return True
        if _content_words(candidate) == _content_words(act):
            return True
    if _arabic_leftover_rotation_of_actual(text, actual):
        return True
    return False


def _advance_covered(idx: int, covered: Sequence[bool]) -> int:
    while idx < len(covered) and covered[idx]:
        idx += 1
    return idx


def _walk_complete_statement(
        want: Sequence[str],
        have: Sequence[str],
        act_set: Sequence[str] | set,
        actual: str = '',
        *,
        allow_out_of_order_spans: bool = True,
) -> Tuple[bool, List[str]]:
    """Walk have consuming want in order. Return (complete, extras).

    A leftover visual run may omit separately drawn identity Latin.
    Classify that leftover run before consuming a reverse chip of the
    next expected word, so leftover identity digits stay with the
    leftover presentation instead of being treated as a new value.
    A contiguous 3+ word ActualText span may appear out of wrap order.
    Single-word relocation is not a span and does not reorder the statement.
    """
    known = set(act_set)
    ref = _layout_norm(actual) or ' '.join(want)
    want_words = list(want)
    covered = [False] * len(want_words)
    idx = 0
    consumed: List[str] = []
    extras: List[str] = []
    i = 0
    while i < len(have):
        word = have[i]
        idx = _advance_covered(idx, covered)
        if idx < len(want_words) and _bare_word(word) == _bare_word(want_words[idx]):
            consumed.append(want_words[idx])
            covered[idx] = True
            idx = _advance_covered(idx + 1, covered)
            i += 1
            continue
        skipped = False
        for end in range(len(have), i, -1):
            run = have[i:end]
            if not run:
                continue
            if idx < len(want_words) and _bare_word(want_words[idx]) in {
                    _bare_word(item) for item in run}:
                continue
            if _is_visual_leftover_run(' '.join(run), ref):
                extras.extend(run)
                i = end
                skipped = True
                break
            span_at = (
                _contiguous_span_index(run, want_words)
                if allow_out_of_order_spans else -1
            )
            if span_at >= 0 and not all(covered[span_at:span_at + len(run)]):
                for offset in range(len(run)):
                    covered[span_at + offset] = True
                extras.extend(run)
                idx = _advance_covered(idx, covered)
                i = end
                skipped = True
                break
            if allow_out_of_order_spans and _contiguous_span_index(
                    run, _content_words(ref)) >= 0:
                extras.extend(run)
                i = end
                skipped = True
                break
        if skipped:
            continue
        if idx < len(want_words) and _bare_word(word)[::-1] == _bare_word(want_words[idx]):
            consumed.append(want_words[idx])
            covered[idx] = True
            idx = _advance_covered(idx + 1, covered)
            i += 1
            continue
        if _bare_word(word) not in {_bare_word(item) for item in known} and (
                _bare_word(word)[::-1] in {_bare_word(item) for item in known}):
            extras.append(word)
            i += 1
            continue
        if word in consumed and word in known:
            extras.append(word)
            i += 1
            continue
        if _is_identity_latin_token(word) and (
                word in known or word in _content_words(ref)):
            extras.append(word)
            for pos, item in enumerate(want_words):
                if not covered[pos] and item == word:
                    covered[pos] = True
                    break
            idx = _advance_covered(idx, covered)
            i += 1
            continue
        if _is_identity_digit_remnant(word, ref):
            prev_have = have[i - 1] if i else ''
            prev_consumed = consumed[-1] if consumed else ''
            if (
                    _identifier_owns_digit_chip(prev_have, word)
                    or _identifier_owns_digit_chip(prev_consumed, word)
            ):
                extras.append(word)
                i += 1
                continue
        extras.append(word)
        i += 1
        if idx < len(want_words):
            return False, extras
    idx = _advance_covered(idx, covered)
    return idx == len(want_words) and all(covered), extras


def _consume_complete_statement(
        want: Sequence[str],
        have: Sequence[str],
        act_set: Sequence[str] | set,
        actual: str = '',
) -> bool:
    """Consume want in order. Skip leftover visual runs and duplicate identity.

    An interrupting word that changes the statement fails. A complete
    bag, opening clause, or prefix is not enough. A later start at the
    first expected word is allowed so a preceding leftover of another
    environment paragraph does not hide a complete later paint.
    """
    if not want:
        return False
    starts = [0]
    target = _bare_word(want[0])
    for index, word in enumerate(have):
        if index == 0:
            continue
        if _bare_word(word) == target:
            starts.append(index)
    for start in starts:
        complete, _extras = _walk_complete_statement(
            want, have[start:], act_set, actual)
        if complete:
            return True
    return False


def _extras_are_accounted(
        extras: Sequence[str],
        act_set: Sequence[str] | set,
        actual: str,
) -> bool:
    known = set(act_set)
    if not extras:
        return True
    if all(
            word in known
            or word[::-1] in known
            or _is_identity_latin_token(word)
            or _is_identity_digit_remnant(word, actual)
            or _is_reverse_fragment(word, known)
            or _is_visual_leftover_run(word, actual)
            for word in extras
    ):
        return True
    extra_text = ' '.join(extras)
    if _is_visual_leftover_run(extra_text, actual):
        return True
    leftovers = [
        word for word in extras
        if word not in known and not _is_identity_latin_token(word)
    ]
    if leftovers and _is_visual_leftover_run(' '.join(leftovers), actual):
        return True
    act_words = _content_words(actual)
    residue: List[str] = []
    i = 0
    extra_words = list(extras)
    while i < len(extra_words):
        skipped = False
        for end in range(len(extra_words), i, -1):
            run = extra_words[i:end]
            if _is_visual_leftover_run(' '.join(run), actual):
                i = end
                skipped = True
                break
            if _contiguous_span_index(run, act_words) >= 0:
                i = end
                skipped = True
                break
        if skipped:
            continue
        residue.append(extra_words[i])
        i += 1
    if not residue:
        return True
    if all(
            word in known
            or word[::-1] in known
            or _is_identity_latin_token(word)
            or _is_identity_digit_remnant(word, actual)
            or _is_reverse_fragment(word, known)
            or _is_visual_leftover_run(word, actual)
            or _is_short_reverse_debris(word, known)
            for word in residue
    ):
        return True
    return False


def _ordered_statement_present(para: str, painted: str, actual: str = '') -> bool:
    """True when the complete ordered statement survives justified mapping."""
    want = _content_words(para)
    if not want:
        return False
    ref = actual or para
    sources = [painted, _undo_extracted_arabic_visual(painted)]
    arabic_already_logical = _arabic_already_logical(painted, para)
    if not arabic_already_logical:
        sources.extend((
            _undo_visible_lines(painted),
            _undo_extracted_arabic_visual(_undo_visible_lines(painted)),
        ))
    if (
            not arabic_already_logical
            and _mixed_script_layout_equivalent(painted, para)
    ):
        return True
    known = set(want)
    for source in sources:
        if _consume_complete_statement(
                want, _content_words(source), known, ref):
            return True
    return False


def _complete_paragraph_in_painted(
        para: str, painted: str, actual: str = '',
) -> bool:
    """True when the associated paint still has this paragraph's content.

    The complete ordered statement must remain. Sharing an opening
    clause, a word bag, or a character prefix is not enough.
    """
    want = _layout_norm(para)
    have = _layout_norm(painted)
    if want and have and want in have:
        return True
    if want and have and _complete_layout_equivalent(painted, para):
        return True
    return _ordered_statement_present(para, painted, actual)


def _leftover_visual_of_actual(painted: str, para: str, actual: str) -> bool:
    """True when extra paint besides the complete paragraph is leftover visual of AT."""
    have = _layout_norm(painted)
    want = _layout_norm(para)
    act = _layout_norm(actual)
    if not have:
        return False
    if _complete_layout_equivalent(painted, para) or _complete_layout_equivalent(painted, act):
        return True
    if not _arabic_already_logical(painted, para) and _mixed_script_layout_equivalent(painted, act or para):
        return True
    sources = [have, _undo_extracted_arabic_visual(painted)]
    if not _arabic_already_logical(painted, para):
        sources.append(_layout_norm(_undo_visible_lines(painted)))
    act_set = set(_content_words(actual) or _content_words(para))
    for source in sources:
        if want and source and want in source:
            remainder = (
                source[:source.find(want)] + source[source.find(want) + len(want):]
            ).strip()
            if not remainder:
                return True
            if _is_visual_presentation_of_actual(remainder, act):
                return True
            extras = [
                word for word in _content_words(remainder)
                if word not in act_set
                and word[::-1] not in act_set
                and not _is_identity_latin_token(word)
                and not _is_reverse_fragment(word, act_set)
                and not _is_short_reverse_debris(word, act_set)
            ]
            if not extras:
                return True
            if _is_visual_leftover_run(' '.join(extras), act):
                return True
            if len(_layout_norm(remainder)) < 20 and not _semantic_latin_tokens(remainder):
                return True
    want_words = _content_words(para)
    if not want_words:
        return False
    ref = actual or para
    target = _bare_word(want_words[0])
    for source in sources:
        have_words = _content_words(source)
        starts = [0]
        for index, word in enumerate(have_words):
            if index == 0:
                continue
            if _bare_word(word) == target:
                starts.append(index)
        for start in starts:
            complete, extras = _walk_complete_statement(
                want_words, have_words[start:], act_set, ref)
            prefix = have_words[:start]
            accounted = list(prefix) + list(extras)
            if complete and _extras_are_accounted(accounted, act_set, ref):
                return True
    return False


def _unrelated_painted_sentence(painted: str, para: str, actual: str) -> bool:
    """True when paint is a different sentence, not leftover presentation.

    Visual-reversed leftovers of ActualText are not a new sentence.
    A complete statement plus contradictory additional visible text is.
    """
    have = _layout_norm(painted)
    if not have:
        return False
    if _leftover_visual_of_actual(painted, para, actual):
        return False
    if _complete_paragraph_in_painted(para, painted):
        return True
    extras = [
        word for word in _content_words(painted)
        if word not in set(_content_words(actual))
    ]
    if not extras:
        return False
    return len(have) >= 20


def _incomplete_paragraph_painted(
        para: str, painted: str, actual: str = '',
) -> bool:
    """True when paint shows this paragraph but is missing content.

    A contiguous fragment of the paragraph without its complete words is
    incomplete. Scattered leftover characters or a different sentence are
    not treated as this paragraph.
    """
    if _complete_paragraph_in_painted(para, painted, actual):
        return False
    want = _layout_norm(para)
    have = _layout_norm(painted)
    if not want or not have:
        return False
    min_len = 12
    limit = min(len(want), 80)
    for length in range(limit, min_len - 1, -1):
        for start in range(0, len(want) - length + 1):
            if want[start:start + length] in have:
                return True
    return False


def _is_visual_presentation_of_actual(text: str, actual: str) -> bool:
    """True when paint is a proven visual form of ActualText.

    Character-reversed leftover and mixed visual-order conversion of the
    same complete statement may match. Digit deletion, overlap
    percentages, set crumbs, and reversed acronym membership do not.
    This is leftover classification, not permission to drop a line.
    """
    act = _layout_norm(actual)
    raw = _layout_norm(text)
    if not raw or not act:
        return False
    if raw == act or _complete_layout_equivalent(text, act):
        return True
    if raw in act and len(raw) >= 12:
        return True
    raw_rev = _layout_norm(str(text or '')[::-1])
    if raw_rev == act:
        return True
    if act and raw_rev in act and len(raw) >= 12:
        return True
    candidates = [
        _layout_norm(_undo_visible_lines(text)),
        _undo_extracted_arabic_visual(text),
        _undo_extracted_arabic_visual(_undo_visible_lines(text)),
    ]
    act_words = _content_words(act)
    for candidate in candidates:
        if not candidate:
            continue
        if candidate == act or _content_words(candidate) == act_words:
            return True
        if candidate != raw and candidate in act and len(candidate) >= 12:
            return True
    return False


def _visible_reconciled_to_actual(visible: str, actual: str) -> Tuple[str, str]:
    """Reconcile painted visible to ActualText only from file evidence.

    ActualText is never reversed. Pure English is never reversed.
    Painted text and logical ActualText stay separately attributable.
    Visible is converted only when the complete associated paint matches
    ActualText after layout normalization or proven visual-to-logical
    conversion. Incomplete paint is not replaced by ActualText.
    """
    vis = _layout_norm(visible)
    act = _layout_norm(actual)
    if not vis:
        return vis, 'visible_empty'
    if not act and _is_mixed_arabic_latin(visible):
        undone = _layout_norm(_undo_visible_lines(visible))
        return undone, 'visible_visual_no_actual'
    if act and vis == act:
        raw_visible = str(visible or '')
        if '\n' in raw_visible:
            return vis, 'visible_wraps_actual'
        return vis, 'same'
    if act and _complete_layout_equivalent(visible, act):
        return act, 'visible_visual_matches_actual'
    return vis, 'visible_as_extracted'


def _ordered_subsequence(want: Sequence[str], have: Sequence[str]) -> bool:
    """True when expected cells appear on the same row in order.

    Extra presentation columns may sit between fields. A swapped
    owner/deliverable or missing field still fails.
    """
    if not want:
        return False
    idx = 0
    for cell in have:
        if idx < len(want) and cell == want[idx]:
            idx += 1
    return idx == len(want)


def _row_present(expected: Sequence[str], tables: Sequence[Sequence[Sequence[str]]]) -> bool:
    want = [_norm(cell) for cell in expected]
    if not any(want):
        return False
    for table in tables:
        for row in table:
            have = [_norm(cell) for cell in row]
            if _ordered_subsequence(want, have):
                return True
    return False


def _text_has_row(expected: Sequence[str], text: str) -> bool:
    blob = _norm(text).lower()
    return all(_norm(cell).lower() in blob for cell in expected if _norm(cell))


def expected_rows(model: CanonicalDocument) -> Dict[str, List[List[str]]]:
    tables = projection_tables(model)
    return {
        'roadmap': [list(row.cells()) for row in model.roadmap],
        'confidence': [list(row.cells()) for row in model.confidence],
        'traceability': [list(row.cells()) for row in model.traceability],
        'kpis': [list(row.cells()) for row in model.kpis],
        'gaps': [list(row.cells()) for row in model.gaps],
        'governance': [list(row.cells()) for row in model.governance],
        'risks': [
            [row.risk, row.impact, row.mitigation, row.owner]
            for row in model.risks
        ],
        '_headers': {
            'roadmap': tables['roadmap']['header'],
            'confidence': tables['confidence']['header'],
        },
    }


def compare_traceability_association(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    """Number-keyed row/field check for saved traceability relationships."""
    inv = inventory_docx_bytes(raw)
    blockers: List[str] = []
    expected_rows_list = [list(row.cells()) for row in model.traceability]
    tables = inv.get('tables') or []

    def _score(table: Sequence[Sequence[str]]) -> int:
        score = 0
        for exp in expected_rows_list:
            want = [_norm(cell) for cell in exp]
            for raw_row in table:
                have = [_norm(cell) for cell in raw_row]
                if want and have and have[0] == want[0]:
                    score += sum(1 for cell in want if cell and cell in have)
                    break
        return score

    ranked = sorted(tables, key=_score, reverse=True)
    chosen = ranked[0] if ranked and _score(ranked[0]) else []
    flat_rows = [[_norm(cell) for cell in row] for row in chosen]
    for idx, expected in enumerate(expected_rows_list):
        want = [_norm(cell) for cell in expected]
        number = want[0] if want else ''
        match = None
        for have in flat_rows:
            if have and have[0] == number:
                match = have
                break
        if match is None:
            blockers.append(
                f'docx_missing_traceability_row:{idx}:{number}')
            continue
        if not _ordered_subsequence(want, match):
            blockers.append(
                f'docx_row_field_mismatch:traceability:{idx}:{number}')
    return blockers


def compare_model_to_docx(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    inv = inventory_docx_bytes(raw)
    blockers: List[str] = []
    for family, rows in expected_rows(model).items():
        if family.startswith('_'):
            continue
        for idx, row in enumerate(rows):
            identity = row[:3] if len(row) >= 3 else row
            if not _row_present(identity, inv['tables']) and not _text_has_row(identity, inv['text']):
                blockers.append(f'docx_missing_{family}_row:{idx}:{_norm(row[0] if row else "")}')
                continue
            # Field association: every non-empty cell must appear on the
            # same row, not merely somewhere in the document.
            if not _row_present(row, inv['tables']):
                blockers.append(
                    f'docx_row_field_mismatch:{family}:{idx}:{_norm(row[0] if row else "")}')
    blockers.extend(compare_environment_narrative_to_docx(model, raw))
    blockers.extend(compare_cover_sector_to_docx(model, raw))
    return blockers


def _section_after_heading(
        paragraphs: Sequence[str],
        heading_re: re.Pattern,
        next_re: re.Pattern,
) -> str:
    taking = False
    collected: List[str] = []
    for para in paragraphs:
        text = str(para or '').strip()
        if not text:
            continue
        # Numbered TOC lines share heading words but are not the body.
        if _TOC_LINE_RE.match(text) and len(text) < 80:
            continue
        if heading_re.search(text) and len(text) < 120:
            taking = True
            continue
        if taking and next_re.search(text) and len(text) < 120:
            break
        if taking:
            collected.append(text)
    return '\n'.join(collected)


def docx_environment_section_text(raw: bytes) -> str:
    inv = inventory_docx_bytes(raw)
    return _section_after_heading(inv.get('paragraphs') or [], _ENV_HEAD_RE, _ENV_NEXT_RE)


def docx_cover_sector_value(raw: bytes) -> str:
    inv = inventory_docx_bytes(raw)
    for table in inv.get('tables') or []:
        for row in table:
            if len(row) >= 2 and _norm(row[0]) in _COVER_LABELS:
                return _norm(row[1])
    return ''


def compare_environment_narrative_to_docx(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    narrative = str(model.environment_narrative or '').strip()
    if not narrative:
        return []
    section = docx_environment_section_text(raw)
    blob = _norm(section)
    if not blob:
        return ['docx_environment_section_missing']
    blockers: List[str] = []
    for idx, para in enumerate(environment_narrative_paragraphs(narrative)):
        if _norm(para) not in blob:
            blockers.append(f'docx_environment_narrative_missing:{idx}')
    if (
            _GENERIC_ENV_AR in section
            and _norm(narrative) not in blob
            and 'سياق تشغيلي' in narrative
            and 'سياق تشغيلي' not in section
    ):
        blockers.append('docx_environment_generic_leftover')
    return blockers


def compare_cover_sector_to_docx(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    expected = cover_sector_from_hashed_narrative(
        model.environment_narrative, model.lang)
    cover = docx_cover_sector_value(raw)
    if not expected:
        if _conflicting_cover_value(cover):
            return ['docx_cover_sector_injected']
        return []
    if not cover:
        return ['docx_cover_sector_missing']
    aliases = sector_reference_aliases(expected)
    if cover not in aliases and not any(alias in cover for alias in aliases):
        return [f'docx_cover_sector_mismatch:{cover}']
    return []


def _compact_sector_token(value: str) -> str:
    return re.sub(r'[\s/\-—–]+', '', _norm(value))


def _conflicting_cover_value(value: str) -> bool:
    raw = _norm(value)
    if not raw or raw in ('—', '-', '\u2014'):
        return False
    known = {english for english, _arabic in UI_SECTOR_PAIRS}
    known.update(arabic for _english, arabic in UI_SECTOR_PAIRS)
    if raw in known:
        return True
    compact = _compact_sector_token(raw)
    return bool(compact) and any(
        _compact_sector_token(item) == compact for item in known)


def _normalize_cover_visual(text: str) -> str:
    """Recover logical Arabic from shaped cover glyphs without inventing words."""
    import unicodedata
    return unicodedata.normalize('NFKC', text or '').replace('\x00', '/')


def _page_lines(page: Dict[str, Any]) -> List[str]:
    blob = '\n'.join(
        part for part in (page.get('actual') or '', page.get('visible') or '')
        if part)
    return [line.strip() for line in blob.splitlines() if line.strip()]


def _is_unreadable_paint_fragment(text: str) -> bool:
    """True when a line is failed hex/CID decode, not readable paint."""
    raw = str(text or '')
    if not raw.strip():
        return True
    letters = sum(
        1 for ch in raw
        if ch.isalpha() or '\u0600' <= ch <= '\u06FF'
    )
    controls = sum(
        1 for ch in raw
        if ord(ch) < 32 or 0x01 <= ord(ch) <= 0x08 or 0x80 <= ord(ch) <= 0x9F
    )
    return controls > 0 and controls >= max(1, letters // 2)


def _pdf_line(value: Any) -> str:
    import unicodedata
    text = unicodedata.normalize('NFKC', str(value or '')).replace('\x00', '').strip()
    if text.startswith('\ufeff'):
        text = text[1:].strip()
    if text.startswith('þÿ'):
        text = text[2:].strip()
    return text


def _looks_like_toc_page(lines: Sequence[str]) -> bool:
    if not lines:
        return False
    toc_hits = 0
    for line in lines:
        text = _pdf_line(line)
        if text and _TOC_NUM_RE.match(text) and len(text) < 80:
            toc_hits += 1
    return toc_hits >= 3 and toc_hits >= max(1, len(lines) // 2)


def _is_running_chrome(line: str) -> bool:
    text = _pdf_line(line)
    return bool(text) and bool(_CHROME_RE.search(text))


def _header_band(y0: Any) -> bool:
    y = float(y0 or 0)
    return y < 40 or y > 800


def _complete_actual_spans(
        fragments: Sequence[str],
        stream_actuals: Sequence[str],
) -> List[str]:
    """Prefer whole ActualText spans over wrap-split line fragments.

    A wrap-split or duplicated line box must not block a stream span
    that already contains one substantial associated fragment.
    """
    parts = [_pdf_line(item) for item in fragments if _pdf_line(item)]
    if not parts:
        return []
    completed: List[str] = []
    seen = set()
    for span in stream_actuals:
        text = str(span or '').strip()
        if not text:
            continue
        cspan = re.sub(r'\s+', '', _pdf_line(text))
        if not cspan:
            continue
        for part in parts:
            cpart = re.sub(r'\s+', '', part)
            if len(cpart) < 12:
                continue
            if cpart in cspan or cspan in cpart:
                key = cspan
                if key not in seen:
                    seen.add(key)
                    completed.append(text)
                break
    return completed or parts


def _env_boundary_line(line: str) -> str:
    text = _pdf_line(line)
    if not text:
        return ''
    if _APPENDIX_RE.search(text) and len(text) < 80:
        return 'stop'
    if _ENV_HEAD_RE.search(text) and len(text) < 120:
        # Numbered TOC entries share heading words. They are not the body.
        if _TOC_NUM_RE.match(text) and len(text) < 80:
            return ''
        return 'heading'
    if _ENV_NEXT_RE.search(text) and len(text) < 120:
        return 'next'
    return ''


def _block_boundary(block: Dict[str, Any]) -> str:
    return (
        _env_boundary_line(block.get('visible') or '')
        or _env_boundary_line(block.get('actual') or '')
    )


def _stream_actual_after_heading(
        events: Sequence[Dict[str, str]],
        *,
        already_taking: bool,
) -> List[str]:
    """ActualText after the heading in stream order. Never reverse it."""
    taking = bool(already_taking)
    collected: List[str] = []
    for event in events:
        text = _pdf_line(event.get('text') or '')
        if not text:
            continue
        boundary = _env_boundary_line(text)
        if boundary == 'heading':
            taking = True
            continue
        if taking and boundary in ('next', 'stop'):
            break
        if taking and event.get('kind') == 'actual':
            collected.append(str(event.get('text') or '').strip())
    return [item for item in collected if item]


def _hidden_overlay_reason(text: str, hidden_norms: Sequence[str]) -> str:
    """File-proven non-displayed logical, not a text-guessed overlay."""
    tn = _layout_norm(text)
    if not tn:
        return ''
    for hidden in hidden_norms:
        if not hidden:
            continue
        if tn == hidden:
            return 'non_displayed_render_mode_3'
        if hidden in tn and len(hidden) >= 12:
            return 'non_displayed_render_mode_3_span'
    return ''


def _strip_proven_hidden_span(text: str, hidden_norms: Sequence[str]) -> str:
    """Remove a file-proven hidden logical span; keep surviving paint."""
    raw = str(text or '')
    norm = _layout_norm(raw)
    for hidden in hidden_norms:
        if not hidden:
            continue
        if hidden in raw:
            remainder = _layout_norm(raw.replace(hidden, ' ', 1))
            if remainder and remainder != hidden:
                return remainder
        if hidden in norm:
            remainder = _layout_norm(norm.replace(hidden, ' ', 1))
            if remainder and remainder != hidden:
                return remainder
    return ''


def _drop_overlay_visible(
        lines: Sequence[Any],
        actual: str,
        *,
        events: Sequence[Dict[str, Any]] = (),
        excluded: Optional[List[Dict[str, str]]] = None,
) -> List[str]:
    """Drop only file-proven non-displayed logicals; keep painted sentences.

    A line is not a harmless overlay merely because it matches, reverses,
    or differs only in digits from ActualText. Geometry or a shared
    y-position is not permission to discard conflicting paint.
    """
    del actual  # ActualText text-guess is not overlay provenance.
    hidden_norms = [
        _layout_norm(item) for item in _hidden_logicals_from_events(events)
        if _layout_norm(item)
    ]
    displayed_norms = {
        _layout_norm(str(event.get('text') or ''))
        for event in events
        if event.get('kind') == 'visible'
        and not event.get('non_displayed')
        and _layout_norm(str(event.get('text') or ''))
    }
    kept: List[Tuple[float, str]] = []
    kept_norms: set = set()
    for item in lines:
        if isinstance(item, (tuple, list)) and len(item) >= 2:
            y0, raw = item[0], item[1]
        else:
            y0, raw = 0, item
        text = _pdf_line(raw)
        if not text or _is_running_chrome(text) or _is_unreadable_paint_fragment(text):
            continue
        reason = _hidden_overlay_reason(text, hidden_norms)
        tn = _layout_norm(text)
        if reason:
            remainder = _strip_proven_hidden_span(text, hidden_norms)
            if excluded is not None:
                record = {'text': text, 'reason': reason}
                if remainder:
                    record['remainder'] = remainder
                excluded.append(record)
            if remainder:
                kept.append((float(y0 or 0), remainder))
                kept_norms.add(_layout_norm(remainder))
            elif tn in displayed_norms and tn not in kept_norms:
                kept.append((float(y0 or 0), text))
                kept_norms.add(tn)
            continue
        kept.append((float(y0 or 0), text))
        kept_norms.add(tn)
    kept.sort(key=lambda item: item[0])
    return [text for _y, text in kept]


def pdf_environment_section_text(raw: bytes) -> Tuple[str, Dict[str, Any]]:
    """Environment-section text only: not cover, TOC, summary, or appendix.

    Association is the content range after the environment heading and
    before the next section. Whole-page ActualText is not borrowed.
    Stream ActualText after the heading stays logical and is not reversed.
    """
    pages, meta = extract_pdf_pages(raw)
    detail = dict(meta)
    if not pages:
        detail['associated'] = False
        return '', detail
    if not meta.get('associated'):
        return '', detail
    collected: List[str] = []
    visible_parts: List[str] = []
    actual_parts: List[str] = []
    overlay_excluded: List[Dict[str, str]] = []
    taking = False
    heading_page = None
    repeated_chrome: set = set()
    for page in pages:
        index = int(page.get('index') or 0)
        if index == 0:
            continue
        visible_lines = [
            _pdf_line(line)
            for line in str(page.get('visible') or '').splitlines()
            if _pdf_line(line)
        ]
        if _looks_like_toc_page(visible_lines) and not taking:
            continue
        blocks = list(page.get('blocks') or [])
        if not blocks:
            blocks = [
                {'visible': line, 'actual': '', 'y0': idx}
                for idx, line in enumerate(visible_lines)
            ]
        page_started = taking
        page_vis: List[Any] = []
        page_act_frags: List[str] = []
        page_took = False
        page_chrome: set = set()
        for block in blocks:
            visible = _pdf_line(block.get('visible') or '')
            actual = _pdf_line(block.get('actual') or '')
            y0 = block.get('y0') or 0
            if visible and _header_band(y0):
                page_chrome.add(visible)
            if (
                    visible
                    and _header_band(y0)
                    and (visible in repeated_chrome or _is_running_chrome(visible))
            ):
                continue
            boundary = _block_boundary(block)
            if boundary == 'stop':
                if taking:
                    taking = False
                break
            if boundary == 'heading':
                taking = True
                heading_page = index
                continue
            if taking and boundary == 'next':
                taking = False
                break
            if taking:
                page_took = True
                if visible and not _is_running_chrome(visible):
                    page_vis.append((y0, visible))
                if actual and not _is_running_chrome(actual):
                    page_act_frags.append(actual)
        events_early = list(page.get('events') or [])
        heading_seen = any(
            _env_boundary_line(event.get('text') or '') == 'heading'
            for event in events_early)
        taking_event = heading_seen or page_started
        for event in events_early:
            text = _pdf_line(event.get('text') or '')
            if not text:
                continue
            boundary = _env_boundary_line(text)
            if boundary == 'heading':
                taking_event = True
                continue
            if taking_event and boundary in ('next', 'stop'):
                break
            if (
                    taking_event
                    and event.get('kind') == 'visible'
                    and event.get('logical_paint')
                    and not event.get('non_displayed')
                    and not _is_running_chrome(text)
            ):
                already = {
                    _layout_norm(item[1] if isinstance(item, (tuple, list)) else item)
                    for item in page_vis
                }
                if _layout_norm(text) not in already:
                    page_took = True
                    page_vis.append((0, text))
        repeated_chrome.update(page_chrome)
        if not page_took:
            continue
        events = list(page.get('events') or [])
        stream_after = _stream_actual_after_heading(
            events, already_taking=page_started)
        stream_all = [
            str(event.get('text') or '').strip()
            for event in events
            if event.get('kind') == 'actual' and str(event.get('text') or '').strip()
        ]
        vis_texts = [
            item[1] if isinstance(item, (tuple, list)) else str(item)
            for item in page_vis
        ]
        heading_in_events = any(
            _env_boundary_line(event.get('text') or '') == 'heading'
            for event in events)
        if stream_after and (heading_in_events or not page_started):
            page_act = stream_after
        elif heading_page == index and not heading_in_events:
            # CID/shaped heading is visible in line boxes, not in the
            # content-stream Tj inventory. Page ActualText on that page
            # is the body; pre-heading AT is excluded when the heading
            # itself appears in the stream (stream_after path).
            page_act = stream_all
        else:
            page_act = _complete_actual_spans(
                vis_texts + page_act_frags, stream_all)
        page_vis = _drop_overlay_visible(
            page_vis,
            '\n'.join(page_act),
            events=events,
            excluded=overlay_excluded,
        )
        collected.extend(page_vis)
        collected.extend(page_act)
        if page_vis or page_act:
            visible_parts.append('\n'.join(page_vis))
            actual_parts.append('\n'.join(page_act))
    section = '\n'.join(collected)
    detail['environment_heading_page'] = heading_page
    detail['environment_associated'] = heading_page is not None
    detail['environment_visible'] = '\n'.join(visible_parts)
    detail['environment_actual'] = '\n'.join(actual_parts)
    detail['overlay_excluded'] = overlay_excluded
    detail['hidden_logicals'] = [
        item for page in pages
        for item in _hidden_logicals_from_events(page.get('events') or [])
        if item
    ]
    return section, detail


def _cover_value_matches(value: str, expected: str) -> bool:
    aliases = sector_reference_aliases(expected)
    raw = _norm(value)
    if raw in aliases:
        return True
    compact = _compact_sector_token(value)
    if not compact:
        return False
    return any(
        alias in raw or _compact_sector_token(alias) == compact
        for alias in aliases
    )


def _cover_label_aliases() -> Tuple[str, ...]:
    aliases = []
    for label in _COVER_LABELS:
        aliases.append(label)
        if label.isascii():
            aliases.append(label.upper())
            aliases.append(label.lower())
            aliases.append(label.title())
    return tuple(dict.fromkeys(aliases))


def _labeled_cover_sector(text: str) -> Tuple[str, bool]:
    """Return (value, field_found) for the labeled Sector/القطاع field.

    A missing label is not a successful empty field. A found label with
    no readable value is an unreadable field, not a neutral dash.
    """
    labels = _cover_label_aliases()
    lines = [line.strip() for line in str(text or '').splitlines() if line.strip()]
    for idx, line in enumerate(lines):
        n = _norm(line)
        for label in labels:
            if n == label:
                if idx + 1 < len(lines):
                    nxt = _norm(lines[idx + 1])
                    if nxt in labels:
                        return '', True
                    return nxt, True
                return '', True
            if n.startswith(label + ' ') or n.startswith(label + ':') or n.startswith(label + '\t'):
                rest = n[len(label):].lstrip(' :|\t')
                return rest, True
            if n.startswith(label) and len(n) > len(label):
                rest = n[len(label):].lstrip(' :|\t')
                if rest:
                    return rest, True
    blob = _norm(text)
    for label in labels:
        pos = blob.find(label)
        if pos < 0:
            continue
        after = blob[pos + len(label):].lstrip(' :|\t')
        token = after.split(' ')[0] if after else ''
        return token, True
    return '', False


def pdf_cover_sector_value(raw: bytes) -> Tuple[str, Dict[str, Any]]:
    """Labeled cover-sector value from page 0 only."""
    pages, meta = extract_pdf_pages(raw)
    detail = dict(meta)
    detail['field_found'] = False
    detail['field_readable'] = False
    if not pages or not meta.get('associated'):
        return '', detail
    cover = pages[0]
    actual = cover.get('actual') or ''
    visual = _normalize_cover_visual(cover.get('visible') or '')
    value, found = _labeled_cover_sector(actual)
    source = 'actual'
    if not found:
        value, found = _labeled_cover_sector(visual)
        source = 'visible'
    detail['field_found'] = found
    detail['field_source'] = source if found else ''
    detail['field_readable'] = bool(found and value)
    detail['cover_page_index'] = 0
    return value, detail


def pdf_cover_and_body_text(raw: bytes) -> Tuple[str, str, Dict[str, Any]]:
    """Keep cover visual evidence distinct from the environment section."""
    pages, meta = extract_pdf_pages(raw)
    cover = ''
    if pages:
        cover = pages[0].get('visible') or ''
    section, env_meta = pdf_environment_section_text(raw)
    merged = dict(meta)
    merged.update(env_meta)
    merged['environment_page_present'] = bool(env_meta.get('environment_associated'))
    return cover, section, merged


def compare_environment_narrative_to_pdf(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    narrative = str(model.environment_narrative or '').strip()
    if not narrative:
        return []
    section, meta = pdf_environment_section_text(raw)
    if raw.startswith(b'%PDF') and not meta.get('reliable') and not section.strip():
        return ['pdf_extraction_unreliable']
    if not meta.get('associated'):
        return ['pdf_environment_section_unassociated']
    if not meta.get('environment_associated'):
        return ['pdf_environment_section_unassociated']
    blob = _norm(section)
    if not blob:
        return ['pdf_environment_text_empty']
    visible = meta.get('environment_visible') or ''
    actual = meta.get('environment_actual') or ''
    blockers: List[str] = []
    for idx, para in enumerate(environment_narrative_paragraphs(narrative)):
        blockers.extend(_paragraph_pdf_blockers(
            idx, para, section, visible=visible, actual=actual))
    return blockers


_LATIN_TOKEN_RE = re.compile(r'[A-Za-z][A-Za-z0-9_-]*')


def _latin_tokens(text: str) -> List[str]:
    return _LATIN_TOKEN_RE.findall(str(text or ''))


def _semantic_latin_tokens(text: str) -> List[str]:
    """Identity-bearing Latin tokens, not ordinary wrapped English words."""
    out: List[str] = []
    for token in _latin_tokens(text):
        if token.isupper() and len(token) >= 2:
            out.append(token)
        elif any(ch.isdigit() for ch in token) and any(ch.isalpha() for ch in token):
            out.append(token)
        elif len(token) >= 4 and sum(1 for ch in token if ch.isupper()) >= 2:
            out.append(token)
    return out


def _latin_forward_window(
        vis_lat: Sequence[str],
        act_lat: Sequence[str],
) -> bool:
    """Visible tokens follow ActualText order, including wrap extras."""
    act = list(act_lat)
    vis = list(vis_lat)
    if not vis:
        return True
    if not act:
        return False
    for start in range(len(act)):
        idx = start
        consumed = list(act[start:start])
        extra = False
        ok = True
        for pos, token in enumerate(vis):
            if extra:
                if token not in consumed:
                    ok = False
                    break
                continue
            if idx < len(act) and token == act[idx]:
                consumed.append(act[idx])
                idx += 1
                if idx >= len(act):
                    extra = True
                continue
            if token in consumed:
                continue
            if (
                    pos == len(vis) - 1
                    and idx < len(act)
                    and len(token) >= 3
                    and len(token) < len(act[idx])
                    and act[idx].startswith(token)
            ):
                consumed.append(act[idx])
                extra = True
                continue
            ok = False
            break
        if ok:
            return True
    return False


def _latin_vis_conflicts_act(
        vis_lat: Sequence[str],
        act_lat: Sequence[str],
) -> bool:
    """True when visible latin identity/order contradicts ActualText.

    A final wrap-split prefix (PDP from PDPL) is not a contradiction.
    A swapped or substituted token is.
    """
    if not vis_lat:
        return False
    vis = list(vis_lat)
    act = list(act_lat)
    if vis == act:
        return False
    if _latin_forward_window(vis, act):
        return False
    return True


def _ordered_latin_match(want: Sequence[str], have: Sequence[str]) -> bool:
    if not want:
        return True
    idx = 0
    for token in have:
        if idx < len(want) and token == want[idx]:
            idx += 1
    return idx == len(want)


def _stream_usable(text: str, para: str) -> bool:
    """False when a stream lost the script the paragraph actually uses."""
    raw = str(text or '')
    if not _layout_norm(raw):
        return False
    want_ar = any('\u0600' <= ch <= '\u06FF' for ch in para)
    have_ar = any('\u0600' <= ch <= '\u06FF' for ch in raw)
    want_lat = _semantic_latin_tokens(para)
    have_lat = _semantic_latin_tokens(raw)
    if want_ar and not have_ar:
        return False
    if want_lat and not have_lat:
        return False
    return True


def _latin_relation(text: str) -> List[str]:
    return _semantic_latin_tokens(_layout_norm(text))


def _paragraph_pdf_blockers(
        idx: int,
        para: str,
        section: str,
        *,
        visible: str = '',
        actual: str = '',
) -> List[str]:
    """Compare one saved paragraph inside the associated environment section.

    Visible and ActualText stay distinct. A match in one stream cannot
    override a material disagreement in the other. Logical ActualText is
    never reversed to create a match.
    """
    want = _layout_norm(para)
    if not want:
        return []
    vis_raw = str(visible or '')
    act_raw = str(actual or '')
    streams_omitted = not vis_raw.strip() and not act_raw.strip()
    if streams_omitted:
        vis_raw = str(section or '')
        act_raw = str(section or '')
    elif not vis_raw.strip() and act_raw.strip():
        return [f'pdf_environment_painted_unestablished:{idx}']
    vis_use = _stream_usable(vis_raw, para)
    act_use = _stream_usable(act_raw, para)
    vis_cmp, how = _visible_reconciled_to_actual(vis_raw, act_raw)
    act_cmp = _layout_norm(act_raw)
    vis_complete = vis_use and _complete_paragraph_in_painted(
        para, vis_raw, act_raw)
    act_complete = act_use and _complete_paragraph_in_painted(
        para, act_raw, act_raw)
    vis_incomplete = vis_use and _incomplete_paragraph_painted(
        para, vis_raw, act_raw)
    leftover_ok = _leftover_visual_of_actual(vis_raw, para, act_raw)
    blockers: List[str] = []
    if vis_use and act_use and vis_incomplete and act_complete:
        blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
        return blockers
    if vis_use and act_use and how not in (
            'same',
            'visible_visual_matches_actual',
            'visible_wraps_actual',
            'visible_visual_no_actual',
    ):
        vis_for_latin = vis_cmp
        if not _arabic_already_logical(vis_raw, para):
            vis_for_latin = _layout_norm(_undo_visible_lines(vis_raw)) or vis_cmp
        vis_lat = _latin_relation(vis_for_latin)
        act_lat = _latin_relation(act_cmp)
        latin_conflict = _latin_vis_conflicts_act(vis_lat, act_lat)
        if latin_conflict and vis_complete and leftover_ok and _ordered_latin_match(
                act_lat, vis_lat):
            latin_conflict = False
        if latin_conflict:
            blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
            return blockers
        if vis_cmp != act_cmp and vis_incomplete:
            blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
            return blockers
        if vis_cmp != act_cmp and _unrelated_painted_sentence(
                vis_raw, para, act_raw):
            blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
            return blockers
        if vis_cmp != act_cmp and not vis_complete:
            blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
            return blockers
        if vis_cmp != act_cmp and vis_complete and not leftover_ok:
            blockers.append(f'pdf_environment_actual_visible_disagree:{idx}')
            return blockers
    targets: List[str] = []
    if vis_use and act_use:
        if how in (
                'same',
                'visible_visual_matches_actual',
                'visible_wraps_actual',
                'visible_visual_no_actual',
        ):
            targets = [act_cmp or vis_cmp]
        elif vis_complete and leftover_ok:
            undone = _layout_norm(_undo_visible_lines(vis_raw))
            targets = [item for item in (undone, vis_cmp) if item]
        else:
            targets = [vis_cmp]
    elif not vis_raw.strip() and act_use:
        return [f'pdf_environment_painted_unestablished:{idx}']
    elif act_use:
        targets = [act_cmp]
    elif vis_use:
        targets = [vis_cmp]
    elif _layout_norm(section):
        targets = [_layout_norm(section)]
    else:
        return [f'pdf_environment_narrative_unreadable:{idx}']
    want_lat = _semantic_latin_tokens(para)
    matched = False
    for target in targets:
        if want_lat:
            have_lat = _latin_relation(target)
            missing = [token for token in want_lat if token not in have_lat]
            if missing:
                continue
            if not _ordered_latin_match(want_lat, have_lat):
                continue
        if want in target:
            matched = True
            break
        if vis_complete and leftover_ok and _ordered_statement_present(
                para, target, act_raw):
            matched = True
            break
    if not matched and vis_complete and leftover_ok and _ordered_statement_present(
            para, vis_raw, act_raw):
        matched = True
    if not matched:
        if want_lat:
            have_any = []
            for target in targets:
                have_any.extend(_latin_relation(target))
            missing = [token for token in want_lat if token not in have_any]
            if missing:
                for token in missing:
                    blockers.append(
                        f'pdf_environment_latin_missing:{idx}:{token}')
            elif not any(
                    _ordered_latin_match(want_lat, _latin_relation(target))
                    for target in targets):
                blockers.append(f'pdf_environment_latin_order:{idx}')
        if want not in ' '.join(targets):
            blockers.append(f'pdf_environment_narrative_missing:{idx}')
    return list(dict.fromkeys(blockers))


def _paragraph_in_pdf_section(para: str, section: str) -> bool:
    return not _paragraph_pdf_blockers(0, para, section)


def compare_cover_sector_to_pdf(
        model: CanonicalDocument,
        raw: bytes,
) -> List[str]:
    expected = cover_sector_from_hashed_narrative(
        model.environment_narrative, model.lang)
    value, meta = pdf_cover_sector_value(raw)
    if raw.startswith(b'%PDF') and not meta.get('associated'):
        return ['pdf_cover_sector_unassociated']
    if not meta.get('field_found'):
        if expected:
            return ['pdf_cover_sector_unreadable']
        return ['pdf_cover_sector_unreadable']
    if not expected:
        if _conflicting_cover_value(value):
            return [f'pdf_cover_sector_injected:{value}']
        return []
    if not value:
        return ['pdf_cover_sector_unreadable']
    if not _cover_value_matches(value, expected):
        return [f'pdf_cover_sector_mismatch:{value}']
    return []


def hashed_narrative_requires_section_parity(model: CanonicalDocument) -> bool:
    """Required nonempty REL37 environment_narrative is compared.

    Applicability is the authorized REL37 model contract, not a keyword
    test for ``سياق تشغيلي`` / ``operating context``. Empty narrative
    follows the existing schema omission contract. Cyber/ERM/Global and
    other non-REL37 domains are not selected here. The route gate still
    requires ``is_rel37_authoritative`` before this helper runs.
    """
    domain = str(getattr(model, 'domain', '') or '').strip().lower()
    if domain not in _REL37_NARRATIVE_DOMAINS:
        return False
    text = str(getattr(model, 'environment_narrative', '') or '').strip()
    return bool(text)


def rel37_returned_bytes_blockers(
        model: CanonicalDocument,
        *,
        docx_bytes: bytes = b'',
        pdf_bytes: bytes = b'',
        route: str = '',
) -> List[str]:
    """Route-local returned-byte blockers. Preview cannot substitute PDF."""
    route_n = str(route or '').strip().lower()
    blockers: List[str] = []
    check_docx = route_n in ('', 'docx', 'docx-async') and docx_bytes
    check_pdf = route_n in ('', 'pdf', 'pdf-async') and pdf_bytes
    if route_n in ('docx', 'docx-async') and not docx_bytes:
        blockers.append('docx_bytes_missing')
    if route_n in ('pdf', 'pdf-async') and not pdf_bytes:
        blockers.append('pdf_bytes_missing')
    require_narrative = hashed_narrative_requires_section_parity(model)
    if check_docx:
        if require_narrative:
            blockers.extend(compare_environment_narrative_to_docx(model, docx_bytes))
        blockers.extend(compare_cover_sector_to_docx(model, docx_bytes))
    if check_pdf:
        if require_narrative:
            blockers.extend(compare_environment_narrative_to_pdf(model, pdf_bytes))
        blockers.extend(compare_cover_sector_to_pdf(model, pdf_bytes))
    return list(dict.fromkeys(blockers))


def gate_rel37_returned_bytes(
        model: CanonicalDocument,
        *,
        docx_bytes: bytes = b'',
        pdf_bytes: bytes = b'',
        route: str = '',
) -> Tuple[bool, List[str]]:
    blockers = rel37_returned_bytes_blockers(
        model, docx_bytes=docx_bytes, pdf_bytes=pdf_bytes, route=route)
    return not blockers, blockers


def compare_model_to_text(
        model: CanonicalDocument,
        text: str,
        *,
        route: str,
) -> List[str]:
    blockers: List[str] = []
    blob = _norm(text)
    if not blob:
        return [f'{route}_text_empty']
    for family, rows in expected_rows(model).items():
        if family.startswith('_'):
            continue
        for idx, row in enumerate(rows):
            if not _text_has_row(row[:3] if len(row) >= 3 else row, blob):
                blockers.append(f'{route}_missing_{family}_row:{idx}:{_norm(row[0] if row else "")}')
    return blockers


def language_blockers(model: CanonicalDocument, text: str) -> List[str]:
    if model.lang != 'en':
        return []
    count = arabic_char_count(text, org_name=model.org_name)
    if count:
        return [f'en_unexpected_arabic_chars:{count}']
    return []


def evaluate_export_parity(
        *,
        model: CanonicalDocument,
        persist_ok: bool,
        download_ok: bool,
        parsed_ok: bool,
        preview_text: str = '',
        docx_bytes: bytes = b'',
        pdf_bytes: bytes = b'',
        source_hash: str = '',
        security: str = 'not_run',
) -> ParityResult:
    result = ParityResult(
        persist_success='passed' if persist_ok else 'failed',
        download_parse_success='passed' if (download_ok and parsed_ok) else 'failed',
        security_acceptance=security,
        model_hash=model.model_hash,
        source_hash=source_hash,
    )
    if source_hash and model.model_hash and source_hash != model.model_hash:
        # Matching/mismatching source hash never overrides actual content.
        result.notes.append('source_hash_recorded_separately')
    if not persist_ok:
        result.blockers.append('persist_failed')
    if not (download_ok and parsed_ok):
        result.blockers.append('download_or_parse_failed')

    preview_blockers = compare_model_to_text(model, preview_text, route='preview') if preview_text else []
    docx_blockers = compare_model_to_docx(model, docx_bytes) if docx_bytes else ['docx_bytes_missing']
    pdf_text, pdf_meta = extract_pdf_text(pdf_bytes) if pdf_bytes else ('', {'reliable': False})
    if pdf_bytes and not pdf_meta.get('reliable'):
        pdf_blockers = ['pdf_extraction_unreliable']
    elif pdf_bytes:
        pdf_blockers = compare_model_to_text(model, pdf_text, route='pdf')
        pdf_blockers.extend(compare_environment_narrative_to_pdf(model, pdf_bytes))
        pdf_blockers.extend(compare_cover_sector_to_pdf(model, pdf_bytes))
    else:
        pdf_blockers = ['pdf_bytes_missing']

    content_blockers = preview_blockers + docx_blockers + pdf_blockers
    result.content_parity = 'failed' if content_blockers else 'passed'
    result.blockers.extend(content_blockers)

    lang_text = '\n'.join([
        preview_text,
        inventory_docx_bytes(docx_bytes)['text'] if docx_bytes else '',
        pdf_text,
    ])
    lang_blockers = language_blockers(model, lang_text)
    result.language_correctness = 'failed' if lang_blockers else 'passed'
    result.blockers.extend(lang_blockers)

    readable = bool(docx_bytes) and bool(pdf_meta.get('reliable'))
    if pdf_bytes and not pdf_meta.get('reliable'):
        result.rendering_readability = 'failed'
        result.blockers.append('pdf_unreadable_or_unextractable')
    elif not readable:
        result.rendering_readability = 'failed'
        result.blockers.append('export_not_readable')
    else:
        result.rendering_readability = 'passed'

    if docx_bytes:
        result.file_sha256 = hashlib.sha256(docx_bytes).hexdigest()
    matching_hash = bool(
        source_hash and model.model_hash and source_hash == model.model_hash)
    docx_content_failed = any(
        item.startswith('docx_row_field_mismatch')
        or item.startswith('docx_missing_')
        for item in docx_blockers)
    result.details = {
        'preview_blockers': preview_blockers,
        'docx_blockers': docx_blockers,
        'pdf_blockers': pdf_blockers,
        'pdf_meta': pdf_meta,
        'source_hash_matches_model': matching_hash,
        'source_hash_does_not_override_content': bool(
            matching_hash and docx_content_failed),
    }
    result.blockers = list(dict.fromkeys(result.blockers))
    return result


def negative_control_arabic_cell_fails(model: CanonicalDocument) -> bool:
    """A readable DOCX with one unexpected generated Arabic cell fails parity."""
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=1, cols=6)
    cells = list(model.roadmap[0].cells())
    cells[4] = 'إطار حوكمة NDMO معتمد'
    for idx, value in enumerate(cells):
        table.cell(0, idx).text = value
    buf = io.BytesIO()
    doc.save(buf)
    blockers = compare_model_to_docx(model, buf.getvalue())
    return any('row_field_mismatch' in item or 'missing_roadmap' in item for item in blockers)


def negative_control_english_wording_fails(model: CanonicalDocument) -> bool:
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=1, cols=6)
    cells = list(model.roadmap[0].cells())
    cells[4] = 'Approved NDMO framework charter'
    for idx, value in enumerate(cells):
        table.cell(0, idx).text = value
    buf = io.BytesIO()
    doc.save(buf)
    blockers = compare_model_to_docx(model, buf.getvalue())
    return any('row_field_mismatch' in item or 'missing_roadmap' in item for item in blockers)
