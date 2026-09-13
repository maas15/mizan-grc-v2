"""REL37 renderer — all targets consume the same CanonicalDocument."""
from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass, field
from typing import Dict, List, Tuple

from release_engine_v3.rel37_canonical_document import (
    CanonicalDocument,
    sha256_text,
)
from release_engine_v3.rel37_schema_registry import (
    REL32_AR_HEADINGS,
    SECTION_TITLES,
    header_line,
    separator_line,
)


def _md_table(kind: str, lang: str, rows: List[Tuple[str, ...]]) -> str:
    lines = [header_line(kind, lang), separator_line(kind, lang)]
    for row in rows:
        lines.append('| ' + ' | '.join(str(c) for c in row) + ' |')
    return '\n'.join(lines)


def _section_heading(key: str, lang: str) -> str:
    if lang == 'ar':
        title = REL32_AR_HEADINGS.get(key) or SECTION_TITLES['ar'].get(key, key)
        return f'## {title}'
    title = SECTION_TITLES['en'].get(key, key)
    return f'## {title}'


def _md_guide(heading: str, lang: str, steps) -> str:
    table = _md_table('guide', lang, [s.cells() for s in steps])
    return f'{heading}\n\n{table}'


def model_to_sections(model: CanonicalDocument) -> Dict[str, str]:
    lang = model.lang
    so_table = _md_table('so', lang, [r.cells() for r in model.strategic_objectives])
    vision = '\n\n'.join([
        _section_heading('vision', lang),
        model.vision,
        so_table,
    ])

    pillar_parts = [_section_heading('pillars', lang)]
    for pillar in model.pillars:
        if lang == 'ar':
            pillar_parts.append(f'### الركيزة {pillar.number}: {pillar.title}')
        else:
            pillar_parts.append(f'### Pillar {pillar.number}: {pillar.title}')
        pillar_parts.append(pillar.description)
        inits = [
            row.cells() for row in model.pillar_initiatives
            if row.pillar_number == pillar.number
        ]
        if inits:
            pillar_parts.append(_md_table('pillar', lang, inits))
    pillars = '\n\n'.join(pillar_parts)

    environment = '\n\n'.join([
        _section_heading('environment', lang),
        model.environment_narrative,
    ])

    from release_engine_v3.rel37_schema_registry import FORMULA_BLOCK, KPI_GUIDES_BLOCK
    gap_parts = [_section_heading('gaps', lang)]
    if lang == 'ar':
        gap_parts.append('يتم إعداد دليل تطبيق لكل فجوة محددة في الجدول التالي.')
    else:
        gap_parts.append('Each identified gap has one implementation guide.')
    gap_parts.append(_md_table('gap', lang, [r.cells() for r in model.gaps]))
    for guide in model.gap_guides:
        gap_parts.append(_md_guide(guide.heading, lang, guide.steps))
    gaps = '\n\n'.join(gap_parts)

    roadmap = '\n\n'.join([
        _section_heading('roadmap', lang),
        _md_table('roadmap', lang, [r.cells() for r in model.roadmap]),
    ])

    kpi_parts = [
        _section_heading('kpis', lang),
        _md_table('kpi_main', lang, [r.cells() for r in model.kpis]),
    ]
    if model.kpi_formula_source:
        kpi_parts.append(FORMULA_BLOCK[lang])
        kpi_parts.append(
            _md_table('kpi_formula', lang, [r.cells() for r in model.kpi_formula_source]))
    kpi_parts.append(KPI_GUIDES_BLOCK[lang])
    for guide in model.kpi_guides:
        kpi_parts.append(_md_guide(guide.heading, lang, guide.steps))
    kpis = '\n\n'.join(kpi_parts)

    if lang == 'ar':
        score_line = f'درجة الثقة: {model.confidence_score}'
        risk_head = '### سجل المخاطر وخطة المعالجة'
        gov_headers = ('الدور', 'المسؤولية', 'التكرار', 'المالك')
        conf_note = model.confidence_justification
    else:
        score_line = f'Confidence Score: {model.confidence_score}'
        risk_head = '### Risk register and treatment plan'
        gov_headers = ('Role', 'Responsibility', 'Cadence', 'Owner')
        conf_note = model.confidence_justification
    confidence = '\n\n'.join([
        _section_heading('confidence', lang),
        score_line,
        conf_note,
        _md_table('csf', lang, [r.cells() for r in model.confidence]),
        risk_head,
        _md_table('risk', lang, [r.cells() for r in model.risks]),
    ])

    gov_table = (
        '| ' + ' | '.join(gov_headers) + ' |\n'
        + '|' + '|'.join(['---'] * 4) + '|\n'
        + '\n'.join(
            '| ' + ' | '.join(r.cells()) + ' |' for r in model.governance)
    )
    governance = '\n\n'.join([
        _section_heading('governance', lang),
        gov_table,
    ])
    traceability = '\n\n'.join([
        _section_heading('traceability', lang),
        _md_table('trace', lang, [r.cells() for r in model.traceability]),
    ])
    return {
        'vision': vision,
        'pillars': pillars,
        'environment': environment,
        'gaps': gaps,
        'roadmap': roadmap,
        'kpis': kpis,
        'confidence': confidence,
        'governance': governance,
        'traceability': traceability,
    }


def model_to_markdown(model: CanonicalDocument) -> str:
    sections = model_to_sections(model)
    order = (
        'vision', 'pillars', 'environment', 'gaps', 'roadmap',
        'kpis', 'confidence', 'governance', 'traceability',
    )
    return '\n\n'.join(sections[k] for k in order)


def _md_to_html(markdown: str) -> str:
    html_parts: List[str] = []
    for raw in (markdown or '').splitlines():
        line = raw.rstrip()
        if not line:
            html_parts.append('')
            continue
        if line.startswith('## '):
            html_parts.append(f'<h2>{_escape(line[3:])}</h2>')
            continue
        if line.startswith('### '):
            html_parts.append(f'<h3>{_escape(line[4:])}</h3>')
            continue
        if line.startswith('#### '):
            html_parts.append(f'<h4>{_escape(line[5:])}</h4>')
            continue
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip('|').split('|')]
            if all(set(c.replace('-', '')) <= {':', ''} and c for c in cells):
                continue
            tag = 'th' if '---' not in line else 'td'
            if any(h in cells for h in (
                    '#', 'Strategic Objective', 'الهدف الاستراتيجي',
                    'KPI Description', 'وصف المؤشر', 'Phase', 'المرحلة',
                    'Step', 'الخطوة')):
                tag = 'th'
            inner = ''.join(f'<{tag}>{_escape(c)}</{tag}>' for c in cells)
            html_parts.append(f'<tr>{inner}</tr>')
            continue
        html_parts.append(f'<p>{_escape(line)}</p>')
    body: List[str] = []
    in_table = False
    for part in html_parts:
        if part.startswith('<tr>'):
            if not in_table:
                body.append('<table>')
                in_table = True
            body.append(part)
            continue
        if in_table:
            body.append('</table>')
            in_table = False
        if part:
            body.append(part)
    if in_table:
        body.append('</table>')
    return '\n'.join(body)


def _escape(text: str) -> str:
    return (
        (text or '')
        .replace('&', '&amp;')
        .replace('<', '&lt;')
        .replace('>', '&gt;')
    )


@dataclass(frozen=True)
class EvidenceProjection:
    model_hash: str
    so_header: str
    kpi_main_header: str
    formula_header: str
    kpi_row_count: int
    kpi_guide_count: int
    gap_row_count: int
    gap_guide_count: int
    roadmap_families: Tuple[str, ...]
    org_name: str
    leakage_hits: Tuple[str, ...] = ()
    source_hash: str = ''

    def to_dict(self) -> Dict[str, object]:
        return {
            'model_hash': self.model_hash,
            'so_header': self.so_header,
            'kpi_main_header': self.kpi_main_header,
            'formula_header': self.formula_header,
            'kpi_row_count': self.kpi_row_count,
            'kpi_guide_count': self.kpi_guide_count,
            'gap_row_count': self.gap_row_count,
            'gap_guide_count': self.gap_guide_count,
            'roadmap_families': list(self.roadmap_families),
            'org_name': self.org_name,
            'leakage_hits': list(self.leakage_hits),
            'source_hash': self.source_hash,
        }


@dataclass
class Rendered:
    target: str
    body: object
    content_hash: str
    evidence: EvidenceProjection
    sections: Dict[str, str] = field(default_factory=dict)
    markdown: str = ''


def evidence_from_model(model: CanonicalDocument) -> EvidenceProjection:
    from release_engine_v3.rel37_canonical_document import detect_leakage
    from release_engine_v3.rel37_schema_registry import leakage_terms
    return EvidenceProjection(
        model_hash=model.model_hash,
        so_header=header_line('so', model.lang),
        kpi_main_header=header_line('kpi_main', model.lang),
        formula_header=(
            header_line('kpi_formula', model.lang)
            if model.kpi_formula_source else ''),
        kpi_row_count=len(model.kpis),
        kpi_guide_count=len(model.kpi_guides),
        gap_row_count=len(model.gaps),
        gap_guide_count=len(model.gap_guides),
        roadmap_families=tuple(r.family for r in model.roadmap),
        org_name=model.org_name,
        leakage_hits=detect_leakage(
            model.generated_text_blob(),
            leakage_terms(model.domain),
            org_name=model.org_name,
        ),
        source_hash=model.model_hash,
    )


def _docx_bytes(markdown: str, model: CanonicalDocument) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    rtl = model.lang == 'ar'
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip('|').split('|')]
            if all(set(c.replace('-', '')) <= {':', ''} for c in cells):
                continue
            table = doc.add_table(rows=1, cols=len(cells))
            for idx, cell in enumerate(cells):
                table.rows[0].cells[idx].text = cell
            continue
        para = doc.add_paragraph(line.lstrip('#').strip())
        if rtl:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(markdown: str, model: CanonicalDocument) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 40
    font_name = 'Helvetica'
    try:
        import os
        for candidate in (
                '/workspace/static/fonts/Amiri-Regular.ttf',
                '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        ):
            if os.path.exists(candidate):
                pdfmetrics.registerFont(TTFont('Rel37Body', candidate))
                font_name = 'Rel37Body'
                break
    except Exception:
        font_name = 'Helvetica'
    c.setFont(font_name, 9)
    for raw in markdown.splitlines():
        line = raw[:140]
        if y < 40:
            c.showPage()
            c.setFont(font_name, 9)
            y = height - 40
        try:
            c.drawString(36, y, line)
        except Exception:
            c.drawString(36, y, line.encode('ascii', 'replace').decode('ascii'))
        y -= 12
    c.save()
    return buf.getvalue()


def render(model: CanonicalDocument, target: str) -> Rendered:
    if not model.model_hash:
        model.compute_hashes()
    sections = model_to_sections(model)
    markdown = model_to_markdown(model)
    evidence = evidence_from_model(model)
    target = str(target or 'md').lower()
    if target in ('md', 'markdown', 'txt', 'print'):
        body = markdown
        digest = sha256_text(markdown)
    elif target in ('html', 'preview'):
        body = _md_to_html(markdown)
        digest = sha256_text(str(body))
    elif target == 'docx':
        body = _docx_bytes(markdown, model)
        digest = hashlib.sha256(body).hexdigest()
    elif target == 'pdf':
        body = _pdf_bytes(markdown, model)
        digest = hashlib.sha256(body).hexdigest()
    else:
        raise ValueError(f'rel37_unknown_render_target:{target}')
    return Rendered(
        target=target,
        body=body,
        content_hash=digest,
        evidence=evidence,
        sections=sections,
        markdown=markdown,
    )
