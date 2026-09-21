"""Bounded REL37 narrative/cover acceptance corrections.

Records the bd8cf01 helper decisions, then asserts the corrected
explicit-clause cover, nonempty REL37 narrative applicability, and
location-bound PDF evidence. Same-source live fixture hash is retained.
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import time
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_ENV_KEYS = (
    'ADMIN_PASSWORD', 'SECRET_KEY', 'DATABASE_PATH', 'DATABASE_URL',
    'OPENAI_API_KEY', 'ANTHROPIC_API_KEY', 'GOOGLE_API_KEY',
    'REL2_SKIP_EXPORT_EVIDENCE', 'REL37_DATA_AI_DT_COMPILER',
)
os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
_ENV_BEFORE = {key: os.environ.get(key) for key in _ENV_KEYS}

_TMP = tempfile.mkdtemp(prefix='test_rel37_narr_accept_')


def _ensure_test_env():
    os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
    os.environ['ADMIN_PASSWORD'] = 'test-admin-password'
    os.environ['SECRET_KEY'] = 'test-secret-key'
    os.environ.setdefault('DATABASE_PATH', os.path.join(_TMP, 'narr_accept.db'))
    os.environ.setdefault(
        'DATABASE_URL', 'sqlite:///' + os.path.join(_TMP, 'narr_accept.db'))
    os.environ['OPENAI_API_KEY'] = ''
    os.environ['ANTHROPIC_API_KEY'] = ''
    os.environ['GOOGLE_API_KEY'] = ''
    os.environ['REL37_DATA_AI_DT_COMPILER'] = '1'


_ensure_test_env()

import app as app_mod  # noqa: E402

from release_engine_v3.rel37_apply import (  # noqa: E402
    REL37_APPLIED_KEY,
    REL37_CANONICAL_FW_KEY,
    REL37_HASH_KEY,
    REL37_MODEL_KEY,
    REL37_ORIGINAL_FW_KEY,
    REL37_SELECTION_REASON_KEY,
    REL37_SELECTION_SUPPORTED_KEY,
    serialize_model,
)
from release_engine_v3.rel37_canonical_document import CanonicalDocument  # noqa: E402
from release_engine_v3.rel37_compilers import compile_for_domain  # noqa: E402
from release_engine_v3.rel37_export_content_parity import (  # noqa: E402
    compare_cover_sector_to_docx,
    compare_cover_sector_to_pdf,
    compare_environment_narrative_to_docx,
    compare_environment_narrative_to_pdf,
    compare_model_to_docx,
    docx_cover_sector_value,
    docx_environment_section_text,
    extract_pdf_pages,
    gate_rel37_returned_bytes,
    hashed_narrative_requires_section_parity,
    inventory_docx_bytes,
    pdf_cover_sector_value,
    pdf_environment_section_text,
    rel37_returned_bytes_blockers,
)
from release_engine_v3.rel37_render import model_to_markdown, model_to_sections  # noqa: E402
from release_engine_v3.rel37_sector_context import (  # noqa: E402
    cover_sector_from_hashed_narrative,
    operating_context_clause,
    present_sector_label,
)

LIVE_FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37' / (
    'data_ar_live_12966c11_canonical.json')
DATA_EN_FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37' / (
    'data_en_saved_canonical_model.json')
LIVE_HASH = (
    'c87b5cc6e1dee58de38657ee721b2e61c509a0e48141e8f6796414293eb99d40')
DATA_EN_HASH = (
    'd73047a541e3a640d2ef0246b58ab50465e1a5464554633220cb9bf36f0bbdab')
GENERIC_OTHER = (
    'The organization reviews catalog coverage and consent under NDMO '
    'without restating the saved operating environment paragraph.')
_UID = {'n': 900}


class ImmediateThread:
    """In-process stand-in. Not the real worker/status/download path."""

    def __init__(self, target=None, args=(), kwargs=None, daemon=None):
        self.target = target
        self.args = args or ()
        self.kwargs = kwargs or {}

    def start(self):
        self.target(*self.args, **self.kwargs)

    def join(self, timeout=None):
        return None


def tearDownModule():
    for key, previous in _ENV_BEFORE.items():
        if previous is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = previous


def _baseline_first_mention(narrative: str, lang: str) -> str:
    """bd8cf01 cover helper: earliest UI-pair substring."""
    from release_engine_v3.rel37_sector_context import UI_SECTOR_PAIRS
    hay = str(narrative or '')
    lang_n = 'ar' if str(lang or '').lower().startswith('ar') else 'en'
    earliest = None
    chosen = ''
    for english, arabic in UI_SECTOR_PAIRS:
        for alias in (english, arabic):
            idx = hay.find(alias)
            if idx < 0:
                continue
            if earliest is None or idx < earliest:
                earliest = idx
                chosen = arabic if lang_n == 'ar' else english
    return chosen


def _baseline_keyword_requires(model: CanonicalDocument) -> bool:
    """bd8cf01 returned-byte narrative applicability."""
    text = str(getattr(model, 'environment_narrative', '') or '')
    return ('سياق تشغيلي' in text) or ('operating context' in text.lower())


def _load_json_model(path: Path) -> CanonicalDocument:
    payload = json.loads(path.read_text(encoding='utf-8'))
    model = CanonicalDocument.from_dict(payload)
    if not model.model_hash:
        model.compute_hashes()
    return model


def _compile(domain, lang, org_name, sector=''):
    return compile_for_domain(domain, {
        'domain': domain,
        'lang': lang,
        'org_name': org_name,
        'sector': sector,
        'document_type': 'strategy',
    })


def _sections_from_model(model: CanonicalDocument) -> dict:
    sections = model_to_sections(model)
    sections[REL37_APPLIED_KEY] = '1'
    sections[REL37_MODEL_KEY] = serialize_model(model)
    sections[REL37_HASH_KEY] = model.model_hash
    sections[REL37_SELECTION_SUPPORTED_KEY] = 'true'
    sections[REL37_SELECTION_REASON_KEY] = 'supported_selection'
    sections[REL37_CANONICAL_FW_KEY] = list(model.selected_frameworks)
    sections[REL37_ORIGINAL_FW_KEY] = list(model.selected_frameworks)
    return sections


def _next_user():
    _UID['n'] += 1
    uid = _UID['n']
    username = f'rel37acc{uid}'
    with app_mod.app.app_context():
        db = app_mod.get_db()
        db.execute(
            'INSERT OR IGNORE INTO users '
            '(id, username, password_hash, role, is_active) '
            'VALUES (?, ?, ?, ?, 1)',
            (uid, username, 'x', 'user'),
        )
        db.commit()
    return uid, username


def _client(uid, username):
    client = app_mod.app.test_client()
    csrf = f'rel37-acc-csrf-{uid}'
    with client.session_transaction() as sess:
        sess['user_id'] = uid
        sess['username'] = username
        sess['role'] = 'user'
        sess['csrf_token'] = csrf
    return client, {
        'X-CSRFToken': csrf,
        'Content-Type': 'application/json',
    }


def _persist(model: CanonicalDocument, *, db_sector='Healthcare') -> dict:
    uid, username = _next_user()
    client, headers = _client(uid, username)
    sections = _sections_from_model(model)
    content = model_to_markdown(model)
    with app_mod.app.app_context():
        db = app_mod.get_db()
        cur = db.execute(
            'INSERT INTO strategies '
            '(user_id, domain, org_name, sector, content, language, '
            ' document_title, sections_json, content_json) '
            'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (
                uid, 'Data Management', model.org_name, db_sector,
                content, model.lang, 'REL37 acceptance',
                json.dumps(sections, ensure_ascii=False),
                json.dumps({'sections': sections}, ensure_ascii=False),
            ),
        )
        sid = cur.lastrowid
        db.commit()
    return {
        'uid': uid,
        'username': username,
        'client': client,
        'headers': headers,
        'strategy_id': sid,
        'model': model,
        'content': content,
        'sections': sections,
        'lang': model.lang,
        'domain': 'Data Management',
    }


def _official_body(saved, *, sector='Healthcare', org_name=None):
    return {
        'content': saved['content'],
        'filename': 'rel37_accept_official',
        'language': saved['lang'],
        'org_name': org_name or 'REL33 P1 Data Management Org',
        'sector': sector,
        'doc_type': 'Strategy Document',
        'domain': 'Data Management',
        'document_type': 'strategy',
        'artifact_type': 'strategy',
        'generation_mode': 'drafting',
        'selected_frameworks': list(saved['model'].selected_frameworks),
        'strategy_id': saved['strategy_id'],
        'artifact_id': saved['strategy_id'],
    }


def _export(saved, body, fmt, *, immediate=True):
    payload = dict(body)
    payload.setdefault('document_type', 'strategy')
    payload.setdefault('artifact_type', 'strategy')
    payload.setdefault('generation_mode', 'drafting')
    payload['strategy_id'] = saved['strategy_id']
    payload['artifact_id'] = saved['strategy_id']
    ctx = patch('threading.Thread', ImmediateThread) if immediate else nullcontext()
    with ctx:
        resp = saved['client'].post(
            f'/api/generate-{fmt}-async',
            json=payload,
            headers=saved['headers'],
        )
    submit = resp.get_json(silent=True) or {}
    tid = submit.get('task_id')
    status = {}
    raw = b''
    download_http = 0
    if tid:
        deadline = time.time() + (2 if immediate else 45)
        while time.time() < deadline:
            status = saved['client'].get(
                f'/api/export-status/{tid}', headers=saved['headers']
            ).get_json(silent=True) or {}
            if status.get('status') in ('done', 'error'):
                break
            if not immediate:
                time.sleep(0.2)
            else:
                break
        if status.get('status') == 'done':
            dl = saved['client'].get(
                f'/api/export-download/{tid}', headers=saved['headers'])
            download_http = dl.status_code
            raw = dl.data or b''
    return {
        'submit_http': resp.status_code,
        'submit': submit,
        'status': status,
        'download_http': download_http,
        'bytes': raw,
        'task_id': tid,
    }


class nullcontext:
    def __enter__(self):
        return None

    def __exit__(self, *args):
        return False


def _mutate_docx_environment(raw: bytes, replacement: str) -> bytes:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    found = False
    taking = False
    for para in doc.paragraphs:
        text = para.text or ''
        if (
                ('البيئة التنظيمية والتهديدات' in text
                 or 'Business Environment' in text
                 or 'Regulatory Environment' in text)
                and len(text) < 80
        ):
            if text[:1].isdigit():
                continue
            taking = True
            continue
        if taking and (
                'تحليل الفجوات' in text or 'Gap Analysis' in text
                or 'Roadmap' in text):
            break
        if taking and text.strip():
            para.text = replacement
            found = True
            break
    if not found:
        raise AssertionError('environment paragraph was not found for mutation')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mutate_docx_cover_sector(raw: bytes, replacement: str) -> bytes:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    found = False
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() in ('القطاع', 'Sector'):
                row.cells[1].text = replacement
                found = True
                break
        if found:
            break
    if not found:
        raise AssertionError('cover sector cell was not found')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_AR_TEST_FONT = {'name': None}


def _ensure_ar_test_font():
    if _AR_TEST_FONT['name']:
        return _AR_TEST_FONT['name']
    path = '/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf'
    if not os.path.exists(path):
        return None
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont('NotoArTest', path))
    _AR_TEST_FONT['name'] = 'NotoArTest'
    return 'NotoArTest'


def _emit_actual_text(canv, logical: str):
    hex_text = 'FEFF' + str(logical or '').encode('utf-16-be').hex().upper()
    canv._code.append(f'/Span <</ActualText <{hex_text}>>> BDC')


def _end_actual_text(canv):
    canv._code.append('EMC')


def _draw_mixed(canv, x, y, text: str):
    """Visible mixed Arabic/Latin. Latin tokens stay on Helvetica."""
    import re as _re
    ar = _ensure_ar_test_font()
    canv.setFont('Helvetica', 10)
    canv.drawString(x, y, str(text or ''))
    cursor = x
    for token in _re.findall(r'[A-Za-z][A-Za-z0-9_-]*|/', str(text or '')):
        canv.setFont('Helvetica', 10)
        canv.drawString(cursor, y - 1, token)
        cursor += canv.stringWidth(token, 'Helvetica', 10) + 4
    if ar and any('\u0600' <= ch <= '\u06FF' for ch in str(text or '')):
        arabic = _re.sub(r'[A-Za-z][A-Za-z0-9._/-]*', ' ', str(text or ''))
        canv.setFont(ar, 10)
        canv.drawRightString(560, y, arabic)


def _pdf_owned(
        *,
        org_name: str,
        cover_sector: str,
        env_heading: str,
        env_paras: list,
        appendix_paras: list | None = None,
        extra_cover: str = '',
        pre_env_heading: str = '',
        pre_env_paras: list | None = None,
        env_continue_paras: list | None = None,
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    buf = io.BytesIO()
    canv = Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 36
    canv.setFont('Helvetica', 10)
    canv.drawString(36, y, f'Organization {org_name}')
    y -= 14
    canv.drawString(36, y, f'Sector {cover_sector}')
    y -= 14
    if extra_cover:
        canv.drawString(36, y, extra_cover)
    canv.showPage()
    if pre_env_heading or pre_env_paras:
        y = height - 36
        canv.setFont('Helvetica', 10)
        canv.drawString(36, y, pre_env_heading or 'Executive Summary')
        y -= 14
        for para in pre_env_paras or []:
            _emit_actual_text(canv, para)
            _draw_mixed(canv, 36, y, para)
            _end_actual_text(canv)
            y -= 14
        canv.showPage()
    y = height - 36
    canv.setFont('Helvetica', 10)
    canv.drawString(36, y, env_heading)
    y -= 14
    for para in env_paras:
        _emit_actual_text(canv, para)
        canv.setFont('Helvetica', 10)
        words = str(para).split()
        line = ''
        for word in words:
            candidate = (line + ' ' + word).strip()
            if canv.stringWidth(candidate, 'Helvetica', 10) > 480 and line:
                canv.drawString(36, y, line)
                y -= 12
                line = word
            else:
                line = candidate
        if line:
            canv.drawString(36, y, line)
            y -= 12
        import re as _re
        from release_engine_v3.rel37_export_content_parity import _semantic_latin_tokens
        latin = _semantic_latin_tokens(para)
        if latin:
            canv.setFont('Helvetica', 10)
            canv.drawString(36, y, ' '.join(latin))
            y -= 12
        _end_actual_text(canv)
    if env_continue_paras:
        canv.showPage()
        y = height - 36
        for para in env_continue_paras:
            _emit_actual_text(canv, para)
            canv.setFont('Helvetica', 10)
            canv.drawString(36, y, para)
            from release_engine_v3.rel37_export_content_parity import _semantic_latin_tokens
            latin = _semantic_latin_tokens(para)
            if latin:
                y -= 12
                canv.drawString(36, y, ' '.join(latin))
            _end_actual_text(canv)
            y -= 14
    canv.showPage()
    y = height - 36
    canv.setFont('Helvetica', 10)
    canv.drawString(36, y, 'Appendix')
    y -= 14
    for para in appendix_paras or []:
        _emit_actual_text(canv, para)
        _draw_mixed(canv, 36, y, str(para))
        _end_actual_text(canv)
        y -= 12
    canv.save()
    return buf.getvalue()


class FindingACoverNotIncidentalTests(unittest.TestCase):
    def test_baseline_first_mention_is_wrong_and_clause_wins(self):
        expected = present_sector_label('Banking/Finance', 'en')
        clause = operating_context_clause('Banking/Finance', 'en')
        model = _compile(
            'data', 'en', 'Healthcare Analytics', 'Banking/Finance')
        self.assertEqual(expected, 'Banking/Finance')
        self.assertIn(clause, model.environment_narrative)
        self.assertEqual(_baseline_first_mention(
            model.environment_narrative, 'en'), 'Healthcare')
        self.assertEqual(
            cover_sector_from_hashed_narrative(
                model.environment_narrative, 'en'),
            expected)
        self.assertNotEqual(
            cover_sector_from_hashed_narrative(
                model.environment_narrative, 'en'),
            model.sector if model.sector == 'Healthcare' else 'Healthcare')

    def test_arabic_org_label_before_banking_clause(self):
        org = 'مؤسسة رعاية صحية الأهلية'
        expected = present_sector_label('بنوك/مالي', 'ar')
        clause = operating_context_clause('بنوك/مالي', 'ar')
        model = _compile('data', 'ar', org, 'بنوك/مالي')
        self.assertEqual(expected, 'بنوك/مالي')
        self.assertIn(clause, model.environment_narrative)
        self.assertEqual(
            _baseline_first_mention(model.environment_narrative, 'ar'),
            'رعاية صحية')
        self.assertEqual(
            cover_sector_from_hashed_narrative(
                model.environment_narrative, 'ar'),
            expected)

    def test_incidental_mention_before_explicit_clause(self):
        expected = present_sector_label('Banking/Finance', 'en')
        narrative = (
            'Unlike Healthcare peers, Acme operates in the Banking/Finance '
            'sector operating context under NDMO and PDPL.')
        self.assertEqual(_baseline_first_mention(narrative, 'en'), 'Healthcare')
        self.assertEqual(
            cover_sector_from_hashed_narrative(narrative, 'en'), expected)

    def test_incidental_only_is_neutral(self):
        narrative = 'Healthcare Analytics reviews NDMO catalog coverage.'
        self.assertEqual(_baseline_first_mention(narrative, 'en'), 'Healthcare')
        self.assertEqual(cover_sector_from_hashed_narrative(narrative, 'en'), '')

    def test_ambiguous_explicit_clauses_are_neutral(self):
        narrative = (
            'It operates in the Banking/Finance sector operating context '
            'and later in the Energy sector operating context.')
        self.assertEqual(cover_sector_from_hashed_narrative(narrative, 'en'), '')

    def test_compiled_healthcare_org_docx_pdf_keep_banking(self):
        model = _compile(
            'data', 'en', 'Healthcare Analytics', 'Banking/Finance')
        before = model.compute_model_hash()
        expected = present_sector_label('Banking/Finance', 'en')
        saved = _persist(model, db_sector='Healthcare')
        official = _official_body(saved, sector='Energy')
        docx = _export(saved, official, 'docx')
        pdf = _export(saved, official, 'pdf')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        self.assertEqual(pdf['download_http'], 200, pdf['status'])
        self.assertIn('Healthcare Analytics', inventory_docx_bytes(docx['bytes'])['text'])
        self.assertEqual(docx_cover_sector_value(docx['bytes']), expected)
        self.assertEqual(compare_cover_sector_to_docx(model, docx['bytes']), [])
        self.assertEqual(compare_environment_narrative_to_docx(model, docx['bytes']), [])
        pages, meta = extract_pdf_pages(pdf['bytes'])
        self.assertTrue(meta.get('associated'), meta)
        self.assertGreaterEqual(meta.get('pages') or 0, 2)
        cover_value, cover_meta = pdf_cover_sector_value(pdf['bytes'])
        self.assertTrue(cover_meta.get('field_found'), cover_meta)
        self.assertIn(cover_value, ('Banking/Finance', 'بنوك/مالي'))
        self.assertEqual(compare_cover_sector_to_pdf(model, pdf['bytes']), [])
        self.assertEqual(compare_environment_narrative_to_pdf(model, pdf['bytes']), [])
        self.assertEqual(model.compute_model_hash(), before)

    def test_correct_banking_cover_not_rejected_for_healthcare_org(self):
        model = _compile(
            'data', 'en', 'Healthcare Analytics', 'Banking/Finance')
        saved = _persist(model, db_sector='Healthcare')
        docx = _export(saved, _official_body(saved, sector='Healthcare'), 'docx')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        self.assertEqual(docx_cover_sector_value(docx['bytes']), 'Banking/Finance')
        allowed, blockers = gate_rel37_returned_bytes(
            model, docx_bytes=docx['bytes'], route='docx')
        self.assertTrue(allowed, blockers)

    def test_legacy_neutral_does_not_acquire_client_or_incidental(self):
        model = _compile('data', 'en', 'Healthcare Analytics', '')
        self.assertEqual(
            cover_sector_from_hashed_narrative(
                model.environment_narrative, 'en'),
            '')
        saved = _persist(model, db_sector='Energy')
        docx = _export(saved, _official_body(saved, sector='Energy'), 'docx')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        cover = docx_cover_sector_value(docx['bytes'])
        self.assertNotEqual(cover, 'Energy')
        self.assertNotEqual(cover, 'Healthcare')
        self.assertIn(cover, ('—', '-', '', 'Government', 'حكومي'))


class FindingBNarrativeApplicabilityTests(unittest.TestCase):
    def test_data_en_nonempty_without_marker_phrases(self):
        model = _load_json_model(DATA_EN_FIXTURE)
        text = model.environment_narrative
        self.assertTrue(text.strip())
        self.assertNotIn('operating context', text.lower())
        self.assertNotIn('سياق تشغيلي', text)
        self.assertFalse(_baseline_keyword_requires(model))
        self.assertTrue(hashed_narrative_requires_section_parity(model))
        self.assertEqual(model.model_hash, DATA_EN_HASH)

    def test_positive_rendered_environment_equals_saved(self):
        model = _load_json_model(DATA_EN_FIXTURE)
        before = model.model_hash
        saved = _persist(model, db_sector='Government')
        docx = _export(saved, _official_body(saved, sector='Government'), 'docx')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        env = docx_environment_section_text(docx['bytes'])
        for para in text_paras(model.environment_narrative):
            self.assertIn(para, env)
        self.assertEqual(compare_environment_narrative_to_docx(model, docx['bytes']), [])
        self.assertEqual(model.compute_model_hash(), before)

    def test_causal_env_swap_fails_corrected_gate_not_keyword(self):
        model = _load_json_model(DATA_EN_FIXTURE)
        saved = _persist(model, db_sector='Government')
        good = _export(saved, _official_body(saved, sector='Government'), 'docx')
        self.assertEqual(good['download_http'], 200, good['status'])
        mutated = _mutate_docx_environment(good['bytes'], GENERIC_OTHER)
        self.assertTrue(inventory_docx_bytes(mutated)['paragraphs'])
        self.assertEqual(docx_cover_sector_value(mutated), docx_cover_sector_value(good['bytes']))
        standalone = compare_environment_narrative_to_docx(model, mutated)
        self.assertTrue(standalone, standalone)
        self.assertFalse(_baseline_keyword_requires(model))
        baseline_blockers = []
        if _baseline_keyword_requires(model):
            baseline_blockers.extend(standalone)
        baseline_blockers.extend(compare_cover_sector_to_docx(model, mutated))
        self.assertFalse(
            any('narrative' in item for item in baseline_blockers),
            baseline_blockers)
        corrected = rel37_returned_bytes_blockers(
            model, docx_bytes=mutated, route='docx')
        self.assertTrue(any('narrative' in item for item in corrected), corrected)
        allowed, gated = gate_rel37_returned_bytes(
            model, docx_bytes=mutated, route='docx')
        self.assertFalse(allowed)
        self.assertFalse(any(
            item in gated for item in (
                'docx_bytes_missing', 'pdf_bytes_missing',
                'pdf_extraction_unreliable')), gated)
        self.assertEqual(model.model_hash, DATA_EN_HASH)

    def test_cyber_erm_global_not_selected(self):
        payload = json.loads(DATA_EN_FIXTURE.read_text(encoding='utf-8'))
        for domain in ('cyber', 'erm', 'global'):
            payload['domain'] = domain
            other = CanonicalDocument.from_dict(payload)
            self.assertFalse(hashed_narrative_requires_section_parity(other), domain)


def text_paras(narrative: str):
    return [part.strip() for part in str(narrative).split('\n\n') if part.strip()]


class FindingCLocationBoundPdfTests(unittest.TestCase):
    def setUp(self):
        self.model = _load_json_model(DATA_EN_FIXTURE)
        self.heading = 'Business Environment and Drivers'
        self.paras = text_paras(self.model.environment_narrative)
        self.good = _pdf_owned(
            org_name=self.model.org_name,
            cover_sector='—',
            env_heading=self.heading,
            env_paras=self.paras,
        )

    def test_matching_pdf_positive(self):
        self.assertTrue(self.good.startswith(b'%PDF'))
        pages, meta = extract_pdf_pages(self.good)
        self.assertTrue(meta.get('associated'), meta)
        self.assertGreaterEqual(len(pages), 2)
        self.assertEqual(compare_environment_narrative_to_pdf(self.model, self.good), [])
        self.assertEqual(compare_cover_sector_to_pdf(self.model, self.good), [])

    def test_appendix_only_narrative_fails(self):
        self.assertEqual(compare_environment_narrative_to_pdf(self.model, self.good), [])
        mutated = _pdf_owned(
            org_name=self.model.org_name,
            cover_sector='—',
            env_heading=self.heading,
            env_paras=['Unrelated environment leftover paragraph.'],
            appendix_paras=self.paras,
        )
        section, meta = pdf_environment_section_text(mutated)
        self.assertTrue(meta.get('environment_associated'), meta)
        self.assertNotIn(self.paras[0], section)
        blockers = compare_environment_narrative_to_pdf(self.model, mutated)
        self.assertTrue(any('narrative' in item for item in blockers), blockers)

    def test_cover_only_narrative_fails(self):
        self.assertEqual(compare_environment_narrative_to_pdf(self.model, self.good), [])
        mutated = _pdf_owned(
            org_name=self.model.org_name,
            cover_sector='—',
            env_heading=self.heading,
            env_paras=['Cover-only leftover environment.'],
            extra_cover=self.paras[0],
        )
        blockers = compare_environment_narrative_to_pdf(self.model, mutated)
        self.assertTrue(blockers, blockers)

    def test_wrong_cover_field_with_sector_in_org_name(self):
        banking = _compile(
            'data', 'en', 'Healthcare Analytics', 'Banking/Finance')
        good = _pdf_owned(
            org_name='Healthcare Analytics',
            cover_sector='Banking/Finance',
            env_heading=self.heading,
            env_paras=text_paras(banking.environment_narrative),
        )
        self.assertEqual(compare_cover_sector_to_pdf(banking, good), [])
        mutated = _pdf_owned(
            org_name='Healthcare Analytics',
            cover_sector='Energy',
            env_heading=self.heading,
            env_paras=text_paras(banking.environment_narrative),
        )
        value, meta = pdf_cover_sector_value(mutated)
        self.assertEqual(value, 'Energy')
        self.assertTrue(meta.get('field_found'), meta)
        blockers = compare_cover_sector_to_pdf(banking, mutated)
        self.assertTrue(any('cover_sector' in item for item in blockers), blockers)

    def test_neutral_energy_injection_rejected(self):
        self.assertEqual(compare_cover_sector_to_pdf(self.model, self.good), [])
        mutated = _pdf_owned(
            org_name=self.model.org_name,
            cover_sector='Energy',
            env_heading=self.heading,
            env_paras=self.paras,
        )
        blockers = compare_cover_sector_to_pdf(self.model, mutated)
        self.assertTrue(any('injected' in item or 'cover_sector' in item
                            for item in blockers), blockers)
        self.assertFalse(any('Healthcare' in item for item in blockers))


class FinalRouteByteRefusalTests(unittest.TestCase):
    def test_fault_injected_bytes_are_not_downloaded(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='docx', lang='ar'):
            if docx_bytes:
                docx_bytes = _mutate_docx_environment(
                    docx_bytes, GENERIC_OTHER)
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(saved, _official_body(saved), 'docx')
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'PK'))
        self.assertEqual(model.model_hash, LIVE_HASH)

    def test_real_thread_worker_status_download_refuses_bad_bytes(self):
        model = _load_json_model(LIVE_FIXTURE)
        saved = _persist(model, db_sector='Healthcare')
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='docx', lang='ar'):
            if docx_bytes:
                docx_bytes = _mutate_docx_cover_sector(docx_bytes, 'Energy')
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'docx', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertIn(result['status'].get('status'), ('error', None, ''))
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'PK'))

    def test_owner_positive_and_csrf_cross_user(self):
        model = _load_json_model(LIVE_FIXTURE)
        saved = _persist(model, db_sector='Healthcare')
        good = _export(saved, _official_body(saved), 'docx')
        self.assertEqual(good['download_http'], 200, good['status'])
        self.assertTrue(good['bytes'].startswith(b'PK'))
        stale = dict(saved['headers'])
        stale['X-CSRFToken'] = 'stale-csrf-token'
        resp = saved['client'].post(
            '/api/generate-docx-async',
            json=_official_body(saved),
            headers=stale,
        )
        self.assertIn(resp.status_code, (400, 403))
        other_uid, other_name = _next_user()
        other_client, other_headers = _client(other_uid, other_name)
        other = dict(saved)
        other['client'] = other_client
        other['headers'] = other_headers
        stolen = _export(other, _official_body(saved), 'docx')
        self.assertNotEqual(stolen['download_http'], 200)
        self.assertFalse(stolen['bytes'].startswith(b'PK'))

    def test_same_source_hash_unchanged(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        self.assertEqual(model.compute_model_hash(), LIVE_HASH)


EXPECTED_AR_ENV = (
    'تعمل الجهة في سياق تشغيلي لقطاع بنوك/مالي، ضمن بيئة تنظيمية '
    'تتطلب حوكمة بيانات وطنية وفق NDMO وحماية بيانات شخصية وفق PDPL.')
REMOVED_LATIN = (
    'تعمل الجهة في سياق تشغيلي لقطاع بنوك/مالي، ضمن بيئة تنظيمية '
    'تتطلب حوكمة بيانات وطنية وفق  وحماية بيانات شخصية وفق .')
SUBSTITUTED_LATIN = (
    'تعمل الجهة في سياق تشغيلي لقطاع بنوك/مالي، ضمن بيئة تنظيمية '
    'تتطلب حوكمة بيانات وطنية وفق ISO9001 وحماية بيانات شخصية وفق GDPR.')
SWAPPED_LATIN = (
    'تعمل الجهة في سياق تشغيلي لقطاع بنوك/مالي، ضمن بيئة تنظيمية '
    'تتطلب حوكمة بيانات وطنية وفق PDPL وحماية بيانات شخصية وفق NDMO.')
ORG_LATIN = EXPECTED_AR_ENV + ' AcmeBankGroup'
ORG_LATIN_CHANGED = EXPECTED_AR_ENV + ' OtherHoldingsLLC'
ENV_HEAD_EN = 'Business Environment and Drivers'
SPECIMEN_EN = 'NDMO is the first reference; PDPL is the second reference.'
SPECIMEN_EN_CHANGED = 'PDPL is the first reference; NDMO is the second reference.'
SPECIMEN_WRONG_ENV = 'The environment sentence is generic leftover prose.'


def _specimen_env_pdf(
        *,
        env_visible: str,
        env_actual: str | None,
        summary_visible: str | None = None,
        summary_actual: str | None = None,
        same_page_summary: bool = False,
        appendix_visible: str | None = None,
        appendix_actual: str | None = None,
        same_page_appendix: bool = False,
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    buf = io.BytesIO()
    canv = Canvas(buf, pagesize=A4)
    canv.setFont('Helvetica', 10)
    canv.drawString(36, 800, 'Organization SpecimenOrg')
    canv.drawString(36, 786, 'Sector Banking/Finance')
    canv.showPage()
    y = 800
    canv.setFont('Helvetica', 10)
    if summary_visible or summary_actual:
        canv.drawString(36, y, 'Executive Summary')
        y -= 16
        if summary_actual:
            _emit_actual_text(canv, summary_actual)
        canv.drawString(36, y, summary_visible or summary_actual or '')
        if summary_actual:
            _end_actual_text(canv)
        y -= 20
        if not same_page_summary:
            canv.showPage()
            y = 800
            canv.setFont('Helvetica', 10)
    canv.drawString(36, y, ENV_HEAD_EN)
    y -= 16
    if env_actual:
        _emit_actual_text(canv, env_actual)
    canv.drawString(36, y, env_visible)
    if env_actual:
        _end_actual_text(canv)
    y -= 20
    if appendix_visible or appendix_actual:
        if not same_page_appendix:
            canv.showPage()
            y = 800
            canv.setFont('Helvetica', 10)
        canv.drawString(36, y, 'Appendix')
        y -= 16
        if appendix_actual:
            _emit_actual_text(canv, appendix_actual)
        canv.drawString(36, y, appendix_visible or appendix_actual or '')
        if appendix_actual:
            _end_actual_text(canv)
    canv.save()
    return buf.getvalue()


def _baseline_ascii_drop(text: str) -> str:
    """2691b50 fallback: strip every Latin token from both sides."""
    import re as _re
    from release_engine_v3.rel37_export_content_parity import _norm
    return _norm(_re.sub(r'[A-Za-z][A-Za-z0-9._/-]*', ' ', text))


class _EnvModel:
    def __init__(self, narrative, lang='ar', domain='data', org='جهة'):
        self.environment_narrative = narrative
        self.lang = lang
        self.domain = domain
        self.org_name = org
        self.model_hash = LIVE_HASH
        self.sector = ''
        self.selected_frameworks = []

    def compute_model_hash(self):
        return self.model_hash


class FindingLatinTokenPdfTests(unittest.TestCase):
    def test_helper_baseline_false_match_and_correction(self):
        from release_engine_v3.rel37_export_content_parity import (
            _paragraph_in_pdf_section,
        )
        cases = {
            'unchanged': (EXPECTED_AR_ENV, EXPECTED_AR_ENV, True),
            'removed': (EXPECTED_AR_ENV, REMOVED_LATIN, False),
            'substituted': (EXPECTED_AR_ENV, SUBSTITUTED_LATIN, False),
            'swapped': (EXPECTED_AR_ENV, SWAPPED_LATIN, False),
            'org_changed': (ORG_LATIN, ORG_LATIN_CHANGED, False),
        }
        for name, (expected, observed, accept) in cases.items():
            baseline = (
                bool(_baseline_ascii_drop(expected))
                and _baseline_ascii_drop(expected)
                in _baseline_ascii_drop(observed))
            corrected = _paragraph_in_pdf_section(expected, observed)
            if name == 'unchanged':
                self.assertTrue(baseline, name)
                self.assertTrue(corrected, name)
            else:
                self.assertTrue(baseline, name)
                self.assertFalse(corrected, name)
            self.assertEqual(corrected, accept, name)

    def test_helper_does_not_reverse_wrong_logical_text(self):
        from release_engine_v3.rel37_export_content_parity import (
            _paragraph_in_pdf_section,
            _paragraph_pdf_blockers,
        )
        expected = 'NDMO هو المرجع المعتمد وليس PDPL'
        observed = 'PDPL هو المرجع المعتمد وليس NDMO'
        self.assertFalse(_paragraph_in_pdf_section(expected, observed))
        blockers = _paragraph_pdf_blockers(
            0, expected, observed, visible=observed, actual=observed)
        self.assertTrue(blockers, blockers)
        visual = (
            '، مع ضغطPDPL  وحماية بيانات شخصية وفقNDMO تعمل الجهة في '
            'بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق')
        logical = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق '
            'NDMO وحماية بيانات شخصية وفق PDPL')
        self.assertTrue(
            not _paragraph_pdf_blockers(
                0, logical, visual, visible=visual, actual=logical),
            _paragraph_pdf_blockers(
                0, logical, visual, visible=visual, actual=logical))

    def _env_pdf(self, env_paras, **kwargs):
        return _pdf_owned(
            org_name='جهة',
            cover_sector='بنوك/مالي',
            env_heading=ENV_HEAD_EN,
            env_paras=env_paras,
            **kwargs,
        )

    def test_matching_mixed_script_pdf_positive(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertTrue(good.startswith(b'%PDF'))
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])

    def test_environment_visual_rtl_reconciles_to_logical_actualtext(self):
        compiled = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق '
            'NDMO وحماية بيانات شخصية وفق PDPL، مع ضغط متزايد على جودة '
            'البيانات والكتالوج وإدارة الموافقات وحقوق أصحاب البيانات.')
        visual = (
            '، مع ضغطPDPL  وحماية بيانات شخصية وفقNDMO تعمل الجهة في '
            'بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق '
            'متزايد على جودة البيانات والكتالوج وإدارة الموافقات وحقوق '
            'أصحاب البيانات.')
        raw = _specimen_env_pdf(
            env_visible=visual, env_actual=compiled)
        model = _EnvModel(compiled)
        self.assertEqual(compare_environment_narrative_to_pdf(model, raw), [])

    def test_missing_acronym_rejected(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = self._env_pdf([REMOVED_LATIN])
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(any('latin_missing' in item or 'narrative' in item
                            for item in blockers), blockers)
        self.assertFalse(any(
            item in blockers for item in (
                'pdf_extraction_unreliable',
                'pdf_environment_section_unassociated',
                'docx_bytes_missing')), blockers)

    def test_substituted_acronym_rejected(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        blockers = compare_environment_narrative_to_pdf(
            model, self._env_pdf([SUBSTITUTED_LATIN]))
        self.assertTrue(any('latin_missing' in item or 'narrative' in item
                            for item in blockers), blockers)

    def test_token_position_swap_rejected(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        blockers = compare_environment_narrative_to_pdf(
            model, self._env_pdf([SWAPPED_LATIN]))
        self.assertTrue(any('latin_order' in item or 'narrative' in item
                            for item in blockers), blockers)

    def test_changed_latin_org_rejected(self):
        model = _EnvModel(ORG_LATIN)
        good = self._env_pdf([ORG_LATIN])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        blockers = compare_environment_narrative_to_pdf(
            model, self._env_pdf([ORG_LATIN_CHANGED]))
        self.assertTrue(any('latin_missing' in item or 'narrative' in item
                            for item in blockers), blockers)

    def test_earlier_page_is_not_environment(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = self._env_pdf(
            ['تعمل الجهة في سياق تشغيلي لقطاع بنوك/مالي دون النص المحفوظ.'],
            pre_env_heading='Executive Summary',
            pre_env_paras=[EXPECTED_AR_ENV],
        )
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(blockers, blockers)

    def test_appendix_only_still_rejected(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = self._env_pdf(
            ['Unrelated environment leftover paragraph.'],
            appendix_paras=[EXPECTED_AR_ENV],
        )
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(blockers, blockers)

    def test_cover_only_still_rejected(self):
        model = _EnvModel(EXPECTED_AR_ENV)
        good = self._env_pdf([EXPECTED_AR_ENV])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = self._env_pdf(
            ['Cover-only leftover environment.'],
            extra_cover=EXPECTED_AR_ENV,
        )
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(blockers, blockers)

    def test_environment_continuation_accepted(self):
        first, second = EXPECTED_AR_ENV, (
            'تشمل المحركات التشغيلية اكتمال التصنيف وفق NDMO.')
        model = _EnvModel(first + '\n\n' + second)
        good = self._env_pdf([first], env_continue_paras=[second])
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])

    def test_real_thread_refuses_latin_stripped_pdf(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='pdf', lang='ar'):
            if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
                pdf_bytes = self._env_pdf([REMOVED_LATIN])
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'pdf', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'%PDF'))
        self.assertEqual(model.model_hash, LIVE_HASH)

    def test_owner_positive_pdf_still_downloads(self):
        model = _load_json_model(LIVE_FIXTURE)
        saved = _persist(model, db_sector='Healthcare')
        good = _export(saved, _official_body(saved), 'pdf')
        self.assertEqual(good['download_http'], 200, good['status'])
        self.assertTrue(good['bytes'].startswith(b'%PDF'))
        self.assertEqual(
            compare_environment_narrative_to_pdf(model, good['bytes']), [])
        self.assertEqual(model.model_hash, LIVE_HASH)

    def test_matching_specimen_accepted_before_negatives(self):
        model = _EnvModel(SPECIMEN_EN, lang='en', org='SpecimenOrg')
        good = _specimen_env_pdf(env_visible=SPECIMEN_EN, env_actual=SPECIMEN_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])

    def test_changed_visible_correct_actualtext_rejected(self):
        model = _EnvModel(SPECIMEN_EN, lang='en', org='SpecimenOrg')
        good = _specimen_env_pdf(env_visible=SPECIMEN_EN, env_actual=SPECIMEN_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _specimen_env_pdf(
            env_visible=SPECIMEN_EN_CHANGED, env_actual=SPECIMEN_EN)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(
            any('actual_visible_disagree' in item for item in blockers),
            blockers)

    def test_both_changed_and_no_actualtext_rejected(self):
        model = _EnvModel(SPECIMEN_EN, lang='en', org='SpecimenOrg')
        good = _specimen_env_pdf(env_visible=SPECIMEN_EN, env_actual=SPECIMEN_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        both = _specimen_env_pdf(
            env_visible=SPECIMEN_EN_CHANGED, env_actual=SPECIMEN_EN_CHANGED)
        none = _specimen_env_pdf(
            env_visible=SPECIMEN_EN_CHANGED, env_actual=None)
        self.assertTrue(compare_environment_narrative_to_pdf(model, both))
        self.assertTrue(compare_environment_narrative_to_pdf(model, none))

    def test_same_page_summary_is_not_environment(self):
        model = _EnvModel(SPECIMEN_EN, lang='en', org='SpecimenOrg')
        good = _specimen_env_pdf(env_visible=SPECIMEN_EN, env_actual=SPECIMEN_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _specimen_env_pdf(
            env_visible=SPECIMEN_WRONG_ENV, env_actual=SPECIMEN_WRONG_ENV,
            summary_visible=SPECIMEN_EN, summary_actual=SPECIMEN_EN,
            same_page_summary=True)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(blockers, blockers)
        self.assertFalse(
            any(item in blockers for item in (
                'pdf_extraction_unreliable',
                'pdf_environment_section_unassociated',
                'docx_bytes_missing')), blockers)

    def test_same_page_appendix_is_not_environment(self):
        model = _EnvModel(SPECIMEN_EN, lang='en', org='SpecimenOrg')
        good = _specimen_env_pdf(env_visible=SPECIMEN_EN, env_actual=SPECIMEN_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _specimen_env_pdf(
            env_visible=SPECIMEN_WRONG_ENV, env_actual=SPECIMEN_WRONG_ENV,
            appendix_visible=SPECIMEN_EN, appendix_actual=SPECIMEN_EN,
            same_page_appendix=True)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self.assertTrue(blockers, blockers)

    def test_real_thread_refuses_visible_actualtext_conflict(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='pdf', lang='ar'):
            if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
                pdf_bytes = _specimen_env_pdf(
                    env_visible=SPECIMEN_EN_CHANGED, env_actual=SPECIMEN_EN)
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'pdf', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'%PDF'))
        self.assertEqual(model.model_hash, LIVE_HASH)

    def test_real_thread_refuses_same_page_summary_only(self):
        model = _load_json_model(LIVE_FIXTURE)
        saved = _persist(model, db_sector='Healthcare')
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='pdf', lang='ar'):
            if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
                pdf_bytes = _specimen_env_pdf(
                    env_visible=SPECIMEN_WRONG_ENV,
                    env_actual=SPECIMEN_WRONG_ENV,
                    summary_visible=SPECIMEN_EN,
                    summary_actual=SPECIMEN_EN,
                    same_page_summary=True)
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'pdf', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'%PDF'))
        self.assertEqual(model.model_hash, LIVE_HASH)


COMPLETE_EN = 'NDMO is not approved for unrestricted sharing.'
PAINTED_EN_DELETED = 'NDMO is approved for unrestricted sharing.'
COMPLETE_AR = 'الجهة لا تشارك البيانات وفق PDPL.'
PAINTED_AR_DELETED = 'الجهة تشارك البيانات وفق PDPL.'
COMPLETE_NUM = 'NDMO requires a target of 100%.'
PAINTED_NUM_CHANGED = 'NDMO requires a target of 10%.'
WRAP_EN = 'NDMO is the first reference; PDPL is the second reference.'


def _complete_repr_pdf(visible: str, actual: str) -> bytes:
    """Test-owned readable PDF with independently painted vs ActualText."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    buf = io.BytesIO()
    canv = Canvas(buf, pagesize=A4)
    canv.setFont('Helvetica', 10)
    canv.drawString(36, 800, 'Organization SpecimenOrg')
    canv.drawString(36, 786, 'Sector Banking/Finance')
    canv.showPage()
    y = 800
    canv.setFont('Helvetica', 10)
    canv.drawString(36, y, ENV_HEAD_EN)
    y -= 16
    if actual:
        _emit_actual_text(canv, actual)
    logical_paint = str(visible or '')
    if logical_paint:
        hex_text = 'FEFF' + logical_paint.encode('utf-16-be').hex().upper()
        canv._code.append(f'<{hex_text}> Tj')
    lines = str(visible or '').splitlines() or ['']
    for line in lines:
        if any('\u0600' <= ch <= '\u06FF' for ch in line):
            import re as _re
            latin = _re.findall(r'[A-Za-z][A-Za-z0-9%._-]*', line)
            if latin:
                canv.setFont('Helvetica', 10)
                canv.drawString(36, y, ' '.join(latin))
                y -= 14
            ar_font = _ensure_ar_test_font()
            arabic = _re.sub(r'[A-Za-z0-9%._/-]+', ' ', line)
            if ar_font and arabic.strip():
                canv.setFont(ar_font, 10)
                canv.drawRightString(560, y, arabic.strip())
                y -= 14
            elif arabic.strip():
                canv.setFont('Helvetica', 10)
                canv.drawString(36, y, arabic.strip())
                y -= 14
        else:
            canv.setFont('Helvetica', 10)
            canv.drawString(36, y, line)
            y -= 14
    if actual:
        _end_actual_text(canv)
    canv.save()
    return buf.getvalue()


class CompleteRepresentationPdfTests(unittest.TestCase):
    def _helper(self, expected, painted, actual):
        from release_engine_v3.rel37_export_content_parity import (
            _paragraph_pdf_blockers,
            _visible_reconciled_to_actual,
        )
        vis_cmp, how = _visible_reconciled_to_actual(painted, actual)
        blockers = _paragraph_pdf_blockers(
            0, expected, painted, visible=painted, actual=actual)
        return vis_cmp, how, blockers

    def _assert_narrative_blocker(self, blockers):
        self.assertTrue(
            any(
                'actual_visible_disagree' in item
                or 'narrative_missing' in item
                or 'latin_missing' in item
                or 'latin_order' in item
                for item in blockers
            ),
            blockers,
        )
        self.assertFalse(any(
            'painted_unestablished' in item
            or item in (
                'pdf_extraction_unreliable',
                'pdf_environment_section_unassociated',
                'docx_bytes_missing',
                'pdf_bytes_missing',
            )
            for item in blockers
        ), blockers)

    def test_helper_complete_content_controls(self):
        vis, how, blockers = self._helper(WRAP_EN, WRAP_EN, WRAP_EN)
        self.assertEqual(how, 'same')
        self.assertFalse(blockers, (how, blockers, vis))
        wrapped = 'NDMO is the first reference;\nPDPL is the second reference.'
        vis, how, blockers = self._helper(WRAP_EN, wrapped, WRAP_EN)
        self.assertIn(how, ('same', 'visible_wraps_actual'))
        self.assertFalse(blockers, (how, blockers, vis))
        vis, how, blockers = self._helper(
            COMPLETE_EN, SPECIMEN_EN_CHANGED, SPECIMEN_EN_CHANGED)
        self.assertTrue(blockers, (how, blockers, vis))

    def test_helper_mashed_visual_leftover_does_not_replace_complete_actual(self):
        para = (
            'تعمل منظمة بيانات عربي في سياق تشغيلي لقطاع حكومي، ضمن بيئة '
            'تنظيمية تتطلب حوكمة بيانات وطنية وفق NDMO وحماية بيانات شخصية '
            'وفق PDPL، مع ضغط متزايد على جودة البيانات والكتالوج وإدارة '
            'الموافقات وحقوق أصحاب البيانات.'
        )
        mashed = (
            'طو تانايب ةمكوح بلطتت ةيميظنت ةئيب نمض ،يموكح عاطقل يليغشت '
            'قايس يف يبرع تانايب ةمظنم لمعتNDMO تعمل منظمة بيانات عربي في '
            'سياق تشغيلي لقطاع حكومي، ضمن بيئة تنظيمية تتطلب حوكمة بيانات '
            'وطنية وفق\n، مع ضغط متزايد على جودة البيانات والكتالوج وإدارة '
            'الموافقات وحقوق أصحاب البيانات.PDPL وحماية بيانات شخصية وفق'
            'NDMO PDPL'
        )
        vis, how, blockers = self._helper(para, mashed, para)
        self.assertFalse(blockers, (how, blockers, vis))
        vis, how, blockers = self._helper(
            COMPLETE_EN, mashed + '\n' + PAINTED_EN_DELETED, COMPLETE_EN)
        self._assert_narrative_blocker(blockers)

    def test_helper_refuses_incomplete_painted_subsequence(self):
        cases = (
            ('english_deletion', COMPLETE_EN, PAINTED_EN_DELETED, COMPLETE_EN),
            ('arabic_deletion', COMPLETE_AR, PAINTED_AR_DELETED, COMPLETE_AR),
            ('numeric_change', COMPLETE_NUM, PAINTED_NUM_CHANGED, COMPLETE_NUM),
        )
        for name, expected, painted, actual in cases:
            vis, how, blockers = self._helper(expected, painted, actual)
            self.assertNotEqual(
                how,
                'visible_wraps_actual',
                (name, how, vis),
            )
            self.assertNotEqual(vis, expected, (name, vis, how))
            self._assert_narrative_blocker(blockers)

    def test_matching_pdf_positive_before_each_painted_mutation(self):
        for expected, lang in (
                (COMPLETE_EN, 'en'),
                (COMPLETE_AR, 'ar'),
                (COMPLETE_NUM, 'en'),
                (WRAP_EN, 'en'),
        ):
            model = _EnvModel(expected, lang=lang, org='SpecimenOrg')
            good = _complete_repr_pdf(expected, expected)
            self.assertTrue(good.startswith(b'%PDF'))
            self.assertEqual(
                compare_environment_narrative_to_pdf(model, good), [],
                expected)

    def test_painted_english_negation_removed_actualtext_correct(self):
        model = _EnvModel(COMPLETE_EN, lang='en', org='SpecimenOrg')
        good = _complete_repr_pdf(COMPLETE_EN, COMPLETE_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(PAINTED_EN_DELETED, COMPLETE_EN)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self._assert_narrative_blocker(blockers)

    def test_painted_arabic_negation_removed_actualtext_correct(self):
        model = _EnvModel(COMPLETE_AR, lang='ar', org='SpecimenOrg')
        good = _complete_repr_pdf(COMPLETE_AR, COMPLETE_AR)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(PAINTED_AR_DELETED, COMPLETE_AR)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self._assert_narrative_blocker(blockers)

    def test_painted_numeric_target_changed_actualtext_correct(self):
        model = _EnvModel(COMPLETE_NUM, lang='en', org='SpecimenOrg')
        good = _complete_repr_pdf(COMPLETE_NUM, COMPLETE_NUM)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(PAINTED_NUM_CHANGED, COMPLETE_NUM)
        blockers = compare_environment_narrative_to_pdf(model, mutated)
        self._assert_narrative_blocker(blockers)

    def test_valid_wrap_mixed_script_and_continuation_still_accepted(self):
        wrap_vis = 'NDMO is the first reference;\nPDPL is the second reference.'
        model = _EnvModel(WRAP_EN, lang='en', org='SpecimenOrg')
        wrapped = _complete_repr_pdf(wrap_vis, WRAP_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, wrapped), [])
        mixed = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق '
            'NDMO وحماية بيانات شخصية وفق PDPL.')
        mixed_model = _EnvModel(mixed, lang='ar', org='جهة')
        mixed_pdf = _complete_repr_pdf(mixed, mixed)
        self.assertEqual(
            compare_environment_narrative_to_pdf(mixed_model, mixed_pdf), [])
        first, second = EXPECTED_AR_ENV, (
            'تشمل المحركات التشغيلية اكتمال التصنيف وفق NDMO.')
        cont_model = _EnvModel(first + '\n\n' + second)
        continued = _pdf_owned(
            org_name='جهة',
            cover_sector='بنوك/مالي',
            env_heading=ENV_HEAD_EN,
            env_paras=[first],
            env_continue_paras=[second],
        )
        self.assertEqual(
            compare_environment_narrative_to_pdf(cont_model, continued), [])

    def test_real_thread_owner_positive_and_painted_mutation_refused(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        good = _export(
            saved, _official_body(saved), 'pdf', immediate=False)
        self.assertEqual(good['download_http'], 200, good['status'])
        self.assertTrue(good['bytes'].startswith(b'%PDF'))
        self.assertEqual(
            compare_environment_narrative_to_pdf(model, good['bytes']), [])
        self.assertEqual(model.model_hash, LIVE_HASH)
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='pdf', lang='ar'):
            if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
                pdf_bytes = _complete_repr_pdf(
                    PAINTED_EN_DELETED, COMPLETE_EN)
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'pdf', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'%PDF'))
        self.assertEqual(model.model_hash, LIVE_HASH)
        self.assertEqual(model.compute_model_hash(), LIVE_HASH)


ORDER_AR_EXP = 'تسمح الجهة بجمع البيانات ولا تسمح بنشرها.'
ORDER_AR_VIS = 'تسمح الجهة بنشرها ولا تسمح بجمع البيانات.'
NEG_MIX_EXP = 'NDMO يجيز مشاركة البيانات.'
NEG_MIX_VIS = 'NDMO لا يجيز مشاركة البيانات.'
OWNER_EXP = (
    'The organization assigns the following work to Alice for approval '
    'and Bob for review.')
OWNER_VIS = (
    'The organization assigns the following work to Bob for approval '
    'and Alice for review.')
NUM_AR_EXP = 'تبلغ القيمة المستهدفة للمؤشر 100% وفق NDMO.'
NUM_AR_VIS = 'تبلغ القيمة المستهدفة للمؤشر 10% وفق NDMO.'
CONTRA_EXTRA = 'The appendix restates an unrelated sharing obligation.'


def _emit_hidden_logical(canv, logical: str):
    canv._code.append('q')
    canv._code.append('3 Tr')
    ar = _ensure_ar_test_font()
    text = str(logical or '')
    if ar and any('\u0600' <= ch <= '\u06FF' for ch in text):
        canv.setFont(ar, 10)
        canv.drawRightString(560, 2, text)
    else:
        canv.setFont('Helvetica', 10)
        canv.drawString(36, 2, text)
    canv._code.append('Q')


def _provenance_pdf(
        *,
        visible: str,
        actual: str,
        hidden_logical: str | None = None,
        extra_visible: str | None = None,
        emit_hidden: bool = False,
) -> bytes:
    """Readable test PDF with optional file-proven hidden logical."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    buf = io.BytesIO()
    canv = Canvas(buf, pagesize=A4)
    canv.setFont('Helvetica', 10)
    canv.drawString(36, 800, 'Organization SpecimenOrg')
    canv.drawString(36, 786, 'Sector Banking/Finance')
    canv.showPage()
    y = 800
    canv.setFont('Helvetica', 10)
    canv.drawString(36, y, ENV_HEAD_EN)
    y -= 18
    if actual:
        _emit_actual_text(canv, actual)
    if emit_hidden:
        _emit_hidden_logical(canv, hidden_logical or actual)
    lines = str(visible or '').splitlines() or ['']
    for line in lines:
        if any('\u0600' <= ch <= '\u06FF' for ch in line):
            import re as _re
            latin = _re.findall(r'[A-Za-z][A-Za-z0-9%._-]*', line)
            if latin:
                canv.setFont('Helvetica', 10)
                canv.drawString(36, y, ' '.join(latin))
                y -= 14
            ar_font = _ensure_ar_test_font()
            arabic = _re.sub(r'[A-Za-z0-9%._/-]+', ' ', line)
            if ar_font and arabic.strip():
                canv.setFont(ar_font, 10)
                canv.drawRightString(560, y, arabic.strip())
                y -= 14
            elif arabic.strip():
                canv.setFont('Helvetica', 10)
                canv.drawString(36, y, arabic.strip())
                y -= 14
        else:
            canv.setFont('Helvetica', 10)
            canv.drawString(36, y, line)
            y -= 14
    if extra_visible:
        canv.setFont('Helvetica', 10)
        canv.drawString(36, y, extra_visible)
        y -= 14
    if actual:
        _end_actual_text(canv)
    canv.save()
    return buf.getvalue()


class OrderedCompleteAndOverlayProvenanceTests(unittest.TestCase):
    """Causal ordered-content and file-provenance overlay contract."""

    def _helper(self, expected, painted, actual):
        from release_engine_v3.rel37_export_content_parity import (
            _paragraph_pdf_blockers,
            _visible_reconciled_to_actual,
        )
        vis_cmp, how = _visible_reconciled_to_actual(painted, actual)
        blockers = _paragraph_pdf_blockers(
            0, expected, painted, visible=painted, actual=actual)
        return vis_cmp, how, blockers

    def _assert_narrative_blocker(self, blockers):
        self.assertTrue(
            any(
                'actual_visible_disagree' in item
                or 'narrative_missing' in item
                or 'latin_missing' in item
                or 'latin_order' in item
                for item in blockers
            ),
            blockers,
        )
        self.assertFalse(any(
            'painted_unestablished' in item
            or item in (
                'pdf_extraction_unreliable',
                'pdf_environment_section_unassociated',
                'docx_bytes_missing',
                'pdf_bytes_missing',
            )
            for item in blockers
        ), blockers)

    def test_helper_rejects_unordered_and_negation_swaps(self):
        cases = (
            ('arabic_relation', ORDER_AR_EXP, ORDER_AR_VIS, ORDER_AR_EXP),
            ('added_negation', NEG_MIX_EXP, NEG_MIX_VIS, NEG_MIX_EXP),
            ('english_owner', OWNER_EXP, OWNER_VIS, OWNER_EXP),
            ('numeric_value', NUM_AR_EXP, NUM_AR_VIS, NUM_AR_EXP),
        )
        for name, expected, painted, actual in cases:
            vis, how, blockers = self._helper(expected, painted, actual)
            self.assertNotEqual(how, 'visible_visual_matches_actual', name)
            self.assertNotEqual(vis, expected, (name, vis, how))
            self._assert_narrative_blocker(blockers)

    def test_helper_controls_still_hold(self):
        vis, how, blockers = self._helper(WRAP_EN, WRAP_EN, WRAP_EN)
        self.assertEqual(how, 'same')
        self.assertFalse(blockers, (how, blockers, vis))
        wrapped = 'NDMO is the first reference;\nPDPL is the second reference.'
        vis, how, blockers = self._helper(WRAP_EN, wrapped, WRAP_EN)
        self.assertIn(how, ('same', 'visible_wraps_actual'))
        self.assertFalse(blockers, (how, blockers, vis))
        vis, how, blockers = self._helper(
            COMPLETE_EN, PAINTED_EN_DELETED, COMPLETE_EN)
        self._assert_narrative_blocker(blockers)
        vis, how, blockers = self._helper(
            COMPLETE_AR, PAINTED_AR_DELETED, COMPLETE_AR)
        self._assert_narrative_blocker(blockers)
        vis, how, blockers = self._helper(
            COMPLETE_EN, SPECIMEN_EN_CHANGED, SPECIMEN_EN_CHANGED)
        self.assertTrue(blockers, (how, blockers, vis))

    def test_numeric_overlay_classification_retains_painted(self):
        from release_engine_v3.rel37_export_content_parity import (
            _drop_overlay_visible,
            _is_visual_presentation_of_actual,
            _paragraph_pdf_blockers,
        )
        direct = _paragraph_pdf_blockers(
            0, NUM_AR_EXP, NUM_AR_VIS, visible=NUM_AR_VIS, actual=NUM_AR_EXP)
        self._assert_narrative_blocker(direct)
        visual_pres = _is_visual_presentation_of_actual(NUM_AR_VIS, NUM_AR_EXP)
        self.assertFalse(visual_pres)
        retained = _drop_overlay_visible([NUM_AR_VIS], NUM_AR_EXP)
        self.assertTrue(retained, retained)
        self.assertTrue(
            any(NUM_AR_VIS in item or '10%' in item for item in retained),
            retained)
        downstream = _paragraph_pdf_blockers(
            0, NUM_AR_EXP, '\n'.join(retained),
            visible='\n'.join(retained), actual=NUM_AR_EXP)
        self._assert_narrative_blocker(downstream)

    def test_actual_pdf_arabic_relation_swap_rejected(self):
        model = _EnvModel(ORDER_AR_EXP)
        good = _complete_repr_pdf(ORDER_AR_EXP, ORDER_AR_EXP)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(ORDER_AR_VIS, ORDER_AR_EXP)
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, mutated))

    def test_actual_pdf_added_negation_rejected(self):
        model = _EnvModel(NEG_MIX_EXP)
        good = _complete_repr_pdf(NEG_MIX_EXP, NEG_MIX_EXP)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(NEG_MIX_VIS, NEG_MIX_EXP)
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, mutated))

    def test_actual_pdf_english_owner_swap_rejected(self):
        model = _EnvModel(OWNER_EXP, lang='en', org='SpecimenOrg')
        good = _complete_repr_pdf(OWNER_EXP, OWNER_EXP)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(OWNER_VIS, OWNER_EXP)
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, mutated))

    def test_actual_pdf_numeric_100_to_10_retained_then_rejected(self):
        from release_engine_v3.rel37_export_content_parity import (
            pdf_environment_section_text,
        )
        model = _EnvModel(NUM_AR_EXP)
        good = _complete_repr_pdf(NUM_AR_EXP, NUM_AR_EXP)
        self.assertEqual(compare_environment_narrative_to_pdf(model, good), [])
        mutated = _complete_repr_pdf(NUM_AR_VIS, NUM_AR_EXP)
        section, meta = pdf_environment_section_text(mutated)
        visible = str(meta.get('environment_visible') or '')
        self.assertIn('10%', visible)
        self.assertFalse(any(
            item.get('reason', '').startswith('non_displayed')
            and '10%' in str(item.get('text') or '')
            for item in meta.get('overlay_excluded') or []
        ), meta.get('overlay_excluded'))
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, mutated))
        del section

    def test_hidden_logical_correct_visible_accepted(self):
        model = _EnvModel(OWNER_EXP, lang='en', org='SpecimenOrg')
        raw = _provenance_pdf(
            visible=OWNER_EXP, actual=OWNER_EXP,
            hidden_logical=OWNER_EXP, emit_hidden=True)
        self.assertEqual(compare_environment_narrative_to_pdf(model, raw), [])

    def test_hidden_logical_wrong_paint_rejected(self):
        model = _EnvModel(OWNER_EXP, lang='en', org='SpecimenOrg')
        raw = _provenance_pdf(
            visible=OWNER_VIS, actual=OWNER_EXP,
            hidden_logical=OWNER_EXP, emit_hidden=True)
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, raw))

    def test_correct_draw_plus_contradictory_visible_rejected(self):
        model = _EnvModel(OWNER_EXP, lang='en', org='SpecimenOrg')
        raw = _provenance_pdf(
            visible=OWNER_EXP, actual=OWNER_EXP, extra_visible=CONTRA_EXTRA)
        self._assert_narrative_blocker(
            compare_environment_narrative_to_pdf(model, raw))

    def test_valid_wrap_mixed_hidden_and_continuation(self):
        wrap_vis = 'NDMO is the first reference;\nPDPL is the second reference.'
        model = _EnvModel(WRAP_EN, lang='en', org='SpecimenOrg')
        wrapped = _complete_repr_pdf(wrap_vis, WRAP_EN)
        self.assertEqual(compare_environment_narrative_to_pdf(model, wrapped), [])
        mixed = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق '
            'NDMO وحماية بيانات شخصية وفق PDPL.')
        mixed_model = _EnvModel(mixed, lang='ar', org='جهة')
        mixed_pdf = _complete_repr_pdf(mixed, mixed)
        self.assertEqual(
            compare_environment_narrative_to_pdf(mixed_model, mixed_pdf), [])
        hidden_ok = _provenance_pdf(
            visible=mixed, actual=mixed, hidden_logical=mixed, emit_hidden=True)
        self.assertEqual(
            compare_environment_narrative_to_pdf(mixed_model, hidden_ok), [])
        first, second = EXPECTED_AR_ENV, (
            'تشمل المحركات التشغيلية اكتمال التصنيف وفق NDMO.')
        cont_model = _EnvModel(first + '\n\n' + second)
        continued = _pdf_owned(
            org_name='جهة',
            cover_sector='بنوك/مالي',
            env_heading=ENV_HEAD_EN,
            env_paras=[first],
            env_continue_paras=[second],
        )
        self.assertEqual(
            compare_environment_narrative_to_pdf(cont_model, continued), [])

    def test_no_borrowing_from_summary_cover_toc_appendix(self):
        model = _EnvModel(OWNER_EXP, lang='en', org='SpecimenOrg')
        borrowed = _specimen_env_pdf(
            env_visible=SPECIMEN_WRONG_ENV, env_actual=SPECIMEN_WRONG_ENV,
            summary_visible=OWNER_EXP, summary_actual=OWNER_EXP,
            same_page_summary=True)
        blockers = compare_environment_narrative_to_pdf(model, borrowed)
        self.assertTrue(blockers, blockers)
        cover_only = _pdf_owned(
            org_name='SpecimenOrg',
            cover_sector='Banking/Finance',
            env_heading=ENV_HEAD_EN,
            env_paras=[SPECIMEN_WRONG_ENV],
            extra_cover=OWNER_EXP,
        )
        self.assertTrue(compare_environment_narrative_to_pdf(model, cover_only))
        appendix_only = _pdf_owned(
            org_name='SpecimenOrg',
            cover_sector='Banking/Finance',
            env_heading=ENV_HEAD_EN,
            env_paras=[SPECIMEN_WRONG_ENV],
            appendix_paras=[OWNER_EXP],
        )
        self.assertTrue(
            compare_environment_narrative_to_pdf(model, appendix_only))

    def test_official_overlay_exclusions_are_file_proven(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        good = _export(saved, _official_body(saved), 'pdf')
        self.assertEqual(good['download_http'], 200, good['status'])
        self.assertTrue(good['bytes'].startswith(b'%PDF'))
        from release_engine_v3.rel37_export_content_parity import (
            extract_pdf_pages,
            pdf_environment_section_text,
        )
        pages, _meta = extract_pdf_pages(good['bytes'])
        section, detail = pdf_environment_section_text(good['bytes'])
        self.assertTrue(str(detail.get('environment_visible') or '').strip(), detail)
        self.assertTrue(str(detail.get('environment_actual') or '').strip(), detail)
        for item in detail.get('overlay_excluded') or []:
            self.assertTrue(
                str(item.get('reason') or '').startswith('non_displayed_render_mode_3'),
                item)
        self.assertEqual(
            compare_environment_narrative_to_pdf(model, good['bytes']), [])
        self.assertTrue(section.strip())
        self.assertTrue(pages)
        self.assertEqual(model.model_hash, LIVE_HASH)

    def test_real_worker_owner_pdf_and_conflicting_paint_refused(self):
        model = _load_json_model(LIVE_FIXTURE)
        self.assertEqual(model.model_hash, LIVE_HASH)
        saved = _persist(model, db_sector='Healthcare')
        good = _export(saved, _official_body(saved), 'pdf', immediate=False)
        self.assertEqual(good['download_http'], 200, good['status'])
        self.assertTrue(good['bytes'].startswith(b'%PDF'))
        self.assertEqual(
            compare_environment_narrative_to_pdf(model, good['bytes']), [])
        real_gate = app_mod._rel37_gate_saved_export_bytes

        def injecting_gate(*, docx_bytes=None, pdf_bytes=None, sections=None,
                           route='pdf', lang='ar'):
            if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
                pdf_bytes = _complete_repr_pdf(OWNER_VIS, OWNER_EXP)
            return real_gate(
                docx_bytes=docx_bytes, pdf_bytes=pdf_bytes,
                sections=sections, route=route, lang=lang)

        with patch.object(app_mod, '_rel37_gate_saved_export_bytes', injecting_gate):
            result = _export(
                saved, _official_body(saved), 'pdf', immediate=False)
        self.assertNotEqual(result['status'].get('status'), 'done', result['status'])
        self.assertNotEqual(result['download_http'], 200)
        self.assertFalse(result['bytes'].startswith(b'%PDF'))
        self.assertEqual(model.model_hash, LIVE_HASH)
        self.assertEqual(model.compute_model_hash(), LIVE_HASH)


if __name__ == '__main__':
    unittest.main()
