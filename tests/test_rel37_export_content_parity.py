"""REL37 full-document preview/DOCX/PDF content parity.

Uses the original Data EN saved model copy and compiled Data/AI/DT
fixtures. No provider calls. Exercises the real local export routes.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_ENV_KEYS = (
    'ADMIN_PASSWORD', 'SECRET_KEY', 'DATABASE_PATH', 'DATABASE_URL',
    'OPENAI_API_KEY', 'ANTHROPIC_API_KEY', 'GOOGLE_API_KEY',
    'REL2_SKIP_EXPORT_EVIDENCE', 'REL37_DATA_AI_DT_COMPILER',
)

_TMP = tempfile.mkdtemp(prefix='test_rel37_parity_')
os.environ['ADMIN_PASSWORD'] = 'test-admin-password'
os.environ['SECRET_KEY'] = 'test-secret-key'
os.environ['DATABASE_PATH'] = os.path.join(_TMP, 'rel37_parity.db')
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.join(_TMP, 'rel37_parity.db')
os.environ['OPENAI_API_KEY'] = ''
os.environ['ANTHROPIC_API_KEY'] = ''
os.environ['GOOGLE_API_KEY'] = ''
os.environ['REL37_DATA_AI_DT_COMPILER'] = '1'
# Do not set REL2_SKIP_EXPORT_EVIDENCE here. That flag disables REL2/cyber
# byte-evidence collection and DQS enforcement. This file's required
# content-parity gate is evaluate_export_parity on returned bytes.
# Capture restore state AFTER clearing so tearDown cannot re-enable skip.
os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
_ENV_BEFORE = {key: os.environ.get(key) for key in _ENV_KEYS}
_ENV_BEFORE['REL2_SKIP_EXPORT_EVIDENCE'] = None

import app as app_mod  # noqa: E402

from release_engine_v3.rel37_apply import (  # noqa: E402
    REL37_APPLIED_KEY,
    REL37_CANONICAL_FW_KEY,
    REL37_HASH_KEY,
    REL37_MODEL_KEY,
    REL37_ORIGINAL_FW_KEY,
    REL37_SELECTION_REASON_KEY,
    REL37_SELECTION_SUPPORTED_KEY,
    apply_rel37_to_sections,
    load_model,
    rel37_bind_export_sections,
    serialize_model,
)
from release_engine_v3.rel37_canonical_document import CanonicalDocument  # noqa: E402
from release_engine_v3.rel37_export_content_parity import (  # noqa: E402
    arabic_char_count,
    compare_model_to_docx,
    evaluate_export_parity,
    expected_rows,
    extract_pdf_text,
    inventory_docx_bytes,
    negative_control_arabic_cell_fails,
    negative_control_english_wording_fails,
)
from release_engine_v3.rel37_professional_projection import (  # noqa: E402
    apply_rel37_projection_to_blocks,
    load_validated_rel37_model,
)
from release_engine_v3.rel37_render import model_to_markdown, model_to_sections  # noqa: E402

FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37' / 'data_en_saved_canonical_model.json'
# Optional archival paths. Mandatory regressions must not read these.
ORIGINAL_DOCX = ROOT / 'qa_outputs' / 'rel37_07_staging' / 'live' / 'data_en' / 'export.docx'
ORIGINAL_MODEL = ROOT / 'qa_outputs' / 'rel37_07_staging' / 'live' / 'data_en' / 'saved_canonical_model.json'
EXPECTED_MODEL_HASH = 'd73047a541e3a640d2ef0246b58ab50465e1a5464554633220cb9bf36f0bbdab'
WRONG_REPLACEMENT = 'إطار حوكمة NDMO معتمد'
ON_TIME = (
    'On-time data-subject request closure',
    'On-time eligible-incident notification',
)

_UID = {'n': 200}


class ImmediateThread:
    def __init__(self, target=None, args=(), kwargs=None, daemon=None):
        self.target = target
        self.args = args or ()
        self.kwargs = kwargs or {}

    def start(self):
        self.target(*self.args, **self.kwargs)

    def join(self, timeout=None):
        return None


def _load_saved_model() -> CanonicalDocument:
    payload = json.loads(FIXTURE.read_text(encoding='utf-8'))
    model = CanonicalDocument.from_dict(payload)
    if not model.model_hash:
        model.compute_hashes()
    return model


def _docx_package_readable(raw: bytes) -> bool:
    if not raw.startswith(b'PK'):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            names = zf.namelist()
            if 'word/document.xml' not in names:
                return False
            zf.read('word/document.xml')
        inventory_docx_bytes(raw)
    except Exception:
        return False
    return True


def _docx_from_model(model: CanonicalDocument) -> bytes:
    """Test-owned readable DOCX built from the tracked fixture model."""
    from docx import Document
    doc = Document()
    doc.add_paragraph(model.vision or model.org_name or 'REL37 fixture')
    heading = (
        'البيئة التنظيمية والتهديدات' if model.lang == 'ar'
        else 'Business Environment and Drivers')
    doc.add_paragraph(heading)
    for para in str(model.environment_narrative or '').split('\n\n'):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.add_paragraph('Gap Analysis')
    for family, rows in expected_rows(model).items():
        if family.startswith('_') or not rows:
            continue
        cols = max(len(row) for row in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        for ridx, row in enumerate(rows):
            for cidx, value in enumerate(row):
                table.cell(ridx, cidx).text = str(value)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_from_model(model: CanonicalDocument) -> bytes:
    """Test-owned extractable PDF containing the model's substantive cells."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas
    buf = io.BytesIO()
    canv = Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 36
    canv.setFont('Helvetica', 8)
    canv.drawString(36, y, f'{model.org_name} {model.domain} {model.lang}')
    y -= 12
    canv.drawString(36, y, 'Sector —')
    canv.showPage()
    y = height - 36
    canv.setFont('Helvetica', 8)
    env_heading = (
        'البيئة التنظيمية والتهديدات' if model.lang == 'ar'
        else 'Business Environment and Drivers')
    canv.drawString(36, y, env_heading)
    y -= 12
    for para in str(model.environment_narrative or '').split('\n\n'):
        words = para.strip().split()
        if not words:
            continue
        line = ''
        for word in words:
            candidate = (line + ' ' + word).strip()
            if len(candidate) > 96 and line:
                if y < 48:
                    canv.showPage()
                    canv.setFont('Helvetica', 8)
                    y = height - 36
                canv.drawString(36, y, line)
                y -= 10
                line = word
            else:
                line = candidate
        if line:
            if y < 48:
                canv.showPage()
                canv.setFont('Helvetica', 8)
                y = height - 36
            canv.drawString(36, y, line)
            y -= 10
    for family, rows in expected_rows(model).items():
        if family.startswith('_'):
            continue
        for row in rows:
            line = ' | '.join(str(cell) for cell in row)
            if y < 48:
                canv.showPage()
                canv.setFont('Helvetica', 8)
                y = height - 36
            canv.drawString(36, y, line[:120])
            y -= 10
    canv.save()
    return buf.getvalue()


def _mutate_roadmap_deliverable(raw: bytes, model: CanonicalDocument, replacement: str) -> bytes:
    from docx import Document
    original = model.roadmap[0].deliverable
    doc = Document(io.BytesIO(raw))
    mutated = False
    target_row = None
    for table in doc.tables:
        for row in table.rows:
            texts = [cell.text for cell in row.cells]
            if original in texts:
                for cell in row.cells:
                    if cell.text == original:
                        cell.text = replacement
                        mutated = True
                        target_row = [c.text for c in row.cells]
                        break
            if mutated:
                break
        if mutated:
            break
    if not mutated:
        raise AssertionError('roadmap deliverable cell was not found for mutation')
    if original in (target_row or []) or replacement not in (target_row or []):
        raise AssertionError('mutation missed the intended deliverable cell')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def tearDownModule():
    for key, previous in _ENV_BEFORE.items():
        if previous is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = previous


def _sections_from_model(model: CanonicalDocument, original_fw=None) -> dict:
    sections = model_to_sections(model)
    sections[REL37_APPLIED_KEY] = '1'
    sections[REL37_MODEL_KEY] = serialize_model(model)
    sections[REL37_HASH_KEY] = model.model_hash
    sections[REL37_SELECTION_SUPPORTED_KEY] = 'true'
    sections[REL37_SELECTION_REASON_KEY] = 'supported_selection'
    sections[REL37_CANONICAL_FW_KEY] = list(model.selected_frameworks)
    sections[REL37_ORIGINAL_FW_KEY] = list(
        original_fw or model.selected_frameworks)
    return sections


def _next_user():
    _UID['n'] += 1
    uid = _UID['n']
    username = f'rel37parity{uid}'
    with app_mod.app.app_context():
        db = app_mod.get_db()
        db.execute(
            'INSERT OR IGNORE INTO users '
            '(id, username, password_hash, role, is_active) '
            'VALUES (?, ?, ?, ?, 1)',
            (uid, username, 'x', 'user'),
        )
        db.commit()
    return uid, username


def _client(uid, username):
    client = app_mod.app.test_client()
    csrf = f'rel37-parity-csrf-{uid}'
    with client.session_transaction() as sess:
        sess['user_id'] = uid
        sess['username'] = username
        sess['role'] = 'user'
        sess['csrf_token'] = csrf
    headers = {'X-CSRFToken': csrf, 'Content-Type': 'application/json'}
    return client, headers


def _persist_model(model: CanonicalDocument, domain_label: str, lang: str) -> dict:
    uid, username = _next_user()
    client, headers = _client(uid, username)
    sections = _sections_from_model(model)
    content = model_to_markdown(model)
    with app_mod.app.app_context():
        db = app_mod.get_db()
        cur = db.execute(
            'INSERT INTO strategies '
            '(user_id, domain, org_name, sector, content, language, '
            ' document_title, sections_json, content_json) '
            'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (
                uid, domain_label, model.org_name, 'Government',
                content, lang, f'{domain_label} strategy',
                json.dumps(sections, ensure_ascii=False),
                json.dumps({'sections': sections}, ensure_ascii=False),
            ),
        )
        sid = cur.lastrowid
        db.commit()
    return {
        'uid': uid,
        'username': username,
        'client': client,
        'headers': headers,
        'strategy_id': sid,
        'sections': sections,
        'content': content,
        'model': model,
        'domain': domain_label,
        'lang': lang,
    }


def _export(saved, fmt):
    body = {
        'content': saved['content'],
        'filename': f'rel37_parity_{fmt}',
        'language': saved['lang'],
        'domain': saved['domain'],
        'doc_type': 'Strategy Document',
        'document_type': 'strategy',
        'artifact_type': 'strategy',
        'generation_mode': 'drafting',
        'strategy_id': saved['strategy_id'],
        'artifact_id': saved['strategy_id'],
        'selected_frameworks': list(saved['model'].selected_frameworks),
        'frameworks': list(saved['model'].selected_frameworks),
        'org_name': saved['model'].org_name,
    }
    with patch('threading.Thread', ImmediateThread):
        resp = saved['client'].post(
            f'/api/generate-{fmt}-async', json=body, headers=saved['headers'])
    submit = resp.get_json(silent=True) or {}
    tid = submit.get('task_id')
    status = {}
    if tid:
        status = saved['client'].get(
            f'/api/export-status/{tid}', headers=saved['headers']
        ).get_json(silent=True) or {}
    raw = b''
    download_http = 0
    if tid and status.get('status') == 'done':
        dl = saved['client'].get(
            f'/api/export-download/{tid}', headers=saved['headers'])
        download_http = dl.status_code
        raw = dl.data or b''
    return {
        'submit_http': resp.status_code,
        'task_id': tid,
        'status': status,
        'download_http': download_http,
        'bytes': raw,
        'sha256': hashlib.sha256(raw).hexdigest() if raw else '',
    }


def _preview_text(saved) -> str:
    latest = saved['client'].get(
        f"/api/strategy/latest?domain={saved['domain']}"
        f"&lang={saved['lang']}&document_type=strategy"
        f"&strategy_id={saved['strategy_id']}",
        headers=saved['headers'],
    )
    body = latest.get_json(silent=True) or {}
    sections = body.get('sections') or {}
    if not isinstance(sections, dict):
        sections = {}
    return '\n\n'.join(str(sections.get(k) or '') for k in (
        'vision', 'pillars', 'environment', 'gaps', 'roadmap',
        'kpis', 'confidence', 'governance', 'traceability',
    ))


def _compile(domain, lang, org_name, frameworks):
    out, _repairs = apply_rel37_to_sections(
        {'vision': 'placeholder'},
        domain=domain,
        lang=lang,
        document_type='strategy',
        selected_frameworks=frameworks,
        org_name=org_name,
    )
    model = load_model(out)
    if model is None:
        raise AssertionError(f'compile failed for {domain} {lang}')
    return model, out


class BaselineReproductionTests(unittest.TestCase):
    def test_original_evidence_untouched(self):
        self.assertTrue(FIXTURE.exists())
        self.assertEqual(
            _load_saved_model().model_hash, EXPECTED_MODEL_HASH)
        if not ORIGINAL_DOCX.exists() and not ORIGINAL_MODEL.exists():
            # Archival check is not_run. Absence is not verification.
            return
        if ORIGINAL_MODEL.exists():
            self.assertEqual(
                hashlib.sha256(ORIGINAL_MODEL.read_bytes()).digest(),
                hashlib.sha256(FIXTURE.read_bytes()).digest())
        if ORIGINAL_DOCX.exists():
            self.assertEqual(
                hashlib.sha256(ORIGINAL_DOCX.read_bytes()).hexdigest(),
                'd145425f6cbdd45323c1440355df9e076d2ac3a6332d831a5439b847611e49bd')
            inv = inventory_docx_bytes(ORIGINAL_DOCX.read_bytes())
            self.assertIn(WRONG_REPLACEMENT, inv['text'])
            self.assertIn('Approved NDMO policy', inv['text'])
            self.assertTrue(any(
                WRONG_REPLACEMENT in ' | '.join(row)
                for table in inv['tables'] for row in table
            ))

    def test_legacy_fill_roadmap_still_replaces_without_authority(self):
        from professional_strategy_render import _fill_roadmap_row
        filled, _meta = _fill_roadmap_row(
            ['Phase 1: Establish (1-6 months)', '1-6 months',
             'NDMO governance program', 'CDO', 'Approved NDMO policy', 'NDMO'],
            lang='en', domain='data')
        self.assertIn(WRONG_REPLACEMENT, filled[4])

    def test_saved_model_hash_and_deliverable(self):
        model = _load_saved_model()
        self.assertEqual(model.model_hash, EXPECTED_MODEL_HASH)
        self.assertEqual(model.roadmap[0].deliverable, 'Approved NDMO policy')
        self.assertEqual(model.confidence[0].factor, 'Approved-framework completeness')
        self.assertEqual(arabic_char_count(model.generated_text_blob(), org_name=model.org_name), 0)


class ProjectionAuthorityTests(unittest.TestCase):
    def test_projection_keeps_saved_values(self):
        model = _load_saved_model()
        sections = _sections_from_model(model)
        loaded, blockers = load_validated_rel37_model(
            sections,
            domain='data',
            lang='en',
            document_type='strategy',
            org_name=model.org_name,
            selected_frameworks=list(model.selected_frameworks),
        )
        self.assertEqual(blockers, [])
        self.assertIsNotNone(loaded)
        blocks = apply_rel37_projection_to_blocks({}, loaded)
        road = blocks['roadmap']['tables'][0]['rows'][0]
        self.assertEqual(road[4], 'Approved NDMO policy')
        self.assertNotIn(WRONG_REPLACEMENT, json.dumps(blocks, ensure_ascii=False))
        self.assertEqual(
            blocks['confidence_risk_register']['tables'][0]['rows'][0][0],
            'Approved-framework completeness')
        self.assertEqual(loaded.model_hash, EXPECTED_MODEL_HASH)

    def test_professional_enrich_uses_projection(self):
        from professional_strategy_render import enrich_professional_blocks
        model = _load_saved_model()
        sections = _sections_from_model(model)
        rendered = enrich_professional_blocks(
            {'domain': 'data', 'lang': 'en', 'org_name': model.org_name,
             'selected_frameworks': list(model.selected_frameworks),
             'document_type': 'strategy', 'blocks': {},
             'strategy_id': '1'},
            sections,
            {'org_name': model.org_name, 'domain': 'data',
             '_rel37_source_sections': sections,
             'selected_frameworks': list(model.selected_frameworks),
             'strategy_id': '1'},
            'en',
        )
        blob = json.dumps(rendered.get('blocks') or {}, ensure_ascii=False)
        self.assertIn('Approved NDMO policy', blob)
        self.assertNotIn(WRONG_REPLACEMENT, blob)
        self.assertIn('Approved-framework completeness', blob)
        self.assertEqual(model.compute_model_hash(), EXPECTED_MODEL_HASH)

    def test_forged_flag_does_not_preserve(self):
        from professional_strategy_render import enrich_professional_blocks
        forged = {'_rel37_applied': 'true', 'roadmap': '| Phase | Period |\n'}
        rendered = enrich_professional_blocks(
            {'domain': 'cyber', 'lang': 'ar', 'document_type': 'strategy',
             'blocks': {}},
            forged,
            {'org_name': 'x', 'domain': 'cyber'},
            'ar',
        )
        self.assertTrue(rendered.get('blocks'))

    def test_hash_mismatch_blocks_bind_and_render(self):
        model = _load_saved_model()
        sections = _sections_from_model(model)
        sections[REL37_HASH_KEY] = '0' * 64
        bound = rel37_bind_export_sections(
            sections,
            {'roadmap': 'legacy fallback'},
            domain='data',
            lang='en',
            document_type='strategy',
            org_name=model.org_name,
            selected_frameworks=list(model.selected_frameworks),
        )
        self.assertIn('_rel37_render_blocked', bound)
        self.assertNotEqual(bound.get('roadmap'), 'legacy fallback')
        from professional_strategy_render import enrich_professional_blocks
        rendered = enrich_professional_blocks(
            {'domain': 'data', 'lang': 'en', 'org_name': model.org_name,
             'selected_frameworks': list(model.selected_frameworks),
             'document_type': 'strategy', 'blocks': {}},
            bound,
            {'org_name': model.org_name, 'domain': 'data',
             'selected_frameworks': list(model.selected_frameworks)},
            'en',
        )
        blob = json.dumps(rendered.get('blocks') or {}, ensure_ascii=False)
        self.assertTrue(rendered.get('rel37_authority_blocked'))
        self.assertNotIn(WRONG_REPLACEMENT, blob)
        self.assertNotIn('legacy fallback', blob)


class NegativeControlTests(unittest.TestCase):
    def test_unexpected_arabic_cell_fails(self):
        self.assertTrue(negative_control_arabic_cell_fails(_load_saved_model()))

    def test_english_wording_change_fails(self):
        self.assertTrue(negative_control_english_wording_fails(_load_saved_model()))

    def test_changed_confidence_weight_fails(self):
        model = _load_saved_model()
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        cells = list(model.confidence[0].cells())
        cells[1] = '99%'
        for idx, value in enumerate(cells):
            table.cell(0, idx).text = value
        buf = io.BytesIO()
        doc.save(buf)
        blockers = compare_model_to_docx(model, buf.getvalue())
        self.assertTrue(any('confidence' in item for item in blockers))

    def test_missing_roadmap_row_fails(self):
        model = _load_saved_model()
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=6)
        for idx, value in enumerate(model.roadmap[1].cells()):
            table.cell(0, idx).text = value
        buf = io.BytesIO()
        doc.save(buf)
        blockers = compare_model_to_docx(model, buf.getvalue())
        self.assertTrue(any('roadmap_row:0' in item for item in blockers))

    def test_missing_traceability_row_fails(self):
        model = _load_saved_model()
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=5)
        for idx, value in enumerate(model.traceability[1].cells()):
            table.cell(0, idx).text = value
        buf = io.BytesIO()
        doc.save(buf)
        blockers = compare_model_to_docx(model, buf.getvalue())
        self.assertTrue(any('traceability_row:0' in item for item in blockers))

    def test_wrong_column_fails(self):
        model = _load_saved_model()
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=6)
        cells = list(model.roadmap[0].cells())
        cells[3], cells[4] = cells[4], cells[3]
        for idx, value in enumerate(cells):
            table.cell(0, idx).text = value
        buf = io.BytesIO()
        doc.save(buf)
        blockers = compare_model_to_docx(model, buf.getvalue())
        self.assertTrue(any('row_field_mismatch' in item for item in blockers))

    def test_source_hash_does_not_override_divergent_docx(self):
        model = _load_saved_model()
        before_hash = model.model_hash
        preview = model_to_markdown(model)
        matching_docx = _docx_from_model(model)
        matching_pdf = _pdf_from_model(model)
        self.assertTrue(_docx_package_readable(matching_docx))
        self.assertTrue(matching_docx.startswith(b'PK'))
        self.assertTrue(matching_pdf.startswith(b'%PDF'))
        self.assertGreater(len(matching_docx), 64)
        positive = evaluate_export_parity(
            model=model,
            persist_ok=True,
            download_ok=True,
            parsed_ok=True,
            preview_text=preview,
            docx_bytes=matching_docx,
            pdf_bytes=matching_pdf,
            source_hash=before_hash,
        )
        self.assertEqual(positive.content_parity, 'passed', positive.blockers)
        self.assertEqual(positive.language_correctness, 'passed', positive.blockers)
        self.assertFalse(any(
            item.startswith('pdf_') or item.startswith('docx_bytes')
            for item in positive.blockers))
        self.assertEqual(model.compute_model_hash(), before_hash)

        mutated = _mutate_roadmap_deliverable(
            matching_docx, model, 'Approved NDMO framework charter')
        self.assertTrue(_docx_package_readable(mutated))
        inv = inventory_docx_bytes(mutated)
        self.assertIn('Approved NDMO framework charter', inv['text'])
        self.assertTrue(any(
            'Approved NDMO framework charter' in ' | '.join(row)
            and model.roadmap[0].initiative in ' | '.join(row)
            for table in inv['tables'] for row in table
        ))
        docx_only = compare_model_to_docx(model, mutated)
        self.assertTrue(
            any(item.startswith('docx_row_field_mismatch:roadmap:0')
                for item in docx_only),
            docx_only)

        result = evaluate_export_parity(
            model=model,
            persist_ok=True,
            download_ok=True,
            parsed_ok=True,
            preview_text=preview,
            docx_bytes=mutated,
            pdf_bytes=matching_pdf,
            source_hash=before_hash,
        )
        self.assertEqual(result.content_parity, 'failed', result.blockers)
        self.assertTrue(
            any(item.startswith('docx_row_field_mismatch:roadmap:0')
                for item in result.blockers),
            result.blockers)
        self.assertFalse(any(
            item in result.blockers for item in (
                'pdf_bytes_missing', 'pdf_extraction_unreliable',
                'docx_bytes_missing', 'download_or_parse_failed')),
            result.blockers)
        self.assertEqual(result.details.get('pdf_blockers'), [])
        self.assertEqual(result.details.get('preview_blockers'), [])
        self.assertTrue(result.details['source_hash_does_not_override_content'])
        self.assertTrue(result.details['source_hash_matches_model'])
        self.assertEqual(model.compute_model_hash(), before_hash)
        self.assertEqual(result.model_hash, before_hash)

    def test_malformed_docx_is_not_accepted(self):
        model = _load_saved_model()
        matching_pdf = _pdf_from_model(model)
        raw = b'PK'
        with self.assertRaises(zipfile.BadZipFile):
            inventory_docx_bytes(raw)
        raised = None
        result = None
        try:
            result = evaluate_export_parity(
                model=model,
                persist_ok=True,
                download_ok=True,
                parsed_ok=True,
                preview_text=model_to_markdown(model),
                docx_bytes=raw,
                pdf_bytes=matching_pdf,
                source_hash=model.model_hash,
            )
        except zipfile.BadZipFile as exc:
            raised = exc
        if raised is not None:
            self.assertIsInstance(raised, zipfile.BadZipFile)
            self.assertIsNone(result)
        else:
            self.assertIsNotNone(result)
            self.assertNotEqual(result.content_parity, 'passed')
            self.assertNotEqual(result.rendering_readability, 'passed')

    def test_process_does_not_set_rel2_evidence_skip(self):
        self.assertNotEqual(os.environ.get('REL2_SKIP_EXPORT_EVIDENCE'), '1')

    def test_pdf_zero_arabic_without_expected_text_fails(self):
        model = _load_saved_model()
        result = evaluate_export_parity(
            model=model,
            persist_ok=True,
            download_ok=True,
            parsed_ok=True,
            preview_text=model_to_markdown(model),
            docx_bytes=b'',
            pdf_bytes=b'%PDF-1.4 empty',
            source_hash=model.model_hash,
        )
        self.assertEqual(result.content_parity, 'failed')
        self.assertIn('pdf_extraction_unreliable', result.blockers)

    def test_arabic_org_name_is_permitted(self):
        self.assertEqual(arabic_char_count('Hello شركة مثال world', org_name='شركة مثال'), 0)
        self.assertGreater(arabic_char_count('Hello إطار حوكمة world', org_name='شركة مثال'), 0)

    def test_missing_required_canonical_content_blocks(self):
        model = _load_saved_model()
        payload = model.to_dict()
        payload['roadmap'] = []
        payload['model_hash'] = model.model_hash
        empty = CanonicalDocument.from_dict(payload)
        sections = _sections_from_model(empty)
        _loaded, blockers = load_validated_rel37_model(
            sections,
            domain='data',
            lang='en',
            document_type='strategy',
            org_name=empty.org_name,
            selected_frameworks=list(empty.selected_frameworks),
        )
        self.assertTrue(blockers)
        bound = rel37_bind_export_sections(
            sections,
            {'roadmap': 'legacy fallback'},
            domain='data',
            lang='en',
            document_type='strategy',
            org_name=empty.org_name,
            selected_frameworks=list(empty.selected_frameworks),
        )
        self.assertIn('_rel37_render_blocked', bound)
        self.assertNotEqual(bound.get('roadmap'), 'legacy fallback')


class RouteParityTests(unittest.TestCase):
    def test_data_en_saved_model_route_parity(self):
        model = _load_saved_model()
        before_hash = model.model_hash
        saved = _persist_model(model, 'Data Management', 'en')
        preview = _preview_text(saved)
        self.assertIn('Approved NDMO policy', preview)
        self.assertNotIn(WRONG_REPLACEMENT, preview)
        docx = _export(saved, 'docx')
        pdf = _export(saved, 'pdf')
        self.assertEqual(docx['status'].get('status'), 'done', docx['status'])
        self.assertEqual(pdf['status'].get('status'), 'done', pdf['status'])
        self.assertEqual(docx['download_http'], 200)
        self.assertEqual(pdf['download_http'], 200)
        inv = inventory_docx_bytes(docx['bytes'])
        self.assertIn('Approved NDMO policy', inv['text'])
        self.assertNotIn(WRONG_REPLACEMENT, inv['text'])
        self.assertIn('Approved-framework completeness', inv['text'])
        for name in ON_TIME:
            self.assertIn(name, inv['text'])
            self.assertIn('100%', inv['text'])
        pdf_text, pdf_meta = extract_pdf_text(pdf['bytes'])
        self.assertTrue(pdf_meta.get('reliable'), pdf_meta)
        self.assertIn('Approved NDMO policy', pdf_text)
        self.assertNotIn(WRONG_REPLACEMENT, pdf_text)
        result = evaluate_export_parity(
            model=model,
            persist_ok=True,
            download_ok=True,
            parsed_ok=True,
            preview_text=preview,
            docx_bytes=docx['bytes'],
            pdf_bytes=pdf['bytes'],
            source_hash=before_hash,
        )
        self.assertEqual(result.persist_success, 'passed')
        self.assertEqual(result.download_parse_success, 'passed')
        self.assertEqual(result.content_parity, 'passed', result.blockers)
        self.assertEqual(result.language_correctness, 'passed', result.blockers)
        self.assertEqual(result.rendering_readability, 'passed', result.blockers)
        self.assertEqual(model.compute_model_hash(), before_hash)
        self.assertNotEqual(docx['sha256'],
                            'd145425f6cbdd45323c1440355df9e076d2ac3a6332d831a5439b847611e49bd')

    def test_arabic_org_english_document(self):
        model, _sections = _compile(
            'Data Management', 'en', 'شركة مثال',
            ['PDPL (Personal Data Protection Law)',
             'NDMO Data Governance Framework'])
        self.assertEqual(model.org_name, 'شركة مثال')
        saved = _persist_model(model, 'Data Management', 'en')
        docx = _export(saved, 'docx')
        self.assertEqual(docx['download_http'], 200)
        inv = inventory_docx_bytes(docx['bytes'])
        self.assertIn('شركة مثال', inv['text'])
        self.assertEqual(arabic_char_count(inv['text'], org_name='شركة مثال'), 0)
        self.assertIn(model.roadmap[0].deliverable, inv['text'])


class SixRouteMatrixTests(unittest.TestCase):
    CASES = (
        ('Data Management', 'en', 'REL37 Data EN UI Org',
         ['PDPL (Personal Data Protection Law)',
          'NDMO Data Governance Framework']),
        ('Data Management', 'ar', 'منظمة بيانات عربي',
         ['نظام حماية البيانات الشخصية (PDPL)',
          'إطار حوكمة البيانات - مكتب إدارة البيانات الوطنية (NDMO)']),
        ('Artificial Intelligence', 'en', 'REL37 AI EN UI Org',
         ['SDAIA AI Ethics Principles']),
        ('Artificial Intelligence', 'ar', 'منظمة ذكاء اصطناعي',
         ['SDAIA AI Ethics Principles']),
        ('Digital Transformation', 'en', 'REL37 DT EN UI Org',
         ['DGA Digital Government Policy']),
        ('Digital Transformation', 'ar', 'منظمة تحول رقمي',
         ['DGA Digital Government Policy']),
    )

    def test_six_route_preview_docx_pdf_parity(self):
        results = []
        for domain, lang, org, frameworks in self.CASES:
            model, _sections = _compile(domain, lang, org, frameworks)
            before = model.model_hash
            saved = _persist_model(model, domain, lang)
            preview = _preview_text(saved)
            docx = _export(saved, 'docx')
            pdf = _export(saved, 'pdf')
            result = evaluate_export_parity(
                model=model,
                persist_ok=True,
                download_ok=docx['download_http'] == 200 and pdf['download_http'] == 200,
                parsed_ok=bool(docx['bytes']) and bool(pdf['bytes']),
                preview_text=preview,
                docx_bytes=docx['bytes'],
                pdf_bytes=pdf['bytes'],
                source_hash=before,
            )
            results.append({
                'domain': domain, 'lang': lang,
                'content_parity': result.content_parity,
                'language': result.language_correctness,
                'blockers': result.blockers[:8],
                'model_hash_unchanged': model.compute_model_hash() == before,
            })
            self.assertEqual(result.content_parity, 'passed', results[-1])
            self.assertEqual(result.language_correctness, 'passed', results[-1])
            self.assertTrue(results[-1]['model_hash_unchanged'])
            self.assertIn(model.roadmap[0].deliverable, inventory_docx_bytes(docx['bytes'])['text'])
            if lang == 'en':
                self.assertNotIn(WRONG_REPLACEMENT, inventory_docx_bytes(docx['bytes'])['text'])
        out = Path(_TMP) / 'six_route_matrix.json'
        out.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')


_LIVE_DATA_EN_AR_ORG_HASH = (
    'ebb6b54215c0ce4f506295cfa55c3d14a801a99ca982e55e535942b34519b135'
)


def _export_real_thread(saved, fmt, *, gate_wrapper=None):
    """Public async export on a real threading.Thread. No ImmediateThread."""
    import time
    captured = {}
    real_gate = app_mod._rel37_gate_saved_export_bytes

    def _recording_gate(**kwargs):
        blob = kwargs.get('pdf_bytes') or kwargs.get('docx_bytes') or b''
        if gate_wrapper is not None:
            kwargs = gate_wrapper(dict(kwargs))
            blob = kwargs.get('pdf_bytes') or kwargs.get('docx_bytes') or b''
        if blob:
            captured['sha256'] = hashlib.sha256(blob).hexdigest()
            captured['nbytes'] = len(blob)
        return real_gate(**kwargs)

    body = {
        'content': saved['content'],
        'filename': f'rel37_ar_org_{fmt}',
        'language': saved['lang'],
        'domain': saved['domain'],
        'doc_type': 'Strategy Document',
        'document_type': 'strategy',
        'artifact_type': 'strategy',
        'generation_mode': 'drafting',
        'strategy_id': saved['strategy_id'],
        'artifact_id': saved['strategy_id'],
        'selected_frameworks': list(saved['model'].selected_frameworks),
        'frameworks': list(saved['model'].selected_frameworks),
        'org_name': saved['model'].org_name,
    }
    app_mod._rel37_gate_saved_export_bytes = _recording_gate
    try:
        resp = saved['client'].post(
            f'/api/generate-{fmt}-async', json=body, headers=saved['headers'])
        submit = resp.get_json(silent=True) or {}
        tid = submit.get('task_id')
        status = {}
        deadline = time.monotonic() + 90
        while tid and time.monotonic() < deadline:
            status = saved['client'].get(
                f'/api/export-status/{tid}', headers=saved['headers']
            ).get_json(silent=True) or {}
            if status.get('status') in ('done', 'error', 'not_found'):
                break
            time.sleep(0.25)
        raw = b''
        download_http = 0
        if tid and status.get('status') == 'done':
            dl = saved['client'].get(
                f'/api/export-download/{tid}', headers=saved['headers'])
            download_http = dl.status_code
            raw = dl.data or b''
        elif tid:
            dl = saved['client'].get(
                f'/api/export-download/{tid}', headers=saved['headers'])
            download_http = dl.status_code
            raw = dl.data or b''
    finally:
        app_mod._rel37_gate_saved_export_bytes = real_gate
    return {
        'submit_http': resp.status_code,
        'task_id': tid,
        'status': status,
        'download_http': download_http,
        'bytes': raw if status.get('status') == 'done' else b'',
        'download_body': raw,
        'sha256': hashlib.sha256(raw).hexdigest() if raw and status.get('status') == 'done' else '',
        'candidate_sha256': captured.get('sha256') or '',
        'candidate_bytes': captured.get('nbytes') or 0,
    }


def _redact_body_arabic(pdf_bytes, *, replacement='', skip_pages=()):
    """Remove body-band Arabic glyph runs from a builder PDF.

    Header bands (y < 40) stay. ``skip_pages`` keeps the cover (page 0)
    so a later refusal is not explained by deleting the cover name.
    The redact box stops short of the next Latin line so NDMO/PDPL
    are not clipped.
    """
    import pymupdf
    doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    for index, page in enumerate(doc):
        if index in skip_pages:
            continue
        rects = []
        for block in page.get_text('rawdict').get('blocks') or []:
            for line in block.get('lines') or []:
                for span in line.get('spans') or []:
                    if 'Noto' not in str(span.get('font') or '') and (
                            'Arabic' not in str(span.get('font') or '')):
                        continue
                    chars = span.get('chars') or []
                    if not chars:
                        continue
                    y0 = min(ch['bbox'][1] for ch in chars)
                    if y0 < 40:
                        continue
                    rects.append(pymupdf.Rect(
                        min(ch['bbox'][0] for ch in chars) - 0.4,
                        y0,
                        max(ch['bbox'][2] for ch in chars) + 0.4,
                        y0 + 12,
                    ))
        for rect in rects:
            page.add_redact_annot(rect, text=replacement or None, fill=(1, 1, 1))
        if rects:
            page.apply_redactions(images=0)
            if replacement:
                # apply_redactions text is not reliable for Arabic; stamp Latin.
                if replacement.isascii():
                    page.insert_text(
                        (rects[0].x0, rects[0].y0 + 9),
                        replacement, fontsize=9, fontname='helv',
                    )
    return doc.tobytes()


class ArabicOrgEnglishPdfTests(unittest.TestCase):
    """Saved English Data/AI/DT PDFs must paint an Arabic org name in place.

    The compiled fixture is not the private live source. Its hash must not
    be reported as the live model hash.
    """

    def _saved(self, domain, org, frameworks):
        model, _sections = _compile(domain, 'en', org, frameworks)
        self.assertEqual(model.org_name, org)
        self.assertNotEqual(model.model_hash, _LIVE_DATA_EN_AR_ORG_HASH)
        before = model.model_hash
        saved = _persist_model(model, domain, 'en')
        self.assertEqual(saved['model'].model_hash, before)
        return saved, before

    def _assert_painted_org(self, pdf_bytes, org, paragraph):
        import pymupdf
        from release_engine_v3.rel37_export_content_parity import (
            _layout_norm,
            pdf_environment_section_text,
        )
        _section, meta = pdf_environment_section_text(pdf_bytes)
        visible = _layout_norm(meta.get('environment_visible') or '')
        self.assertIn(_layout_norm(paragraph), visible)
        self.assertIn(org, visible)
        doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
        found = False
        for page in doc:
            for block in page.get_text('rawdict').get('blocks') or []:
                for line in block.get('lines') or []:
                    spans = line.get('spans') or []
                    for span in spans:
                        chars = span.get('chars') or []
                        if not chars or chars[0]['bbox'][1] < 40:
                            continue
                        text = ''.join(ch.get('c') or '' for ch in chars)
                        if 'Noto' not in str(span.get('font') or ''):
                            continue
                        if not any(
                                '\u0600' <= ch <= '\u06FF'
                                or '\uFB50' <= ch <= '\uFEFF'
                                for ch in text):
                            continue
                        xs = [
                            ch['bbox'][0] for ch in chars
                            if (ch.get('c') or '').strip()
                        ]
                        if len(xs) < 2:
                            continue
                        # Shaped Arabic is drawn right-to-left: the first
                        # stored character sits to the right of the last.
                        self.assertGreater(xs[0], xs[-1])
                        found = True
                        break
                if found:
                    break
            if found:
                break
        self.assertTrue(found, 'Arabic org run was not painted with Noto')

    def test_data_en_arabic_org_real_thread_pdf(self):
        import pymupdf
        from release_engine_v3.rel37_export_content_parity import (
            _layout_norm,
            environment_narrative_paragraphs,
            gate_rel37_returned_bytes,
            pdf_environment_section_text,
        )
        org = 'شركة مثال'
        saved, before = self._saved(
            'Data Management', org,
            ['PDPL (Personal Data Protection Law)',
             'NDMO Data Governance Framework'])
        paragraph = environment_narrative_paragraphs(
            saved['model'].environment_narrative)[0]
        self.assertIn(org, paragraph)
        self.assertIn('operates', paragraph)
        markup = app_mod._english_pdf_mixed_script_markup(paragraph)
        self.assertIn('operates', markup)
        self.assertLess(markup.find('ArabicFont'), markup.find('operates'))
        self.assertNotIn(paragraph[::-1], markup)
        docx = _export_real_thread(saved, 'docx')
        self.assertEqual(docx['status'].get('status'), 'done', docx['status'])
        self.assertIn(org, inventory_docx_bytes(docx['bytes'])['text'])
        pdf = _export_real_thread(saved, 'pdf')
        self.assertEqual(pdf['submit_http'], 200, pdf['status'])
        self.assertEqual(pdf['status'].get('status'), 'done', pdf['status'])
        self.assertEqual(pdf['download_http'], 200)
        self.assertTrue(pdf['bytes'].startswith(b'%PDF'))
        self.assertEqual(pdf['candidate_sha256'], pdf['sha256'])
        self.assertEqual(
            hashlib.sha256(pdf['bytes']).hexdigest(), pdf['candidate_sha256'])
        allowed, blockers = gate_rel37_returned_bytes(
            saved['model'], pdf_bytes=pdf['bytes'], route='pdf')
        self.assertTrue(allowed, blockers)
        self._assert_painted_org(pdf['bytes'], org, paragraph)
        self.assertEqual(saved['model'].compute_model_hash(), before)
        self.assertEqual(
            arabic_char_count(
                extract_pdf_text(pdf['bytes'])[0], org_name=org), 0)

        removed = _redact_body_arabic(pdf['bytes'], skip_pages=(0,))
        ok, blockers = gate_rel37_returned_bytes(
            saved['model'], pdf_bytes=removed, route='pdf')
        self.assertFalse(ok)
        self.assertTrue(any(
            item.startswith('pdf_environment_narrative_missing:')
            for item in blockers), blockers)
        self.assertFalse(any(
            'unassociated' in item or 'csrf' in item or 'unreliable' in item
            for item in blockers), blockers)
        cover_doc = pymupdf.open(stream=removed, filetype='pdf')
        cover_text = ''.join(
            ch.get('c') or ''
            for block in cover_doc[0].get_text('rawdict').get('blocks') or []
            for line in block.get('lines') or []
            for span in line.get('spans') or []
            for ch in (span.get('chars') or [])
        )
        self.assertTrue(any(
            '\u0600' <= ch <= '\u06FF' or '\uFB50' <= ch <= '\uFEFF'
            for ch in cover_text))
        _section, removed_meta = pdf_environment_section_text(removed)
        self.assertNotIn(
            org, _layout_norm(removed_meta.get('environment_visible') or ''))

        replaced = _redact_body_arabic(
            pdf['bytes'], replacement='Other Org', skip_pages=(0,))
        ok, blockers = gate_rel37_returned_bytes(
            saved['model'], pdf_bytes=replaced, route='pdf')
        self.assertFalse(ok)
        self.assertTrue(any(
            item.startswith('pdf_environment_narrative_missing:')
            for item in blockers), blockers)

        def _swap(kwargs):
            kwargs['pdf_bytes'] = removed
            return kwargs

        refused = _export_real_thread(saved, 'pdf', gate_wrapper=_swap)
        self.assertEqual(refused['status'].get('status'), 'error', refused['status'])
        self.assertNotEqual(refused['status'].get('status'), 'done')
        self.assertFalse((refused['download_body'] or b'').startswith(b'%PDF'))
        self.assertIn(
            'validation failed',
            str(refused['status'].get('error') or ''))

    def test_org_name_variation_and_sibling_domains(self):
        from release_engine_v3.rel37_export_content_parity import (
            environment_narrative_paragraphs,
            gate_rel37_returned_bytes,
        )
        cases = (
            ('Data Management', 'مؤسسة الاختبار',
             ['PDPL (Personal Data Protection Law)',
              'NDMO Data Governance Framework']),
            ('Artificial Intelligence', 'شركة مثال',
             ['SDAIA AI Ethics Principles']),
            ('Digital Transformation', 'هيئة التجربة',
             ['DGA Digital Government Policy']),
        )
        for domain, org, frameworks in cases:
            saved, before = self._saved(domain, org, frameworks)
            paragraph = environment_narrative_paragraphs(
                saved['model'].environment_narrative)[0]
            self.assertIn(org, paragraph)
            pdf = _export_real_thread(saved, 'pdf')
            self.assertEqual(
                pdf['status'].get('status'), 'done', (domain, pdf['status']))
            self.assertEqual(pdf['candidate_sha256'], pdf['sha256'])
            allowed, blockers = gate_rel37_returned_bytes(
                saved['model'], pdf_bytes=pdf['bytes'], route='pdf')
            self.assertTrue(allowed, (domain, blockers))
            self._assert_painted_org(pdf['bytes'], org, paragraph)
            self.assertEqual(saved['model'].compute_model_hash(), before)
            self.assertEqual(
                arabic_char_count(
                    extract_pdf_text(pdf['bytes'])[0], org_name=org), 0)

    def test_actualtext_cannot_replace_missing_org_paint(self):
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.pagesizes import A4
        from release_engine_v3.rel37_export_content_parity import (
            environment_narrative_paragraphs,
            gate_rel37_returned_bytes,
        )
        from release_engine_v3.rel37_sector_context import (
            cover_sector_from_hashed_narrative,
        )
        org = 'شركة مثال'
        saved, _before = self._saved(
            'Data Management', org,
            ['PDPL (Personal Data Protection Law)',
             'NDMO Data Governance Framework'])
        model = saved['model']
        paragraph = environment_narrative_paragraphs(
            model.environment_narrative)[0]
        visible = paragraph.replace(org, '').strip()
        sector = cover_sector_from_hashed_narrative(
            model.environment_narrative, 'en')
        buf = io.BytesIO()
        canv = Canvas(buf, pagesize=A4)
        canv.setFont('Helvetica', 11)
        canv.drawString(48, 800, 'Organization')
        canv.drawString(48, 784, sector or 'Government')
        canv.showPage()
        canv.setFont('Helvetica', 12)
        canv.drawString(48, 800, 'Environment and Drivers')
        hx = 'FEFF' + paragraph.encode('utf-16-be').hex().upper()
        canv._code.append(f'/Span <</ActualText <{hx}>>> BDC')
        canv.setFont('Helvetica', 10)
        canv.drawString(48, 760, visible[:180])
        canv._code.append('EMC')
        canv.save()
        raw = buf.getvalue()
        allowed, blockers = gate_rel37_returned_bytes(
            model, pdf_bytes=raw, route='pdf')
        self.assertFalse(allowed)
        self.assertTrue(any(
            item.startswith('pdf_environment_narrative_missing:')
            for item in blockers), blockers)
        self.assertFalse(any(
            'unassociated' in item or item == 'pdf_extraction_unreliable'
            for item in blockers), blockers)

    def test_unrelated_arabic_outside_org_name_is_refused(self):
        import pymupdf
        from release_engine_v3.rel37_export_content_parity import (
            environment_narrative_paragraphs,
            gate_rel37_returned_bytes,
        )
        org = 'شركة مثال'
        saved, _before = self._saved(
            'Data Management', org,
            ['PDPL (Personal Data Protection Law)',
             'NDMO Data Governance Framework'])
        pdf = _export_real_thread(saved, 'pdf')
        self.assertEqual(pdf['status'].get('status'), 'done', pdf['status'])
        paragraph = environment_narrative_paragraphs(
            saved['model'].environment_narrative)[0]
        font = '/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf'
        if not Path(font).exists():
            font = str(ROOT / 'static/fonts/NotoSansArabic-Regular.ttf')
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.pagesizes import A4
        pdfmetrics.registerFont(TTFont('NARExtra', font))
        extra = io.BytesIO()
        canv = Canvas(extra, pagesize=A4)
        canv.setFont('NARExtra', 12)
        canv.drawString(48, 400, app_mod._shape_arabic_run('ملاحظة خارج الاسم'))
        canv.save()
        host = pymupdf.open(stream=pdf['bytes'], filetype='pdf')
        addon = pymupdf.open(stream=extra.getvalue(), filetype='pdf')
        host.insert_pdf(addon)
        mutated = host.tobytes()
        allowed, blockers = gate_rel37_returned_bytes(
            saved['model'], pdf_bytes=mutated, route='pdf')
        self.assertFalse(allowed)
        self.assertTrue(any(
            item.startswith('en_unexpected_arabic_chars:')
            for item in blockers), blockers)
        self.assertIn(org, paragraph)


if __name__ == '__main__':
    unittest.main()
