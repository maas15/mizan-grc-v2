"""Saved REL37 environment narrative and cover-sector authority.

Replays the live Data AR same-saved-source sequence from the retained
canonical fixture. Inspects actual returned DOCX/PDF bytes. Official
and saved-ID request shapes must not lose the hashed operating-context
clause or show a client Healthcare cover.
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_ENV_KEYS = (
    'ADMIN_PASSWORD', 'SECRET_KEY', 'DATABASE_PATH', 'DATABASE_URL',
    'OPENAI_API_KEY', 'ANTHROPIC_API_KEY', 'GOOGLE_API_KEY',
    'REL2_SKIP_EXPORT_EVIDENCE', 'REL37_DATA_AI_DT_COMPILER',
)
os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
_ENV_BEFORE = {key: os.environ.get(key) for key in _ENV_KEYS}

_TMP = tempfile.mkdtemp(prefix='test_rel37_narr_cover_')


def _ensure_test_env():
    os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
    os.environ['ADMIN_PASSWORD'] = 'test-admin-password'
    os.environ['SECRET_KEY'] = 'test-secret-key'
    os.environ.setdefault('DATABASE_PATH', os.path.join(_TMP, 'narr_cover.db'))
    os.environ.setdefault(
        'DATABASE_URL', 'sqlite:///' + os.path.join(_TMP, 'narr_cover.db'))
    os.environ['OPENAI_API_KEY'] = ''
    os.environ['ANTHROPIC_API_KEY'] = ''
    os.environ['GOOGLE_API_KEY'] = ''
    os.environ['REL37_DATA_AI_DT_COMPILER'] = '1'


_ensure_test_env()

import app as app_mod  # noqa: E402

from release_engine_v3.rel37_apply import (  # noqa: E402
    REL37_APPLIED_KEY,
    REL37_CANONICAL_FW_KEY,
    REL37_HASH_KEY,
    REL37_MODEL_KEY,
    REL37_ORIGINAL_FW_KEY,
    REL37_SELECTION_REASON_KEY,
    REL37_SELECTION_SUPPORTED_KEY,
    apply_rel37_to_sections,
    load_model,
    overlay_rel37_authority,
    serialize_model,
)
from release_engine_v3.rel37_canonical_document import CanonicalDocument  # noqa: E402
from release_engine_v3.rel37_export_content_parity import (  # noqa: E402
    compare_cover_sector_to_docx,
    compare_cover_sector_to_pdf,
    compare_environment_narrative_to_docx,
    compare_environment_narrative_to_pdf,
    compare_model_to_docx,
    docx_cover_sector_value,
    docx_environment_section_text,
    evaluate_export_parity,
    extract_pdf_text,
    gate_rel37_returned_bytes,
    inventory_docx_bytes,
    pdf_cover_and_body_text,
)
from release_engine_v3.rel37_professional_projection import (  # noqa: E402
    apply_rel37_projection_to_blocks,
)
from release_engine_v3.rel37_render import model_to_markdown, model_to_sections  # noqa: E402
from release_engine_v3.rel37_sector_context import (  # noqa: E402
    cover_sector_from_hashed_narrative,
)

LIVE_FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37' / (
    'data_ar_live_12966c11_canonical.json')
EXPECTED_HASH = (
    'c87b5cc6e1dee58de38657ee721b2e61c509a0e48141e8f6796414293eb99d40')
CLAUSE_AR = 'في سياق تشغيلي لقطاع بنوك/مالي'
GENERIC_AR = 'تعمل الجهة في بيئة تنظيمية'
_UID = {'n': 800}


class ImmediateThread:
    def __init__(self, target=None, args=(), kwargs=None, daemon=None):
        self.target = target
        self.args = args or ()
        self.kwargs = kwargs or {}

    def start(self):
        self.target(*self.args, **self.kwargs)

    def join(self, timeout=None):
        return None


def tearDownModule():
    for key, previous in _ENV_BEFORE.items():
        if previous is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = previous


def _load_live_model() -> CanonicalDocument:
    payload = json.loads(LIVE_FIXTURE.read_text(encoding='utf-8'))
    model = CanonicalDocument.from_dict(payload)
    if not model.model_hash:
        model.compute_hashes()
    return model


def _sections_from_model(model: CanonicalDocument) -> dict:
    sections = model_to_sections(model)
    sections[REL37_APPLIED_KEY] = '1'
    sections[REL37_MODEL_KEY] = serialize_model(model)
    sections[REL37_HASH_KEY] = model.model_hash
    sections[REL37_SELECTION_SUPPORTED_KEY] = 'true'
    sections[REL37_SELECTION_REASON_KEY] = 'supported_selection'
    sections[REL37_CANONICAL_FW_KEY] = list(model.selected_frameworks)
    sections[REL37_ORIGINAL_FW_KEY] = list(model.selected_frameworks)
    return sections


def _next_user():
    _UID['n'] += 1
    uid = _UID['n']
    username = f'rel37narr{uid}'
    with app_mod.app.app_context():
        db = app_mod.get_db()
        db.execute(
            'INSERT OR IGNORE INTO users '
            '(id, username, password_hash, role, is_active) '
            'VALUES (?, ?, ?, ?, 1)',
            (uid, username, 'x', 'user'),
        )
        db.commit()
    return uid, username


def _client(uid, username):
    client = app_mod.app.test_client()
    csrf = f'rel37-narr-csrf-{uid}'
    with client.session_transaction() as sess:
        sess['user_id'] = uid
        sess['username'] = username
        sess['role'] = 'user'
        sess['csrf_token'] = csrf
    return client, {
        'X-CSRFToken': csrf,
        'Content-Type': 'application/json',
    }


def _persist(model: CanonicalDocument, *, db_sector='Healthcare') -> dict:
    uid, username = _next_user()
    client, headers = _client(uid, username)
    sections = _sections_from_model(model)
    content = model_to_markdown(model)
    with app_mod.app.app_context():
        db = app_mod.get_db()
        cur = db.execute(
            'INSERT INTO strategies '
            '(user_id, domain, org_name, sector, content, language, '
            ' document_title, sections_json, content_json) '
            'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (
                uid, 'Data Management', model.org_name, db_sector,
                content, model.lang, 'REL37 saved narrative',
                json.dumps(sections, ensure_ascii=False),
                json.dumps({'sections': sections}, ensure_ascii=False),
            ),
        )
        sid = cur.lastrowid
        db.commit()
    return {
        'uid': uid,
        'client': client,
        'headers': headers,
        'strategy_id': sid,
        'model': model,
        'content': content,
        'sections': sections,
        'lang': model.lang,
        'domain': 'Data Management',
    }


def _export(saved, body, fmt):
    payload = dict(body)
    payload.setdefault('document_type', 'strategy')
    payload.setdefault('artifact_type', 'strategy')
    payload.setdefault('generation_mode', 'drafting')
    payload['strategy_id'] = saved['strategy_id']
    payload['artifact_id'] = saved['strategy_id']
    with patch('threading.Thread', ImmediateThread):
        resp = saved['client'].post(
            f'/api/generate-{fmt}-async',
            json=payload,
            headers=saved['headers'],
        )
    submit = resp.get_json(silent=True) or {}
    tid = submit.get('task_id')
    status = {}
    raw = b''
    download_http = 0
    if tid:
        status = saved['client'].get(
            f'/api/export-status/{tid}', headers=saved['headers']
        ).get_json(silent=True) or {}
        if status.get('status') == 'done':
            dl = saved['client'].get(
                f'/api/export-download/{tid}', headers=saved['headers'])
            download_http = dl.status_code
            raw = dl.data or b''
    return {
        'submit_http': resp.status_code,
        'submit': submit,
        'status': status,
        'download_http': download_http,
        'bytes': raw,
        'task_id': tid,
    }


def _browser_body(saved):
    return {
        'content': saved['content'],
        'filename': 'data_ar_rel37_23',
        'language': 'ar',
        'org_name': saved['model'].org_name,
        'sector': 'بنوك/مالي',
        'doc_type': 'Strategy Document',
        'domain': 'Data Management',
        'selected_frameworks': list(saved['model'].selected_frameworks),
    }


def _official_body(saved):
    return {
        'content': saved['content'],
        'filename': 'rel33_data_strategy',
        'language': 'ar',
        'org_name': 'REL33 P1 Data Management Org',
        'sector': 'Healthcare',
        'doc_type': 'Strategy Document',
        'domain': 'Data Management',
        'selected_frameworks': ['NDMO', 'PDPL'],
    }


def _saved_id_body(_saved):
    return {
        'language': 'ar',
        'filename': 'saved_id_only',
        'domain': 'Data Management',
    }


def _assert_docx_narrative_cover(test, model, raw, *, label):
    test.assertTrue(raw.startswith(b'PK'), label)
    env = docx_environment_section_text(raw)
    test.assertIn(CLAUSE_AR, env, label)
    test.assertNotIn(GENERIC_AR, env, label)
    test.assertEqual(compare_environment_narrative_to_docx(model, raw), [], label)
    test.assertEqual(compare_cover_sector_to_docx(model, raw), [], label)
    cover = docx_cover_sector_value(raw)
    test.assertIn(cover, ('بنوك/مالي', 'Banking/Finance'), (label, cover))
    test.assertNotEqual(cover, 'Healthcare', label)
    test.assertEqual(compare_model_to_docx(model, raw), [], label)


def _assert_pdf_narrative_cover(test, model, raw, *, label):
    test.assertTrue(raw.startswith(b'%PDF'), label)
    test.assertEqual(compare_environment_narrative_to_pdf(model, raw), [], label)
    test.assertEqual(compare_cover_sector_to_pdf(model, raw), [], label)
    cover, body, meta = pdf_cover_and_body_text(raw)
    test.assertTrue(meta.get('pages'), label)
    test.assertIn(CLAUSE_AR, body, label)
    test.assertNotIn('Healthcare', cover, label)


def _mutate_docx_environment(raw: bytes, replacement: str) -> bytes:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    found = False
    taking = False
    for para in doc.paragraphs:
        text = para.text or ''
        if 'البيئة التنظيمية والتهديدات' in text and len(text) < 80:
            if text[:1].isdigit():
                continue
            taking = True
            continue
        if taking and ('تحليل الفجوات' in text or 'Gap Analysis' in text):
            break
        if taking and CLAUSE_AR in text:
            para.text = replacement
            found = True
            break
        if taking and text.strip() and replacement == GENERIC_AR:
            para.text = GENERIC_AR
            found = True
            break
    if not found:
        raise AssertionError('environment paragraph was not found for mutation')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _mutate_docx_cover_sector(raw: bytes, replacement: str) -> bytes:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    found = False
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() in ('القطاع', 'Sector'):
                row.cells[1].text = replacement
                found = True
                break
        if found:
            break
    if not found:
        raise AssertionError('cover sector cell was not found')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class OverlayProjectionTests(unittest.TestCase):
    def test_leftover_generic_environment_does_not_outrank_saved(self):
        model = _load_live_model()
        saved = _sections_from_model(model)
        leftover = dict(saved)
        leftover['environment'] = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق NDMO.')
        overlaid = overlay_rel37_authority(
            leftover, saved, domain='data', lang='ar',
            org_name=model.org_name)
        restored = load_model(overlaid)
        self.assertIsNotNone(restored)
        self.assertIn(CLAUSE_AR, overlaid.get('environment', ''))
        self.assertIn(CLAUSE_AR, restored.environment_narrative)
        self.assertEqual(restored.model_hash, EXPECTED_HASH)

    def test_projection_writes_saved_narrative_into_environment_block(self):
        model = _load_live_model()
        blocks = apply_rel37_projection_to_blocks(
            {'environment_context': {
                'paragraphs': [GENERIC_AR],
                'content': GENERIC_AR,
            }},
            model,
        )
        paras = blocks['environment_context']['paragraphs']
        joined = '\n'.join(paras)
        self.assertIn(CLAUSE_AR, joined)
        self.assertNotIn(GENERIC_AR, joined)
        self.assertEqual(model.model_hash, EXPECTED_HASH)

    def test_frozen_renderer_bind_restores_saved_environment(self):
        from types import SimpleNamespace
        from release_engine_v3.rel32_docx_renderer import (
            bind_rel32_docx_renderer_input,
        )
        model = _load_live_model()
        saved = _sections_from_model(model)
        leftover_visible = {
            key: value for key, value in saved.items()
            if isinstance(value, str) and not str(key).startswith('_')
        }
        leftover_visible['environment'] = (
            'تعمل الجهة في بيئة تنظيمية تتطلب حوكمة بيانات وطنية وفق NDMO.')
        frozen = SimpleNamespace(
            legacy_sections=leftover_visible,
            canonical_sections={},
            document_type='strategy',
            artifact_type='strategy',
        )
        tree = SimpleNamespace(
            markdown_view=leftover_visible['environment'],
            nodes=[],
        )
        _content, sections, _meta = bind_rel32_docx_renderer_input(
            frozen, tree,
            backend={'_rel37_source_sections': saved},
            artifact_dict={'sections': leftover_visible},
        )
        self.assertIn(CLAUSE_AR, sections.get('environment', ''))
        self.assertNotIn(GENERIC_AR, sections.get('environment', ''))

    def test_raw_sector_provenance_cannot_change_cover_binding(self):
        model = _load_live_model()
        before = model.model_hash
        tampered = deepcopy(model)
        tampered.sector = 'Healthcare'
        self.assertEqual(tampered.compute_model_hash(), before)
        self.assertEqual(
            cover_sector_from_hashed_narrative(
                tampered.environment_narrative, 'ar'),
            'بنوك/مالي')
        self.assertNotEqual(
            cover_sector_from_hashed_narrative(
                tampered.environment_narrative, 'ar'),
            tampered.sector)


class SameSourceRouteSequenceTests(unittest.TestCase):
    def setUp(self):
        _ensure_test_env()

    def test_live_fixture_identity(self):
        model = _load_live_model()
        self.assertEqual(model.model_hash, EXPECTED_HASH)
        self.assertIn(CLAUSE_AR, model.environment_narrative)
        self.assertEqual(model.org_name, 'منظمة بيانات نسبي ٣٧٫٢٠')
        self.assertEqual(model.lang, 'ar')
        self.assertEqual(model.domain, 'data')

    def test_browser_official_saved_id_repeat_sequence(self):
        model = _load_live_model()
        before = model.model_hash
        saved = _persist(model, db_sector='Healthcare')
        order = (
            ('browser_docx', _browser_body(saved), 'docx'),
            ('browser_pdf', _browser_body(saved), 'pdf'),
            ('official_docx', _official_body(saved), 'docx'),
            ('official_pdf', _official_body(saved), 'pdf'),
            ('saved_id_docx', _saved_id_body(saved), 'docx'),
            ('repeat_official_docx', _official_body(saved), 'docx'),
            ('repeat_ui_pdf', _browser_body(saved), 'pdf'),
        )
        results = {}
        for name, body, fmt in order:
            results[name] = _export(saved, body, fmt)
            self.assertEqual(results[name]['submit_http'], 200, name)
            self.assertEqual(
                results[name]['status'].get('status'), 'done',
                (name, results[name]['status']))
            self.assertEqual(results[name]['download_http'], 200, name)
            if fmt == 'docx':
                _assert_docx_narrative_cover(
                    self, model, results[name]['bytes'], label=name)
            else:
                _assert_pdf_narrative_cover(
                    self, model, results[name]['bytes'], label=name)
        self.assertEqual(model.compute_model_hash(), before)
        self.assertEqual(model.model_hash, EXPECTED_HASH)

    def test_english_equivalent_cover_and_narrative(self):
        compiled = load_model(apply_rel37_to_sections(
            {'vision': 'placeholder'},
            domain='data',
            lang='en',
            document_type='strategy',
            selected_frameworks=['ndmo', 'pdpl'],
            org_name='Data Narrative EN Org',
            request={'sector': 'Banking/Finance'},
        )[0])
        self.assertIsNotNone(compiled)
        self.assertIn('Banking/Finance', compiled.environment_narrative)
        saved = _persist(compiled, db_sector='Healthcare')
        official = {
            'content': saved['content'],
            'filename': 'rel37_en_official',
            'language': 'en',
            'org_name': 'REL33 P1 Data Management Org',
            'sector': 'Healthcare',
            'doc_type': 'Strategy Document',
            'domain': 'Data Management',
            'selected_frameworks': ['NDMO', 'PDPL'],
        }
        docx = _export(saved, official, 'docx')
        pdf = _export(saved, official, 'pdf')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        self.assertEqual(pdf['download_http'], 200, pdf['status'])
        env = docx_environment_section_text(docx['bytes'])
        self.assertIn('Banking/Finance sector operating context', env)
        self.assertEqual(docx_cover_sector_value(docx['bytes']), 'Banking/Finance')
        cover, body, _meta = pdf_cover_and_body_text(pdf['bytes'])
        self.assertIn('Banking/Finance', body)
        self.assertNotIn('Healthcare', cover)
        self.assertNotRegex(env, r'[\u0600-\u06FF]')


class CausalNegativeTests(unittest.TestCase):
    def setUp(self):
        _ensure_test_env()
        self.model = _load_live_model()
        self.saved = _persist(self.model, db_sector='Healthcare')
        self.good = _export(self.saved, _official_body(self.saved), 'docx')
        self.assertEqual(self.good['download_http'], 200, self.good['status'])
        self.good_raw = self.good['bytes']

    def test_remove_only_operating_context_clause(self):
        mutated = _mutate_docx_environment(
            self.good_raw,
            self.model.environment_narrative.replace(CLAUSE_AR, 'ضمن').strip())
        blockers = compare_environment_narrative_to_docx(self.model, mutated)
        self.assertTrue(blockers, blockers)
        self.assertTrue(
            any('narrative' in item or 'generic' in item for item in blockers),
            blockers)
        allowed, gated = gate_rel37_returned_bytes(
            self.model, docx_bytes=mutated, route='docx')
        self.assertFalse(allowed)
        self.assertTrue(gated)
        self.assertFalse(any(
            item in gated for item in (
                'docx_bytes_missing', 'pdf_bytes_missing',
                'pdf_extraction_unreliable')), gated)
        self.assertEqual(self.model.model_hash, EXPECTED_HASH)

    def test_generic_leftover_replaces_environment(self):
        mutated = _mutate_docx_environment(self.good_raw, GENERIC_AR)
        env = docx_environment_section_text(mutated)
        self.assertIn(GENERIC_AR, env)
        self.assertNotIn(CLAUSE_AR, env)
        blockers = compare_environment_narrative_to_docx(self.model, mutated)
        self.assertTrue(any('narrative' in item or 'generic' in item
                            for item in blockers), blockers)

    def test_clause_only_outside_environment_section_fails(self):
        mutated = _mutate_docx_environment(self.good_raw, GENERIC_AR)
        from docx import Document
        doc = Document(io.BytesIO(mutated))
        doc.add_paragraph(CLAUSE_AR)
        buf = io.BytesIO()
        doc.save(buf)
        mutated = buf.getvalue()
        inv = inventory_docx_bytes(mutated)
        self.assertIn(CLAUSE_AR, inv['text'])
        env = docx_environment_section_text(mutated)
        self.assertNotIn(CLAUSE_AR, env)
        blockers = compare_environment_narrative_to_docx(self.model, mutated)
        self.assertTrue(blockers, blockers)

    def test_cover_only_healthcare_fails(self):
        mutated = _mutate_docx_cover_sector(self.good_raw, 'Healthcare')
        env = docx_environment_section_text(mutated)
        self.assertIn(CLAUSE_AR, env)
        blockers = compare_cover_sector_to_docx(self.model, mutated)
        self.assertTrue(any('cover_sector' in item for item in blockers), blockers)
        allowed, gated = gate_rel37_returned_bytes(
            self.model, docx_bytes=mutated, route='docx')
        self.assertFalse(allowed)
        self.assertTrue(any('cover_sector' in item for item in gated), gated)

    def test_matching_hash_does_not_excuse_mutated_bytes(self):
        mutated = _mutate_docx_cover_sector(self.good_raw, 'Healthcare')
        result = evaluate_export_parity(
            model=self.model,
            persist_ok=True,
            download_ok=True,
            parsed_ok=True,
            preview_text=self.saved['content'],
            docx_bytes=mutated,
            pdf_bytes=b'',
            source_hash=EXPECTED_HASH,
        )
        self.assertEqual(result.content_parity, 'failed', result.blockers)
        self.assertTrue(any('cover_sector' in item for item in result.blockers))
        self.assertEqual(result.model_hash, EXPECTED_HASH)
        self.assertEqual(self.model.compute_model_hash(), EXPECTED_HASH)

    def test_docx_pass_does_not_substitute_for_pdf(self):
        allowed, blockers = gate_rel37_returned_bytes(
            self.model, docx_bytes=self.good_raw, route='pdf')
        self.assertFalse(allowed)
        self.assertIn('pdf_bytes_missing', blockers)

    def test_server_path_blocks_wrong_cover_before_download(self):
        mutated = _mutate_docx_cover_sector(self.good_raw, 'Healthcare')
        allowed, blockers, gate = app_mod._rel37_gate_saved_export_bytes(
            docx_bytes=mutated,
            sections=self.saved['sections'],
            route='docx',
            lang='ar',
        )
        self.assertFalse(allowed)
        self.assertTrue(blockers)
        self.assertEqual(gate.get('model_hash'), EXPECTED_HASH)


class PreviewReloadTests(unittest.TestCase):
    def test_preview_keeps_saved_narrative(self):
        model = _load_live_model()
        saved = _persist(model, db_sector='Healthcare')
        self.assertIn(CLAUSE_AR, saved['content'])
        with app_mod.app.app_context():
            db = app_mod.get_db()
            row = db.execute(
                'SELECT content, sections_json FROM strategies '
                'WHERE id = ? AND user_id = ?',
                (saved['strategy_id'], saved['uid']),
            ).fetchone()
        self.assertIsNotNone(row)
        self.assertIn(CLAUSE_AR, row['content'] or '')
        self.assertIn(CLAUSE_AR, row['sections_json'] or '')
        self.assertEqual(model.model_hash, EXPECTED_HASH)


class OfficialDifferentialHealthcareTests(unittest.TestCase):
    def test_official_shaped_en_without_ui_pair_does_not_inject_healthcare(self):
        out, _repairs = apply_rel37_to_sections(
            {'vision': 'placeholder'},
            domain='data',
            lang='en',
            document_type='strategy',
            selected_frameworks=['ndmo', 'pdpl'],
            org_name='REL33 P1 Data Management Org',
        )
        model = load_model(out)
        self.assertIsNotNone(model)
        saved = _persist(model, db_sector='Healthcare')
        official = {
            'content': saved['content'],
            'filename': 'rel37_official_en',
            'language': 'en',
            'org_name': 'client-other-org',
            'sector': 'Healthcare',
            'doc_type': 'Strategy Document',
            'domain': 'Data Management',
            'selected_frameworks': list(model.selected_frameworks),
        }
        docx = _export(saved, official, 'docx')
        self.assertEqual(docx['download_http'], 200, docx['status'])
        cover = docx_cover_sector_value(docx['bytes'])
        self.assertNotEqual(cover, 'Healthcare')
        self.assertIn(cover, ('—', '-', '', 'Government', 'حكومي'))


if __name__ == '__main__':
    unittest.main()
