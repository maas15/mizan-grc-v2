"""Repeated Arabic Data traceability association on a compiled fixture.

Not official staging strategy 11. New implementation evidence only.
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

_ENV_KEYS = (
    'ADMIN_PASSWORD', 'SECRET_KEY', 'DATABASE_PATH', 'DATABASE_URL',
    'OPENAI_API_KEY', 'ANTHROPIC_API_KEY', 'GOOGLE_API_KEY',
    'REL2_SKIP_EXPORT_EVIDENCE', 'REL37_DATA_AI_DT_COMPILER',
)
os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
_ENV_BEFORE = {key: os.environ.get(key) for key in _ENV_KEYS}

_TMP = tempfile.mkdtemp(prefix='test_rel37_ar_cell_')


def _ensure_test_env():
    os.environ.pop('REL2_SKIP_EXPORT_EVIDENCE', None)
    os.environ['ADMIN_PASSWORD'] = 'test-admin-password'
    os.environ['SECRET_KEY'] = 'test-secret-key'
    os.environ.setdefault('DATABASE_PATH', os.path.join(_TMP, 'ar_cell.db'))
    os.environ.setdefault('DATABASE_URL', 'sqlite:///' + os.path.join(_TMP, 'ar_cell.db'))
    os.environ['OPENAI_API_KEY'] = ''
    os.environ['ANTHROPIC_API_KEY'] = ''
    os.environ['GOOGLE_API_KEY'] = ''
    os.environ['REL37_DATA_AI_DT_COMPILER'] = '1'


_ensure_test_env()

import app as app_mod  # noqa: E402

from release_engine_v3.rel32_docx_renderer import (  # noqa: E402
    sections_from_frozen_artifact,
)
from release_engine_v3.rel37_apply import (  # noqa: E402
    REL37_APPLIED_KEY,
    REL37_CANONICAL_FW_KEY,
    REL37_HASH_KEY,
    REL37_MODEL_KEY,
    REL37_ORIGINAL_FW_KEY,
    REL37_SELECTION_REASON_KEY,
    REL37_SELECTION_SUPPORTED_KEY,
    apply_rel37_to_sections,
    is_rel37_authoritative,
    load_model,
    overlay_rel37_authority,
    rel37_hash_identity_blockers,
    serialize_model,
)
from release_engine_v3.rel37_export_content_parity import (  # noqa: E402
    _paragraph_pdf_blockers,
    arabic_char_count,
    compare_model_to_docx,
    compare_traceability_association,
    expected_rows,
    extract_pdf_text,
    inventory_docx_bytes,
)
from release_engine_v3.rel37_render import model_to_markdown, model_to_sections  # noqa: E402
from tests._export_failure_diagnostics import (  # noqa: E402
    StdoutCapture,
    build_export_diagnostic,
    persist_export_diagnostic,
)

_DEJAVU_FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37_clean_runner_dejavu_env.json'
_AMIRI_DT_FIXTURE = ROOT / 'tests' / 'fixtures' / 'rel37_clean_runner_amiri_dt_env.json'
_REMOTE_FIRST_NODE_HASH = (
    '97a4831fe99430c4bba90e7c4c38064aee20b559af7363298f3ad496000df5b3'
)
_DT_AR_HASH = (
    '190e32b6d70b560cb6b7dc38562a4419ce79fd717d6c3a59b72ae0af33c108ca'
)


class ImmediateThread:
    def __init__(self, target=None, args=(), kwargs=None, daemon=None):
        self.target = target
        self.args = args or ()
        self.kwargs = kwargs or {}

    def start(self):
        self.target(*self.args, **self.kwargs)

    def join(self, timeout=None):
        return None


def tearDownModule():
    for key in ('REL2_SKIP_EXPORT_EVIDENCE', 'OPENAI_API_KEY',
                'ANTHROPIC_API_KEY', 'GOOGLE_API_KEY'):
        previous = _ENV_BEFORE.get(key)
        if previous is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = previous


def _compile_ar():
    out, _repairs = apply_rel37_to_sections(
        {'vision': 'placeholder'},
        domain='data',
        lang='ar',
        document_type='strategy',
        selected_frameworks=['ndmo', 'pdpl'],
        org_name='منظمة إدارة البيانات',
    )
    model = load_model(out)
    if model is None:
        raise AssertionError('arabic compile failed')
    return model, out


def _sections(model):
    sections = model_to_sections(model)
    sections[REL37_APPLIED_KEY] = '1'
    sections[REL37_MODEL_KEY] = serialize_model(model)
    sections[REL37_HASH_KEY] = model.model_hash
    sections[REL37_SELECTION_SUPPORTED_KEY] = 'true'
    sections[REL37_SELECTION_REASON_KEY] = 'supported_selection'
    sections[REL37_CANONICAL_FW_KEY] = list(model.selected_frameworks)
    sections[REL37_ORIGINAL_FW_KEY] = list(model.selected_frameworks)
    return sections


def _persist(model):
    sections = _sections(model)
    content = model_to_markdown(model)
    with app_mod.app.app_context():
        db = app_mod.get_db()
        db.execute(
            'INSERT OR IGNORE INTO users '
            '(id, username, password_hash, role, is_active) '
            'VALUES (901, ?, ?, ?, 1)',
            ('rel37ar901', 'x', 'user'),
        )
        cur = db.execute(
            'INSERT INTO strategies '
            '(user_id, domain, org_name, sector, content, language, '
            ' document_title, sections_json, content_json) '
            'VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (
                901, 'Data Management', model.org_name, 'Government',
                content, 'ar', 'Data AR cell fixture',
                json.dumps(sections, ensure_ascii=False),
                json.dumps({'sections': sections}, ensure_ascii=False),
            ),
        )
        sid = cur.lastrowid
        db.commit()
    client = app_mod.app.test_client()
    csrf = 'rel37-ar-cell-csrf'
    with client.session_transaction() as sess:
        sess['user_id'] = 901
        sess['username'] = 'rel37ar901'
        sess['role'] = 'user'
        sess['csrf_token'] = csrf
    return {
        'client': client,
        'headers': {'X-CSRFToken': csrf, 'Content-Type': 'application/json'},
        'strategy_id': sid,
        'model': model,
        'content': content,
        'sections': sections,
    }


def _export(saved, *, org_name, fmt, include_content=True):
    body = {
        'filename': f'rel37_ar_{fmt}',
        'language': 'ar',
        'domain': 'Data Management',
        'doc_type': 'Strategy Document',
        'document_type': 'strategy',
        'artifact_type': 'strategy',
        'generation_mode': 'drafting',
        'strategy_id': saved['strategy_id'],
        'artifact_id': saved['strategy_id'],
        'selected_frameworks': list(saved['model'].selected_frameworks),
        'org_name': org_name,
    }
    if include_content:
        body['content'] = saved['content']
    captured = StdoutCapture()
    with captured:
        with patch('threading.Thread', ImmediateThread):
            resp = saved['client'].post(
                f'/api/generate-{fmt}-async', json=body, headers=saved['headers'])
        submit = resp.get_json(silent=True) or {}
        tid = submit.get('task_id')
        status = {}
        if tid:
            status = saved['client'].get(
                f'/api/export-status/{tid}', headers=saved['headers']
            ).get_json(silent=True) or {}
        raw = b''
        if tid and status.get('status') == 'done':
            raw = saved['client'].get(
                f'/api/export-download/{tid}', headers=saved['headers']
            ).data or b''
    diagnostic = build_export_diagnostic(
        submit_http=resp.status_code,
        submit=submit,
        status=status,
        raw=raw,
        stdout_text=captured.text,
        model_hash=str(saved['model'].model_hash or ''),
        font_path=str(getattr(app_mod, '_ARABIC_PDF_FONT_PATH', '') or ''),
        fmt=fmt,
    )
    if fmt == 'pdf' and (not raw or status.get('status') == 'error'):
        persist_export_diagnostic(
            diagnostic,
            f'{fmt}_{tid or "no_task"}_{diagnostic.get("model_hash", "")[:12]}',
        )
    return resp.status_code, diagnostic, raw


class _Frozen:
    def __init__(self, legacy):
        self.legacy_sections = legacy
        self.canonical_sections = {}


class ArabicCellAssociationTests(unittest.TestCase):
    def setUp(self):
        _ensure_test_env()

    def test_compiled_row_six_present(self):
        model, _secs = _compile_ar()
        self.assertGreaterEqual(len(model.traceability), 6)
        row = model.traceability[5]
        self.assertEqual(row.cells()[0], '6')
        self.assertIn('تصنيف البيانات الشخصية', row.cells()[1])

    def test_frozen_leftover_without_overlay_loses_authority(self):
        model, _secs = _compile_ar()
        full = _sections(model)
        visible = {
            key: value for key, value in full.items()
            if isinstance(value, str) and not str(key).startswith('_')
        }
        self.assertFalse(is_rel37_authoritative(visible))
        restored = overlay_rel37_authority(visible, full)
        self.assertTrue(is_rel37_authoritative(restored))
        leftover = sections_from_frozen_artifact(_Frozen(full))
        self.assertTrue(is_rel37_authoritative(leftover), leftover.keys())

    def test_hash_tamper_not_repaired_by_recompute(self):
        model, _secs = _compile_ar()
        full = _sections(model)
        full[REL37_HASH_KEY] = '0' * 64
        self.assertTrue(rel37_hash_identity_blockers(full))
        overlaid = overlay_rel37_authority({'traceability': 'x'}, full)
        self.assertTrue(overlaid.get('rel37_authority_blocked'))
        self.assertIn('rel37_model_hash_mismatch', overlaid.get('_rel37_render_blocked'))

    def test_repeated_exports_keep_all_relationships(self):
        model, _secs = _compile_ar()
        saved = _persist(model)
        sequence = []
        for org, fmt, include in (
                ('منظمة إدارة البيانات Extra', 'docx', True),
                (model.org_name, 'docx', True),
                (model.org_name, 'docx', False),
                (model.org_name, 'pdf', True),
                ('منظمة إدارة البيانات Extra', 'docx', True),
        ):
            http, body, raw = _export(
                saved, org_name=org, fmt=fmt, include_content=include)
            sequence.append((fmt, http, body, raw))
            self.assertEqual(http, 200, body)
            self.assertTrue(raw, body)
        for fmt, http, body, raw in sequence:
            if fmt == 'docx':
                self.assertTrue(raw.startswith(b'PK'), body)
                blockers = compare_model_to_docx(model, raw)
                self.assertEqual(blockers, [], blockers)
                assoc = compare_traceability_association(model, raw)
                self.assertEqual(assoc, [], assoc)
                inv = inventory_docx_bytes(raw)
                self.assertTrue(inv['tables'])
            else:
                self.assertTrue(raw.startswith(b'%PDF'), body)
                text, meta = extract_pdf_text(raw)
                if not meta.get('reliable') or not str(text or '').strip():
                    self.fail('unreadable PDF extraction must not pass')
                self.assertGreater(arabic_char_count(text), 0, str(text)[:200])

    def test_causal_docx_negatives(self):
        model, _secs = _compile_ar()
        from docx import Document

        source_hash = model.model_hash
        families = expected_rows(model)

        def _docx(trace_rows):
            doc = Document()
            for family, rows in families.items():
                if family.startswith('_'):
                    continue
                use = trace_rows if family == 'traceability' else rows
                if not use:
                    continue
                table = doc.add_table(rows=len(use), cols=max(len(use[0]), 1))
                for ridx, row in enumerate(use):
                    for cidx, value in enumerate(row):
                        table.cell(ridx, cidx).text = value
            buf = io.BytesIO()
            doc.save(buf)
            raw = buf.getvalue()
            # Mutating returned bytes must fail even when hashes are unchanged.
            self.assertEqual(model.model_hash, source_hash)
            return raw

        expected = [list(row.cells()) for row in model.traceability]
        changed = [list(row) for row in expected]
        changed[5][1] = 'قيمة عربية مختلفة'
        blockers = compare_traceability_association(model, _docx(changed))
        self.assertTrue(
            any('docx_row_field_mismatch:traceability:5' in item for item in blockers),
            blockers)

        swapped = [list(row) for row in expected]
        swapped[5][1], swapped[5][2] = swapped[5][2], swapped[5][1]
        blockers = compare_traceability_association(model, _docx(swapped))
        self.assertTrue(
            any('docx_row_field_mismatch:traceability:5' in item for item in blockers),
            blockers)

        moved = [list(row) for row in expected]
        moved[4][1] = expected[5][1]
        blockers = compare_traceability_association(model, _docx(moved))
        self.assertTrue(
            any('docx_row_field_mismatch:traceability:4' in item for item in blockers),
            blockers)

        removed = expected[:5] + expected[6:]
        blockers = compare_traceability_association(model, _docx(removed))
        self.assertTrue(
            any('docx_missing_traceability_row:5' in item for item in blockers),
            blockers)

        duped = [list(row) for row in expected]
        duped[5] = list(expected[4])
        blockers = compare_traceability_association(model, _docx(duped))
        self.assertTrue(
            any(
                'traceability:5' in item or 'traceability_row:5' in item
                for item in blockers
            ),
            blockers,
        )

    def test_compiled_fixture_hash_matches_remote_first_node(self):
        model, _secs = _compile_ar()
        self.assertEqual(model.model_hash, _REMOTE_FIRST_NODE_HASH)

    def test_dejavu_captured_streams_refuse_paragraph_one(self):
        fixture = json.loads(_DEJAVU_FIXTURE.read_text(encoding='utf-8'))
        self.assertEqual(fixture['model_hash'], _REMOTE_FIRST_NODE_HASH)
        visible = fixture['visible']
        actual = fixture['actual']
        paragraphs = fixture['paragraphs']
        all_blockers = []
        for idx, para in enumerate(paragraphs):
            blockers = _paragraph_pdf_blockers(
                idx, para, visible, visible=visible, actual=actual)
            all_blockers.extend(blockers)
        self.assertIn('pdf_environment_actual_visible_disagree:1', all_blockers)
        self.assertFalse(
            any('dejavu' in item.lower() for item in all_blockers),
            all_blockers,
        )
        # ActualText-only must not pass when painted paragraph 1 is incomplete.
        self.assertTrue(paragraphs[1] in actual)
        self.assertNotEqual(visible, actual)

    def test_matching_logical_streams_still_accept_same_paragraphs(self):
        fixture = json.loads(_DEJAVU_FIXTURE.read_text(encoding='utf-8'))
        logical = fixture['actual']
        for idx, para in enumerate(fixture['paragraphs']):
            blockers = _paragraph_pdf_blockers(
                idx, para, logical, visible=logical, actual=logical)
            self.assertEqual(blockers, [], (idx, blockers))

    def test_amiri_dt_captured_streams_refuse_paragraph_one(self):
        fixture = json.loads(_AMIRI_DT_FIXTURE.read_text(encoding='utf-8'))
        self.assertEqual(fixture['model_hash'], _DT_AR_HASH)
        visible = fixture['visible']
        actual = fixture['actual']
        paragraphs = fixture['paragraphs']
        # Classification from the captured representation, not a restored
        # false rejection: the full captured Amiri wrap now associates the
        # same complete statement as Noto. That is a changed positive. The
        # leftover-only negative is a different constructed case: the
        # paragraph-1 line containing "تشمل المحركات" is removed. These
        # are not the same unchanged negative.
        full = []
        for idx, para in enumerate(paragraphs):
            full.extend(_paragraph_pdf_blockers(
                idx, para, visible, visible=visible, actual=actual))
        self.assertEqual(full, [], full)
        leftover_only = '\n'.join(
            line for line in visible.splitlines()
            if 'تشمل المحركات' not in line)
        self.assertNotEqual(leftover_only, visible)
        leftover_blockers = _paragraph_pdf_blockers(
            1, paragraphs[1], leftover_only,
            visible=leftover_only, actual=actual)
        self.assertTrue(leftover_blockers, leftover_blockers)
        self.assertFalse(
            any('amiri' in item.lower() or 'dejavu' in item.lower()
                for item in leftover_blockers),
            leftover_blockers,
        )

    def test_swapped_framework_tokens_still_refused(self):
        fixture = json.loads(_DEJAVU_FIXTURE.read_text(encoding='utf-8'))
        para = fixture['paragraphs'][0]
        mapping = {'NDMO': 'PDPL', 'PDPL': 'NDMO'}
        pattern = re.compile(r'\b(?:NDMO|PDPL)\b')
        swapped = pattern.sub(lambda match: mapping[match.group(0)], para)
        intended = (
            para.replace('NDMO', '\0L\0').replace('PDPL', 'NDMO').replace(
                '\0L\0', 'PDPL')
        )
        self.assertEqual(swapped, intended)
        self.assertEqual(swapped.count('NDMO'), para.count('PDPL'))
        self.assertEqual(swapped.count('PDPL'), para.count('NDMO'))
        self.assertNotIn('HOLD', swapped)
        arabic_src = re.sub(r'[A-Za-z0-9._-]+', '', para)
        arabic_sw = re.sub(r'[A-Za-z0-9._-]+', '', swapped)
        self.assertEqual(arabic_src, arabic_sw)
        blockers = _paragraph_pdf_blockers(
            0, para, swapped, visible=swapped, actual=swapped)
        self.assertTrue(blockers, blockers)
        self.assertFalse(any(
            item in (
                'pdf_bytes_missing', 'pdf_extraction_unreliable',
            ) or item.endswith('painted_unestablished:0')
            for item in blockers
        ), blockers)


if __name__ == '__main__':
    unittest.main()
